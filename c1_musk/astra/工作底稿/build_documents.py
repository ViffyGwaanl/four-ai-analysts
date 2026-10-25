import sys,json
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parent/'python_lib'))
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PI,Pt as PP
from pptx.dml.color import RGBColor as PC
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from lxml.etree import SubElement
D=json.loads(Path('交付/工作底稿/summary.json').read_text());Q=D['q'];BS=D['bs'];S=D['scenarios']
b=lambda n:f'{n/1000:,.2f}';pc=lambda n:f'{n*100:.1f}%'
# Word memo: explicit five-page structure.
doc=Document();sec=doc.sections[0];sec.page_height=Inches(11.69);sec.page_width=Inches(8.27);sec.top_margin=Inches(.73);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.77);sec.right_margin=Inches(.77)
styles=doc.styles
for name in ['Normal','Body Text','Caption']:
 st=styles[name];st.font.name='Arial';st.font.size=Pt(10.5 if name!='Caption' else 8);st._element.rPr.rFonts.set(qn('w:eastAsia'),'PingFang SC');st.font.color.rgb=RGBColor.from_string('203B50')
 styles[name].paragraph_format.space_after=Pt(7);styles[name].paragraph_format.line_spacing=1.19
for name,size in [('Title',28),('Heading 1',21),('Heading 2',12)]:
 st=styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor.from_string('153950');st._element.rPr.rFonts.set(qn('w:eastAsia'),'PingFang SC')
 st.paragraph_format.space_after=Pt(10);st.paragraph_format.space_before=Pt(12)
h=sec.header.paragraphs[0];h.text='ORBIT  /  投资委员会内部讨论                                        2026年10月22日';h.style='Caption'
f=sec.footer.paragraphs[0];f.text='研究与展示用途 · 不构成投资建议                                      ';f.style='Caption';fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
def para(t,style=None):return doc.add_paragraph(t,style)
def title(k,t):para(k,'Subtitle');doc.add_heading(t,0 if k.startswith('01') else 1)
def hd(t):doc.add_heading(t,2)
def table(head,rows,width=None):
 t=doc.add_table(rows=1,cols=len(head));t.autofit=True
 for c,txt in zip(t.rows[0].cells,head):
  c.text=txt;pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'153950');pr.append(sh)
  for r in c.paragraphs[0].runs:r.font.color.rgb=RGBColor(255,255,255);r.font.bold=True;r.font.size=Pt(9)
 for i,row in enumerate(rows):
  cells=t.add_row().cells
  for c,txt in zip(cells,row):
   c.text=str(txt)
   if i%2==0:sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDF4F7');c._tc.get_or_add_tcPr().append(sh)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(4)
    for r in p.runs:r.font.size=Pt(9)
  trPr=t.rows[-1]._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
 return t
def source(t):para('出处：'+t,'Caption')
def page():doc.add_page_break()
title('01 / 结论先行','合并将形成“制造＋连接＋算力”集团，核心约束是资本配置')
para('讨论判断：现有材料支持继续开展交易结构与资金需求尽调，但不足以据此认定合并创造财务协同。两家公司合起来仍以汽车业务贡献收入，以卫星连接支撑当季经营利润，并承担快速扩张的算力资本开支。')
table(['2026Q3有限备考','十亿美元','讨论含义'],[
 ['营收',b(Q['rev']),'已扣除0.221内部Megapack销售'],['经营利润 / 利润率',b(Q['op'])+' / '+pc(Q['op']/Q['rev']),'已扣内部资本化毛利，未计协同'],['自由现金流',b(Q['fcf']),'合并不自动填补资金缺口'],['现金及短期投资 / 有息负债',b(BS['liq'])+' / '+b(BS['debt']),'截至9月30日；债务含融资租赁']])
hd('合并后是什么样的公司')
para('五个正式分部包括Tesla汽车（含服务及其他）、能源发电及储能，以及SpaceX航天、连接、超级智能（SI）。抵销后汽车分部占Q3收入59.4%，SI占17.8%，连接占12.4%，能源占6.3%，航天占4.2%。不能把机器人或Robotaxi单独赋予未披露的收入与利润。')
hd('财务看点与首要风险')
para('连接业务当季经营利润23.67亿美元，占备考整体约68.4%；SI当季已转为经营盈利，但164.44亿美元资本开支大于75.60亿美元收入。Tesla能源对SpaceX的销售体现产业协作，却是合并后需要抵销的内部交易。')
para('两家公司TTM简单相加营收1,356.13亿美元、经营利润27.98亿美元、自由现金流负358.95亿美元；Q3利润改善不等于过去四季现金创造能力改善。短期资金缓冲较大，但必须逐项核对受限资金、承诺投资、融资约束和项目回报。')
hd('本次模型的讨论边界')
para('三种定价方法得到0.2812–0.4234股Tesla／每股SpaceX的零溢价换股区间，属于方法敏感性，不是交易建议或估值结论。尚未确定法律及会计收购方，也没有可审核的购买价分摊、完整债务条款或合并协议。')
source('S3 pp.3–12；T2财务报表及附注13、14；T3 pp.27–31；合并模型「合并备考」「换股比例」。S/T文件全名见第5页。')
page();title('02 / 业务与现金流','收入重心在汽车，利润与投资重心已经分化')
table(['Q3业务','外部收入 / 占比','分部利润（十亿美元）'],[[a['name'],b(a['rev'])+' / '+pc(a['share']),a['profitLabel']+' '+b(a['profit'])] for a in D['segments']])
para('利润不能跨口径直接相加：Tesla汽车与能源披露毛利；SpaceX披露经营利润。汽车毛利45.40亿美元＋能源抵销后毛利9.04亿美元−Tesla总经营费用47.76亿美元＋SpaceX经营利润27.91亿美元＝合并备考经营利润34.59亿美元。')
hd('现金流：两个现金流分类同时抵销，不能虚增FCF')
table(['十亿美元','2026Q3备考','2026前九月备考','TTM简单相加'],[[lab,b(Q[k]),b(D['y'][k]),b(D['ttm'][k])] for k,lab in [('ocf','经营现金流'),('capex','资本开支'),('fcf','自由现金流')]])
para('第三季度内部资本采购现金按SpaceX关联采购现金流9M的4.80亿美元减6M的3.29亿美元，估计为1.51亿美元。将这一笔同时从经营现金流和资本开支扣除，FCF保持不变。与Tesla当季2.21亿美元销售不相等，需要应计与结算对账。')
hd('时间可比性与收入质量')
para('SpaceX招股书已把xAI和X的共同控制重组追溯纳入历史数据，不能再将其重复相加。Cursor于8月14日收购，至9月末贡献6.08亿美元收入；历史TTM没有假设其全年并入。SI增长含新云合同及收购，需核实客户集中度、预付款期限与合同履约。')
para('利润质量也需拆开：SpaceX当季净利润44.04亿美元高于经营利润27.91亿美元，包含其他收益净额10.25亿美元及所得税利益3.90亿美元。不能将净利润跳升直接外推为经常性经营盈利。')
para('FCF口径取披露经营现金流减披露资本开支。SpaceX单列资本化利息、频谱等无形资产支出与返利未重新纳入该口径；FCF因此不是所有投资和并购支出的现金预算。不能简单以期末现金除单季消耗得出确定融资期限。')
source('S1 Note 1、Note 19；S2现金流表及Note 17；S3 pp.3–8、12；T2 Note14；T3 pp.27–31。TTM尚未完整抵销历史关联交易。')
page();title('03 / 抵销与资产负债表','交叉持股要注销，关联交易须按真实对手和资产去向处理')
table(['事项','本期处理','未完成核验'],[
 ['Megapack内部销售','Q3抵销收入221、成本171，内部毛利50；9M对应626、478、148（百万美元）','资产资本化、投用日期、折旧年限、存量毛利与递延税'],
 ['Tesla持有SpaceX','9月末投资账面2,865抵销；约19百万股不获新股；Q3剔除税后重估损失108','3.8百万拆股前股数经5:1调整为估计19；完整证券账户待确认'],
 ['应收及递延收入','应收729含未交货开票588；净额141。递延648扣588后为60','需与SpaceX应付/预付逐项核对，不能直接全部抵销'],
 ['Valor及其他关联方','关联债务12,897和相关利息暂保留；不因“关联”而并表','Q3关联其他收益561的具体对手与交易性质未充分披露']])
hd('账面备考桥：不是购买法权益')
para('9月末两家公司归母权益简单相加为2,660.57亿美元，扣除Tesla对SpaceX投资28.65亿美元及2026年已知未实现内部毛利1.48亿美元，有限备考权益为2,630.44亿美元。未扣2025及以前未知内部毛利；未计相应税收利益、折旧回转、商誉/PPA或融资费用。')
para('总资产备考4,085.49亿美元、总负债1,449.05亿美元，另有Tesla非控股权益及可赎回非控股权益6.00亿美元；与归母权益合计勾稽。现金和有息负债保持1,353.02亿美元与483.67亿美元，净现金869.35亿美元。交叉投资不属于现金，因此抵销它不会减少现金。')
hd('历史交易和合并范围必须对齐')
para('2025年SpaceX披露采购Megapack 506和Cybertruck 131百万美元；Tesla委托书分别披露对SpaceX销售143.3、对xAI销售430.1百万美元，口径与时点并不完全匹配。模型列示差异，没有强行配平，TTM仅简单相加。X广告、飞机及支持服务属于范围内待核对交易；TBC、安保公司和Valor不属于本次合并范围，交易应保留。')
source('T2附注2、13；T4 Related Person Transactions；S1 Note18、F-61；S2 Note17；S3 pp.9、11–12。所有金额和抵销假设可在模型中追溯。')
page();title('04 / 换股与治理','定价方法改变股东分配，双重身份要求独立谈判')
para('结构假设：Tesla作为法律发行主体，每1股外部SpaceX换r股Tesla；原Tesla股东保留每股1股。Tesla原持约19百万SpaceX股先注销。基础股本为Tesla 3,952.649百万股、SpaceX 13,579百万股；后者披露已取整。')
table(['零溢价情景','换股比 r','外部SpaceX / Tesla','Musk合并后'],[[s['name'],f'{s["ratio"]:.4f}',pc(s['spct'])+' / '+pc(s['tpct']),pc(s['musk'])] for s in S])
hd('三种方法的含义')
para('收盘法：使用10月21日SPCX 159.31美元、TSLA 386.47美元。20日法：使用9月24日至10月21日日收盘价按日成交量加权，得到156.34与369.24美元，仅为VWAP近似。倍数法：假设SpaceX EV/TTM营收12×、Tesla 4×；以EV加现金减债务，并在Tesla端加回其SpaceX股权同法价值。倍数未作可比公司校准。')
hd('Musk持股必须统一分子与分母')
para('SpaceX招股书受益持股6,418.548百万股含350百万可行权期权，基础扣除后为6,068.548百万股；假设5月至9月无净变动，对9月期末已发行股份占44.7%。Tesla可支配699.581百万股对全部已发行股份占17.7%；委托书19.8%从分母排除了423.744百万未达标奖励，不能直接套入主情景。两边的限制股均不代表全部已归属。')
para('收盘法合并后Musk占33.5%，较其Tesla口径增加15.8个百分点、较其SpaceX口径减少11.1个百分点。补充情景若SpaceX相关期权全行权，占34.5%；若Tesla已发行未达标奖励全部转为可支配，上界示意38.0%，分母不重复加股。后者未计每股334.09美元抵扣、归属条件或合并触发结果。')
para('SpaceX的IPO后84.4%表决权与经济股份占比不同。合并后多重投票权、激励处理、特别委员会权限和少数股东保护均需协议确认；共同CEO不自动证明两家公司属于同一会计控制主体。')
source('S1 pp.248–249及脚注1、F-61；T4 Ownership of Securities及脚注2、2025 CEO Performance Award；行情来源Stock Analysis，访问2026-10-22；具体网页见README和模型来源登记。')
page();title('05 / 核验清单与来源','进入交易决策前，先解决五项可改变结论的问题')
table(['优先级 / 问题','需要的证据','可能影响'],[
 ['1 现金与承诺','逐实体现金、受限/质押清单；借款条款、担保及控制权变更条款；12–24个月资本预算','可动用资金、融资缺口、偿债安排'],
 ['2 交易与持股','最新股份登记；Musk股权变动；Tesla持有SpaceX明细；所有期权/限制股及归属条款','换股数量、实际稀释与投票权'],
 ['3 内部交易对账','发票、验收、应收应付、预收预付、资本化和折旧；2025历史毛利明细','营收、利润、资产及现金流抵销'],
 ['4 SI与收购质量','云服务客户合同、预付/退款条款、关联客户；Cursor PPA及收购后服务奖励','可持续利润、商誉、未来股份费用'],
 ['5 协同与交易结构','可执行采购/算力协同方案；整合成本；监管和合同同意要求；税务与会计意见','交易是否具有可验证财务利益']])
hd('主要假设与未核实项')
para('模型采用全股票交易、SpaceX统一每股价值、Tesla法律发行主体、零基础溢价；无已确认协同、交易费用、内部毛利税收利益和折旧回转。历史内部毛利、完整往来余额、全部员工激励稀释、PPA及控制权变更费用尚未计量。SpaceX持股数据较旧，20日均价为近似，市场价格与期末股本日期不一致。这些限制均保留在模型和README中。')
hd('材料索引（文件名＋章节）')
for t in [
 'S1：SpaceX_招股书__spaceexplorationtechnologi.htm；财务报表F-3起、Note 18/19、Q1 Note17/18、F-61及pp.248–249。',
 'S2：SpaceX_2026Q2季报__spcx-20260630.htm；财务报表、Note17/18。S3：SpaceX_2026Q3业绩公告__spaceexplorationtechnologi.htm；pp.3–12。',
 'T1：Tesla_2025年报__tsla-20251231.htm；财务报表、Note16。T2：Tesla_2026Q3季报__tsla-20260930.htm；财务报表、Notes2/13/14、Item5。',
 'T3：Tesla_2026Q3业绩公告__exhibit9914.htm；pp.27–31。T4：Tesla_2026股东大会委托书__tm267596-1_def14a.htm；持股、奖励与关联交易章节。']:
 para(t,'Caption')
para('市场数据：Stock Analysis 的 SPCX、TSLA 历史行情页（完整网页地址见模型「来源登记」和README）；收盘日2026-10-21，20日窗口2026-09-24至10-21；访问/取数日期2026-10-22。所有模型数字以随附Excel为准，材料均由用户提供，未对原始披露真实性另行审计。','Caption')
doc.save('交付/投委会备忘录.docx')
# PowerPoint, five slides, all charts editable vector shapes.
r=Presentation();r.slide_width=PI(13.333);r.slide_height=PI(7.5)
C={'ink':'17394E','teal':'07899C','blue':'1264B5','cyan':'22B5D0','purple':'7662DC','orange':'E88A3D','green':'00A88C','muted':'647E8F','bg':'F6FAFC','line':'D9E6ED'}
def box(sl,x,y,w,h,fill,line=None,shape=MSO_SHAPE.RECTANGLE):
 z=sl.shapes.add_shape(shape,PI(x),PI(y),PI(w),PI(h));z.fill.solid();z.fill.fore_color.rgb=PC.from_string(fill);SubElement(z._element.spPr,'{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst');z.line.fill.background() if line is None else None
 if line:z.line.color.rgb=PC.from_string(line)
 return z
def tx(sl,x,y,w,h,text,size=16,color='ink',bold=False,align=None):
 z=sl.shapes.add_textbox(PI(x),PI(y),PI(w),PI(h));tf=z.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';SubElement(p.font._element,'{http://schemas.openxmlformats.org/drawingml/2006/main}ea',typeface='PingFang SC');p.font.size=PP(size);p.font.bold=bold;p.font.color.rgb=PC.from_string(C.get(color,color));p.space_after=PP(6)
  if align is not None:p.alignment=align
 return z
def slide(n,title,sub,foot):
 sl=r.slides.add_slide(r.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=PC.from_string(C['bg']);box(sl,.48,.44,.13,.18,C['teal']);tx(sl,.75,.39,5,.25,'ORBIT  /  投资委员会内部讨论',10,'muted');tx(sl,10,.39,2.7,.25,'2026.10.22    /    '+str(n).zfill(2),10,'muted',align=PP_ALIGN.RIGHT)
 tx(sl,.6,.92,12.1,.65,title,27,'ink',True);tx(sl,.62,1.65,12,.36,sub,11,'muted');box(sl,.6,6.91,12.1,.01,C['line']);tx(sl,.62,7.03,11.9,.28,foot,8,'muted');return sl
sl=slide(1,'合并将形成汽车主导、连接盈利、算力重投入的集团','2026Q3有限备考 · 已抵销Megapack内部销售 · 金额为十亿美元','来源：S3 pp.3–12；T2附注13/14；T3 p.27。文件全名及假设见备忘录第5页 / 合并模型。')
x=.65
for label,key in [('营收','rev'),('经营利润','op'),('自由现金流','fcf')]:
 box(sl,x,2.25,3.93,1.13,'FFFFFF',C['line']);tx(sl,x+.18,2.38,2.8,.5,b(Q[key]),30,'orange' if key=='fcf' else 'ink');tx(sl,x+.18,2.94,3.4,.24,label+'  /  USD bn',11,'muted');x+=4.08
start=.7;totalw=11.95
for a in D['segments']:
 width=a['share']*totalw;box(sl,start,4.13,width,.58,a['color'][1:]);start+=width
tx(sl,.7,3.72,11,.3,'外部收入结构：汽车与服务仍占59.4%',15,'ink',True)
for i,a in enumerate(D['segments']):
 x=.72+i*2.45;box(sl,x,5.0,.12,.12,a['color'][1:]);tx(sl,x+.2,4.95,2.15,.3,a['name'],11,'ink');tx(sl,x+.2,5.35,2.15,.45,pc(a['share']),23,'ink')
box(sl,.7,6.15,11.95,.46,'E9F3F6');tx(sl,.87,6.24,11.6,.25,'收入规模扩大 ≠ 自由现金流改善；尚未确认任何可计入模型的经营协同。',13,'teal',True)
sl=slide(2,'卫星连接支撑当季利润，SI吸收最多资本','利润桥统一到经营利润口径 · 十亿美元 · Tesla分部先列毛利再扣总经营费用','来源：S3分部表pp.3–7；T2附注14；T3 p.27。SI资本开支16.444；利润不是现金。')
values=[4.540,.904,-4.776,-.007,2.367,.431,3.459];labs=['汽车毛利','能源毛利¹','Tesla经营费用','航天经营利润','连接经营利润','SI经营利润','备考经营利润']
baseY=5.75;scale=.40;cum=0
for i,(v,l) in enumerate(zip(values,labs)):
 x=.7+i*1.35;start=0 if i==6 else cum;end=v if i==6 else cum+v
 y=baseY-max(start,end)*scale;h=max(abs(v)*scale,.035);fill=C['teal'] if i==6 else C['orange'] if v<0 else C['cyan'];box(sl,x,y,1.02,h,fill);tx(sl,x-.1,y-.36,1.24,.28,f'{v:+.3f}' if i<6 else f'{v:.3f}',11,'ink',True,PP_ALIGN.CENTER);tx(sl,x-.1,5.95,1.25,.55,l,10,'muted',align=PP_ALIGN.CENTER)
 if i<6:cum+=v
box(sl,10.35,2.45,2.26,3.83,'EAF3F7');tx(sl,10.55,2.7,1.9,.35,'SI / Q3',13,'teal',True);tx(sl,10.55,3.22,1.9,.6,'16.44',30,'orange');tx(sl,10.55,3.89,1.9,.32,'资本开支 / USD bn',10,'muted');tx(sl,10.55,4.57,1.9,.55,'2.2×',29,'ink');tx(sl,10.55,5.17,1.9,.55,'资本开支 / 收入\n收购后期间不完全可比',10,'muted')
tx(sl,.7,2.3,8.8,.48,'Tesla毛利与SpaceX经营利润须经费用桥接，不能直接相加。',14,'ink',True);tx(sl,.7,6.56,9.3,.22,'¹ 能源毛利已扣0.050内部毛利；无折旧回转、税收利益或协同假设。',9,'muted')
sl=slide(3,'1,353亿美元流动资产缓冲，仍需面对持续资金消耗','左：当季现金创造与资本投入；右：9月末现金和融资租赁在内的有息负债 · 十亿美元','来源：S2/S3现金流表；T3 pp.28–31。FCF为OCF减披露资本开支，未覆盖全部投资支出。')
labels=['Tesla','SpaceX','有限备考合计'];ocfs=[6.263,9.631,15.743];caps=[4.732,19.593,24.174]
for i,l in enumerate(labels):
 y=2.65+i*1.12;tx(sl,.7,y,1.8,.3,l,13,'ink',True);box(sl,2.3,y,ocfs[i]*.18,.25,C['cyan']);box(sl,2.3,y+.36,caps[i]*.18,.25,C['orange']);tx(sl,2.36+ocfs[i]*.18,y-.02,.85,.3,f'{ocfs[i]:.2f}',11,'ink');tx(sl,2.36+caps[i]*.18,y+.34,.85,.3,f'{caps[i]:.2f}',11,'ink')
box(sl,.75,6.1,.13,.13,C['cyan']);tx(sl,.98,6.04,1.8,.3,'经营现金流',11,'muted');box(sl,2.8,6.1,.13,.13,C['orange']);tx(sl,3.03,6.04,1.8,.3,'资本开支',11,'muted')
box(sl,8.2,2.3,4.4,4.27,'FFFFFF',C['line']);
for i,(label,k,col) in enumerate([('现金及短期投资','liq','ink'),('有息负债','debt','orange'),('净现金','netcash','teal')]):
 tx(sl,8.48,2.55+i*1.13,2.4,.26,label,11,'muted');tx(sl,8.48,2.88+i*1.13,3.75,.54,b(BS[k]),30,col)
tx(sl,.7,6.6,11.8,.2,'9M26 FCF：−33.09；TTM简单相加FCF：−35.90。不能以单季消耗机械推算融资期限。',10,'muted')
sl=slide(4,'三种定价基础，使外部SpaceX股东占比落在49%–59%','基准溢价0%；每1股外部SpaceX换取Tesla股数；原Tesla股份保留；交叉持股先注销','来源：M1/M2行情；S1持股章节；T4持股脚注2。20日为收盘成交量加权近似；倍数12×/4×均为假设。')
for i,s in enumerate(S):
 x=.7+i*4.07;box(sl,x,2.3,3.92,3.85,'FFFFFF',C['line']);tx(sl,x+.2,2.53,3.5,.34,s['name'],15,'ink',True);tx(sl,x+.2,3.15,3.3,.63,f'{s["ratio"]:.4f}',36,'teal');tx(sl,x+.2,3.88,3.3,.27,'股Tesla / 每股SpaceX',10,'muted');barw=3.5;box(sl,x+.2,4.42,barw*s['tpct'],.28,C['blue']);box(sl,x+.2+barw*s['tpct'],4.42,barw*s['spct'],.28,C['cyan']);tx(sl,x+.2,4.85,3.5,.32,'Tesla '+pc(s['tpct'])+'  /  SpaceX '+pc(s['spct']),11,'muted');tx(sl,x+.2,5.46,3.5,.34,'Musk股份占比  '+pc(s['musk']),15,'ink',True)
tx(sl,.76,6.46,11.8,.25,'Musk股份占比不等于表决权；基础持股包含限制股但剔除尚未行权期权与不可支配未达标奖励。',11,'muted')
sl=slide(5,'先完成持股、对账与资金核验，再判断交易是否创造价值','现阶段的模型提供可复核的讨论起点；不包含购买价分摊，也没有假设已实现协同','来源：T2附注2/13；T4奖励与关联交易；S1/S2关联交易。研究展示用途，不构成投资建议。')
risks=[('01','股份与控制','确认期末股份登记、期权及限制股；\n明确投票权和奖励的合并触发条件。','直接改变换股与稀释'),('02','内部交易与资产','对账729应收与648递延收入；\n核对历史资本化毛利、折旧及税务。','改变利润与净资产'),('03','现金与承诺','核实可动用现金、债务约束、\nSI投资预算及Cursor后续股份费用。','决定融资与整合可行性')]
for i,(n,t,desc,imp) in enumerate(risks):
 y=2.35+i*1.31;box(sl,.7,y,.61,.61,'E4F0F4');tx(sl,.81,y+.13,.4,.3,n,17,'teal',True);tx(sl,1.58,y+.02,2.45,.4,t,18,'ink',True);tx(sl,4.08,y,5.2,.75,desc,14,'muted');box(sl,9.9,y+.05,2.7,.53,'FFFFFF',C['line']);tx(sl,10.05,y+.18,2.4,.28,imp,12,'teal',True)
box(sl,.7,6.44,11.9,.33,'E8F2F6');tx(sl,.84,6.48,11.5,.25,'决策门槛：可验证的协同方案 ＋ 完整现金预算 ＋ 可审核的公平换股与治理安排。',12,'teal',True)

for slide_item in r.slides:
 for shape_item in slide_item.shapes:
  for effect in shape_item._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef'):
   effect.set('idx','0')
r.save('交付/投委会汇报.pptx');print('Created memo (5 designed pages) and deck (5 slides).')
