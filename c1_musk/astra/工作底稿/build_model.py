import sys,json
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parent/'python_lib'))
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.comments import Comment
from openpyxl.chart import BarChart,Reference
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.workbook.properties import CalcProperties
D=json.loads(Path('交付/工作底稿/data.json').read_text()); M=json.loads(Path('交付/工作底稿/market.json').read_text())
w=Workbook();w.remove(w.active)
navy='152E49'; teal='087E8B'; blue='E8F1FF'; green='EAF5F1'; yellow='FFF2CC'; muted='587083'
def sheet(name,title,heads,widths):
 s=w.create_sheet(name);s.append([title]);s.append(['单位：百万美元（股数为百万股；股价为美元/股），除非另有注明 | 取数：2026-10-22'])
 s.append(heads);s.freeze_panes='B4';s.sheet_view.showGridLines=False
 for i,wi in enumerate(widths,1):s.column_dimensions[chr(64+i)].width=wi
 s.row_dimensions[1].height=32;s.row_dimensions[2].height=26;s.row_dimensions[3].height=29
 return s
def note(s,row,text,end=8):
 s.cell(row,1,text);s.merge_cells(start_row=row,start_column=1,end_row=row,end_column=end);s.row_dimensions[row].height=44;s.cell(row,1).alignment=Alignment(wrap_text=True,vertical='center')
s=sheet('说明','SPACE × TESLA | 并购讨论模型',['项目','说明'],[27,130])
info=[('模型性质','研究和展示用途，不构成投资建议。合并尚未确定；此表是账面加总加有限抵销，不是符合监管要求的完整购买法备考财务报表。'),('期间','最近四季度：2025Q4、2026Q1、Q2、Q3；完整财年：2025；TTM=四季度之和。资产负债项目取各期末，TTM列取2026-09-30。'),('SpaceX口径','招股书对SpaceX、xAI、X按共同控制重组追溯合并；沿用其历史比较口径。AI在Q3称SI。Cursor自2026-08-14并入，历史TTM未做全年收购备考。'),('分部利润','SpaceX按分部经营利润；Tesla正式分部为汽车（含服务及其他）和能源，披露分部毛利。不得直接把两类利润相加；备考另列Tesla总经营费用桥接。'),('净利润','主表净利润为合并净利润；另列归属股东净利润。SpaceX Q1两者相差671，不能将归属股东数字当合并净利润。'),('现金与债务','现金及短期投资不含受限现金、数字资产、Tesla长期SpaceX投资；有息负债含融资租赁，不含经营租赁。SpaceX历史可赎回优先股独列，不计普通权益或债务。'),('自由现金流','FCF=经营现金流−披露资本开支。SpaceX资本开支用PP&E采购（不另扣返利、不并入单列资本化利息或频谱支出）；Tesla用披露净资本开支。属分析口径，非全部资金需求。'),('单季推导','SpaceX Q4=FY2025−9M2025；Q2经营现金流/资本开支=6M−Q1；Q3=9M−6M。所有差分由公式计算。'),('合并范围','只合并Tesla与SpaceX及既有子公司（含xAI/X/Cursor）；不额外纳入TBC、Neuralink、Valor或个人安保公司。关联不等于属于合并范围。'),('抵销边界','Q3与9M26按Tesla披露Megapack销售/成本抵销；对资本化内部毛利按尚未对外实现、无折旧回转假设处理。Q4 2025内部毛利未知，不把调整后的9M冒充完整TTM。'),('资产负债备考','仅确认可量化交叉投资与2026内部毛利抵销。往来抵销设可编辑的待核对假设，默认0；未对账≠不存在。没有PPA、商誉重估、融资或递延税全面调整。'),('换股方向','假设Tesla作为法律发行主体，原Tesla每股保留1股；每1股外部SpaceX换取r股Tesla。Tesla自持SpaceX股数先注销，避免向自己发行。法律/会计收购方并未确定。'),('股本与奖励','基础用已发行期末股本，包含已发行但未达标限制股；不使用EPS加权平均股本。Tesla未达标423.744百万股不计入Musk可支配分子；SpaceX350百万可行权期权从分子扣除。'),('股权与表决权','模型只计算按统一分母的股份占比，不代表合并后的投票权。SpaceX招股书84.4%为IPO后表决权，不能与股份占比混用。Musk SpaceX持股沿用5月披露，后续变动待核实。'),('市场数据','采用2026-10-21最近完整交易日收盘，和2026-09-24至10-21的20个交易日。10-22盘中价格不用于定价。收盘价×成交量加权仅为VWAP近似，并非逐笔成交VWAP。'),('三种方法','①收盘市值；②20日收盘成交量加权近似；③假设EV/TTM营收倍数（SpaceX12×、Tesla4×），由EV加现金减债务并加非经营交叉持股还原股权价值。倍数无可比公司校准。'),('颜色','蓝底=披露事实/外部市场输入；黄底=假设或估计输入；浅绿=公式。每个输入标文件/表名；每个公式格批注给出取数路径。'),('刷新与输出','修改假设后Excel公式自动重算；Word/PPT/网页是基准日快照。导出脚本刷新数据后，还须同步复核静态叙述、图表标题与页数，不能视为全套实时联动。'),('访问日期','材料及行情访问/取数日期均为2026-10-22。来源登记含文件全名及章节。所有模型假设集中在假设页。')]
for a,b in info:s.append([a,b]);s.row_dimensions[s.max_row].height=56
src=sheet('来源登记','来源与取数索引',['ID','文件 / 网页','章节 / 用途','访问日期'],[12,85,80,18])
for k,v in D['files'].items():src.append([k,v,'原始材料；逐项章节见原始数据页', '2026-10-22'])
src.append(['M1','https://stockanalysis.com/stocks/spcx/history/','20个已完成交易日 Close / Volume；S&P Global数据，网页转录','2026-10-22'])
src.append(['M2','https://stockanalysis.com/stocks/tsla/history/','20个已完成交易日 Close / Volume；未使用盘中行','2026-10-22'])
src.append(['M3','https://chartexchange.com/symbol/nasdaq-spcx/historical/','10月21日收盘159.31交叉核对；成交量与M1有差异，统一采用M1口径','2026-10-22'])
src.append(['背景','任务.md；T2 Item 5 Other Information','Tesla特别委员会成立可由季报核实；交易未保证完成','2026-10-22'])
raw=sheet('原始数据','可追溯披露输入',['数据ID','公司','期间','字段','披露值','来源文件','章节 / 表名','取数日期'],[35,12,12,24,20,75,95,18]);loc={}
for ob in D['raw']:
 raw.append([ob['id'],ob['co'],ob['period'],ob['metric'],ob['value'],ob['source'],ob['section'],'2026-10-22']);loc[ob['id']]=raw.max_row
 raw.cell(raw.max_row,5).comment=Comment(ob['source']+' | '+ob['section']+' | 2026-10-22','来源')
raw.auto_filter.ref=raw.dimensions

def R(co,p,k):return f"'原始数据'!E{loc[f'{co}_{p}_{k}']}"
ass=sheet('假设','独立于披露事实的建模假设',['假设名称','数值','单位','说明'],[42,23,18,115]);AL={}
def A(k,label,v,unit,why):
 ass.append([label,v,unit,'【假设】'+why]);AL[k]=ass.max_row;ass.cell(ass.max_row,2).fill=PatternFill('solid',fgColor=yellow);ass.row_dimensions[ass.max_row].height=44
A('premium','支付给外部SpaceX股东的溢价',0,'比例','以各方法SpaceX基础每股价格加成；Tesla每股价值不加成。')
A('smult','SpaceX EV / TTM营收',12,'倍','纯演示倍数，未作可比公司校准，不是估值意见。')
A('tmult','Tesla EV / TTM营收',4,'倍','纯演示倍数；不将负净利润用于P/E。')
A('synergy','Q3可实现经营协同',0,'百万美元','尚无已签约可验证协同。')
A('fee','Q3一次性交易费用',0,'百万美元','尚无交易方案，基准不计；可编辑。')
A('tax','内部毛利/协同净调整税率',0,'比例','无完整税务资料，基准不确认税收利益；非适用法定税率。')
A('deprq','Q3内部毛利折旧回转',0,'百万美元','未获资产投用日期/年限；基准0，可能低估利润。')
A('depr9','9M26内部毛利折旧回转',0,'百万美元','不含2025及以前内部毛利余额。')
A('cashq','Q3内部资本采购对应已结算现金',f'={R("X","9M26","related_capex")}-{R("X","6M26","related_capex")}','百万美元','以SpaceX关联PP&E现金流差分151估计对应Tesla收款；与221销售不能直接等同，待现金对账。')
A('cash9','9M26内部资本采购对应已结算现金',f'={R("X","9M26","related_capex")}','百万美元','以480关联PP&E现金流近似；重分类同时减少经营现金流和资本开支，FCF不变。')
A('ar','已匹配往来应收/应付抵销',0,'百万美元','双边余额未核对，默认不抵销；不是认定无往来。Tesla账面729含588未交货开票，净141仅为核对起点。')
A('advance','已匹配预付/递延收入抵销',0,'百万美元','需核实对手预付款；不得自动把648或60全额认作可抵销余额。')
A('oldprofit','2025及以前未实现内部毛利余额',0,'百万美元','缺少成本/折旧明细；0为未计量占位，不表示实际为0。')
A('sdelta','Musk SpaceX披露后持股净变动',0,'百万股','5月持股沿用至9月，未取得完整持股变动；按基础无变动。')
A('sexercise','SpaceX期权行权比例（补充情景）',1,'比例','只用于持股补充情景；不纳入主换股股本，未建行权现金/库存股法。')
A('tvesting','Tesla未达标股份转为Musk可支配比例',1,'比例','仅作上界示意，非合并必然触发归属；未计334.09美元/股抵扣或现金支付影响。')
A('sp_other','SpaceX其他非经营净资产估值调整',0,'百万美元','除现金/债务外不另估数字资产、少数投资；倍数法简化。')
A('ts_other','Tesla其他非经营净资产估值调整',0,'百万美元','SpaceX股权单独加入；数字资产、少数股权、经营租赁等暂不另调整。')
def AR(k):return f"'假设'!B{AL[k]}"
metrics=[('rev','营收'),('cost','营业成本'),('gross','毛利'),('op','经营利润'),('net','合并净利润'),('attrib','归属股东净利润'),('ocf','经营现金流'),('capex','资本开支（正数）'),('fcf','自由现金流'),('liq','现金及短期投资'),('debt','有息负债含融资租赁'),('equity','股东权益（归母）'),('assets','总资产'),('liab','总负债'),('netcash','净现金'),('margin','经营利润率')]
mr={k:i+4 for i,(k,l) in enumerate(metrics)}
periods=['FY25','Q425','Q126','Q226','Q326','TTM'];cmap=dict(zip(periods,'BCDEFG'))
for co,name in [('S','SpaceX财务'),('T','Tesla财务')]:
 s=sheet(name,name+' | 单季、财年与TTM',['项目','FY2025','2025Q4','2026Q1','2026Q2','2026Q3','TTM至26Q3','口径 / 来源'],[34,20,20,20,20,20,22,95])
 for k,label in metrics:
  row=mr[k];s.cell(row,1,label)
  for p,col in cmap.items():
   if k in ['gross','fcf','debt','liq','netcash','margin']:
    if k=='gross':expr=f'{col}{mr["rev"]}-{col}{mr["cost"]}'
    if k=='fcf':expr=f'{col}{mr["ocf"]}-{col}{mr["capex"]}'
    if k=='netcash':expr=f'{col}{mr["liq"]}-{col}{mr["debt"]}'
    if k=='margin':expr=f'{col}{mr["op"]}/{col}{mr["rev"]}'
    if k in ['debt','liq']:
     pp='Q326' if p=='TTM' else ('FY25' if co=='S' and p=='Q425' else p)
     if k=='liq' and co=='T' and pp!='FY25':expr=R(co,pp,'liq')
     else:expr=R(co,pp,'dc' if k=='debt' else 'cash')+'+'+R(co,pp,'dl' if k=='debt' else 'st')
   elif p=='TTM':expr=f'SUM(C{row}:F{row})' if k not in ['equity','assets','liab'] else f'F{row}'
   elif co=='S' and p=='Q425':expr=f'B{row}' if k in ['equity','assets','liab'] else R(co,'FY25',k)+'-'+R(co,'9M25',k)
   elif co=='S' and p in ['Q226','Q326'] and k in ['ocf','capex']:expr=(R(co,'6M26',k)+'-'+R(co,'Q126',k)) if p=='Q226' else R(co,'9M26',k)+'-'+R(co,'6M26',k)
   else:expr=R(co,p,k)
   s[f'{col}{row}']='='+expr
   if k=='margin':s[f'{col}{row}'].number_format='0.0%'
  s.cell(row,8,'公式追溯至原始数据页；TTM流量求和/存量取期末。'+('债务含融资租赁，不含经营租赁。' if k=='debt' else ''))
 note(s,22,'每格公式批注标引用路径。SpaceX历史已含xAI/X共同控制追溯合并；Cursor仅自2026-08-14并入。Tesla主表净利润含非控股权益，归属股东净利润另列。')
 if co=='S':
  s.cell(24,1,'可赎回优先股（夹层）')
  for p,col in cmap.items():s[f'{col}24']='='+R('S',('Q326' if p=='TTM' else 'FY25' if p=='Q425' else p),'pref')
  note(s,26,'可赎回优先股单列，2026Q2后为0。2025末权益2,573不含38,752夹层优先股；不能将IPO前后权益变化全部解释为经营盈利。')
seg=sheet('分部财务','正式披露分部与利润桥接',['业务 / 指标','FY2025','2025Q4','2026Q1','2026Q2','2026Q3','TTM至26Q3','利润口径 / 来源'],[40,18,18,18,18,18,20,105]);SR={}
for co,b,label in [('S','space','航天'),('S','conn','连接'),('S','si','超级智能 SI'),('T','automotive','汽车（含服务及其他）'),('T','energy','能源发电及储能')]:
 for k,lab in [('rev','收入'),('cost','成本'),('gross','毛利')]+([('op','经营利润')] if co=='S' else []):
  row=seg.max_row+1;SR[(b,k)]=row;seg.cell(row,1,label+' / '+lab)
  for p,col in cmap.items():
   if k=='gross':expr=f'{col}{SR[b,"rev"]}-{col}{SR[b,"cost"]}'
   elif p=='TTM':expr=f'SUM(C{row}:F{row})'
   elif co=='S' and p=='Q425':expr=R(co,'FY25',b+'_'+k)+'-'+R(co,'9M25',b+'_'+k)
   elif b=='automotive':expr=R(co,p,'auto_'+k)+'+'+R(co,p,'serv_'+k)
   else:expr=R(co,p,b+'_'+k)
   seg[f'{col}{row}']='='+expr
  seg.cell(row,8,('S1 Note19/18；S3 分部表；AI更名SI。' if co=='S' else 'T1 Note16；T2 Note14；T3 p27。汽车含服务及其他，未披露分部经营利润。'))
for label,exprtype in [('Tesla经营费用（未分配到分部）','opex'),('合计经营利润（桥接后）','op'),('营收勾稽差（应为0）','check')]:
 row=seg.max_row+1;SR[exprtype]=row;seg.cell(row,1,label)
 for col in 'BCDEFG':
  if exprtype=='opex':expr=f"'Tesla财务'!{col}{mr['gross']}-'Tesla财务'!{col}{mr['op']}"
  elif exprtype=='op':expr='+'.join(f'{col}{SR[b,"op"]}' for b in ['space','conn','si'])+f'+{col}{SR["automotive","gross"]}+{col}{SR["energy","gross"]}-{col}{SR["opex"]}'
  else:expr='+'.join(f'{col}{SR[b,"rev"]}' for b in ['space','conn','si','automotive','energy'])+f"-'SpaceX财务'!{col}{mr['rev']}-'Tesla财务'!{col}{mr['rev']}"
  seg[f'{col}{row}']='='+expr
note(seg,seg.max_row+3,'Tesla分部利润为毛利；SpaceX分部利润为经营利润。先扣Tesla未分配经营费用，再与SpaceX经营利润相加。不可将两种分部利润直接相加。')
market=sheet('市场数据','行情输入 | 只使用完整交易日',['日期','SPCX收盘','SPCX成交量','TSLA收盘','TSLA成交量','SPCX收盘×量','TSLA收盘×量','出处'],[18,20,23,20,23,25,25,90])
for i,a in enumerate(M,4):
 market.append([a['date'],a['s_close'],a['s_vol'],a['t_close'],a['t_vol'],f'=B{i}*C{i}',f'=D{i}*E{i}','M1/M2：Stock Analysis历史行情；2026-10-22访问'])
market.cell(25,1,'20日加权近似');market['B25']='=SUM(F4:F23)/SUM(C4:C23)';market['D25']='=SUM(G4:G23)/SUM(E4:E23)'
note(market,27,'近似值=Σ（日收盘价×日成交量）/Σ（日成交量）。不是逐笔VWAP；20日窗口2026-09-24至2026-10-21。若获交易所逐笔数据，应替换。')
# related party review table
rp=sheet('关联交易','交易对账与抵销路线',['交易 / 余额','期间','披露金额','比较金额','单位','来源','处理','局限'],[40,20,20,20,18,60,95,85])
items=[
 ('Megapack销售 / 成本','Q3 2026',R('X','Q326','megapack_rev'),R('X','Q326','megapack_cost'),'T2 Note13','抵销221收入、171成本，50内部毛利从资产及利润扣除','设备在SpaceX资本化；并非冲收入又重复冲221费用'),
 ('Megapack销售 / 成本','9M26',R('X','9M26','megapack_rev'),R('X','9M26','megapack_cost'),'T2 Note13','抵销626收入、478成本；内部毛利148','未考虑折旧、税务回转和历史毛利，单列假设'),
 ('SpaceX关联资本采购现金流','9M26 / 6M26',R('X','9M26','related_capex'),R('X','6M26','related_capex'),'S3 p12；S2 Note17','151为Q3差分；经营现金流与资本开支等额重分类','9M现金480 vs Tesla收入626；需区分应计/现金、验收及预付'),
 ('对SpaceX应收 / 未交货开票','2026-09-30',R('X','SEP26','ar'),R('X','SEP26','undelivered'),'T2 Note13','729−588=141为已交货净核对起点；双边确认后抵销应收/应付','不能将729全部视为已交货；也不能将141当已确认SpaceX应付'),
 ('递延收入 / 未交货开票','2026-09-30',R('X','SEP26','deferred'),R('X','SEP26','undelivered'),'T2 Note13','648−588=60；按已收预付款与对方预付资产逐项核对','当前516/长期132；60并不自动等于可抵销现金预付'),
 ('Tesla持有SpaceX投资','2026-09-30',R('X','SEP26','investment'),None,'T2 Note2；S1 F61','账面投资2,865与权益抵销；模型注销约19百万SpaceX股','股数为拆股后估算；投资账面值不是用于定价的市值'),
 ('SpaceX股权重估收益','Q3 / 9M26',R('X','Q326','fv_gain'),R('X','9M26','fv_gain'),'T2 Note2；T3 p30','税前−142/863应剔除；实际税后−108/655用于净利润调整','合并后不能确认自身权益工具价格变动收益'),
 ('Valor关联债务','2026-09-30',R('X','SEP26','valor_dc'),R('X','SEP26','valor_dl'),'S3 p11；S2 Note17','保留：Valor不在合并范围；总12,897','Q3关联对手归属需以完整Q3季报再确认'),
 ('SpaceX关联利息 / 其他收益','Q3 2026',R('X','Q326','related_interest'),R('X','Q326','related_other_income'),'S3 p9','351利息、561其他收益不可因“关联”二字全部抵销','Q3公告未充分标明561的对手与交易性质'),
 ('Megapack / Cybertruck采购','FY25',R('X','FY25','megapack_purchase'),R('X','FY25','cybertruck_purchase'),'S1 Note18 F57','追溯对账历史资产成本及未实现毛利','与Tesla披露收入存在时点/分类差异，禁止直接硬配'),
 ('Tesla向SpaceX / xAI销售','FY25',R('X','FY25','tesla_sales_spacex'),R('X','FY25','tesla_sales_xai'),'T4 Related Person Transactions','历史收入需去重抵销；未进Q3/9M26本期金额','FY25收入573.4与637采购不一致；历史成本明细不足'),
 ('商业支持支出：SpaceX / xAI','FY25',R('X','FY25','tesla_exp_spacex'),R('X','FY25','tesla_exp_xai'),'T4 Related Person Transactions','合并范围内费用/收入成对抵销，核对重叠','11.4与4.0可能与其他披露包含关系待核'),
 ('SpaceX飞机 / X广告','FY25',R('X','FY25','aircraft'),R('X','FY25','advertising_x'),'T4 Related Person Transactions','同在合并范围，核对后抵销费用与对方收入','不在模型的2026已知抵销额中'),
 ('TBC / Musk安保公司支出','FY25',R('X','FY25','tbc'),R('X','FY25','security'),'T4 Related Person Transactions','保留外部关联方交易；不合并这些企业','关联交易定价与治理需独立审查'),
 ('Redwood支出 / 废料成本冲减','FY25',R('X','FY25','redwood_exp'),R('X','FY25','redwood_scrap'),'T4 Related Person Transactions','保留；Redwood不在范围内，属董事关联公司','不属于Musk控制企业，单列治理事项'),
 ('Terafab合作','2026',None,None,'S1 F61 Tesla Collaboration','仅合作安排，不另添加一个独立收入分部','项目投资分担、知识产权与资金承诺未量化')]
for row in items:
 label,p,x,y,source,action,lim=row;rp.append([label,p,'='+x if x else '未量化','='+y if y else '—','百万美元',source,action,lim]);rp.row_dimensions[rp.max_row].height=64
# tax-adjusted FV observations
for co,p,k,v in [('X','Q326','fv_aftertax',-108),('X','9M26','fv_aftertax',655)]:
 id=f'{co}_{p}_{k}';raw.append([id,co,p,k,v,D['files']['T3'],'p.30 Reconciliation: SpaceX equity investment gain/loss, net of tax','2026-10-22']);loc[id]=raw.max_row
pf=sheet('合并备考','合并备考 | 已知抵销与未量化事项分开',['项目','Tesla','SpaceX','简单相加','抵销 / 假设调整','备考结果','单位','解释'],[37,21,21,23,26,23,18,110])
PR={}
for group,p,caption in [('q','Q326','2026Q3'),('y','9M26','2026前九月'),('t','TTM','TTM：仅简单相加')]:
 note(pf,pf.max_row+1,caption)
 for k,label in [('rev','营收'),('gross','毛利'),('op','经营利润'),('net','合并净利润'),('ocf','经营现金流'),('capex','资本开支'),('fcf','自由现金流')]:
  row=pf.max_row+1;PR[group,k]=row;pf.cell(row,1,label)
  for co,col in [('T','B'),('S','C')]:
   if p in ['Q326','TTM']:expr=f"'{('Tesla' if co=='T' else 'SpaceX')}财务'!{('F' if p=='Q326' else 'G')}{mr[k]}"
   elif k=='gross':expr=R(co,p,'rev')+'-'+R(co,p,'cost')
   elif k=='fcf':expr=R(co,p,'ocf')+'-'+R(co,p,'capex')
   else:expr=R(co,p,k)
   pf[f'{col}{row}']='='+expr
  pf[f'D{row}']=f'=B{row}+C{row}';pf.cell(row,7,'百万美元')
  if group=='t':pf[f'E{row}']='未完整抵销';pf[f'F{row}']=f'=D{row}';pf.cell(row,8,'Q4 2025内部毛利/交易缺少明细；本列仅展示加总，不代表完整合并TTM。')
  else:
   profit=f'({R("X",p,"megapack_rev")}-{R("X",p,"megapack_cost")})';dep=AR('deprq' if group=='q' else 'depr9')
   if k=='rev':adj='-'+R('X',p,'megapack_rev')
   elif k=='gross':adj='-'+profit+'+'+dep
   elif k=='op':adj='-'+profit+'+'+dep+(f'+{AR("synergy")}-{AR("fee")}' if group=='q' else '')
   elif k=='net':adj=f'(-{profit}+{dep}'+(f'+{AR("synergy")}-{AR("fee")}' if group=='q' else '')+f')*(1-{AR("tax")})-{R("X",p,"fv_aftertax")}'
   elif k in ['ocf','capex']:adj='-'+AR('cashq' if group=='q' else 'cash9')
   else:adj=f'E{PR[group,"ocf"]}-E{PR[group,"capex"]}'
   pf[f'E{row}']='='+adj;pf[f'F{row}']=f'=D{row}+E{row}'
   pf.cell(row,8,{'net':'剔除已披露税后投资重估损益；内部毛利及协同用假设税率。','fcf':'内部现金经营/投资重分类同时调整OCF和资本开支，FCF不因此改善。','op':'内部资本化毛利未对外实现；折旧回转、协同与费用按假设。'}.get(k,'含已知Megapack抵销；其他未量化关联交易尚待对账。'))
note(pf,pf.max_row+2,'2026-09-30资产负债表 | 账面基础，未做收购价分摊');BS={}
for k,label in [('liq','现金及短期投资'),('debt','有息负债'),('netcash','净现金'),('assets','总资产'),('liab','总负债'),('equity','归母股东权益')]:
 row=pf.max_row+1;BS[k]=row;pf.cell(row,1,label)
 for co,col in [('Tesla','B'),('SpaceX','C')]:pf[f'{col}{row}']=f"='{co}财务'!F{mr[k]}"
 pf[f'D{row}']=f'=B{row}+C{row}'
 if k in ['liq','debt','netcash']:adj='0'
 elif k=='assets':adj=f'-{R("X","SEP26","investment")}-({R("X","9M26","megapack_rev")}-{R("X","9M26","megapack_cost")})+{AR("depr9")}-{AR("oldprofit")}-{AR("ar")}-{AR("advance")}'
 elif k=='liab':adj=f'-{AR("ar")}-{AR("advance")}'
 else:adj=f'-{R("X","SEP26","investment")}-({R("X","9M26","megapack_rev")}-{R("X","9M26","megapack_cost")})+{AR("depr9")}-{AR("oldprofit")}'
 pf[f'E{row}']='='+adj;pf[f'F{row}']=f'=D{row}+E{row}';pf.cell(row,8,'资产负债暂不确认相关递延税；未核对往来默认不调整。股东权益不含非控股权益。')
note(pf,pf.max_row+2,'Q3营收结构按已知221内部销售抵销后计；Tesla能源分部利润为毛利，SpaceX为经营利润。')
BR={}
for b,label in [('automotive','汽车（含服务及其他）'),('energy','能源发电及储能'),('space','航天'),('conn','连接'),('si','超级智能 SI')]:
 row=pf.max_row+1;BR[b]=row;pf.cell(row,1,label);pf[f'B{row}']=f"='分部财务'!F{SR[b,'rev']}"+(f'-{R("X","Q326","megapack_rev")}' if b=='energy' else '');pf[f'C{row}']=f'=B{row}/F{PR["q","rev"]}';pf[f'C{row}'].number_format='0.0%'
 pf[f'D{row}']=f"='分部财务'!F{SR[b,('gross' if b in ['automotive','energy'] else 'op')]}"+(f'-({R("X","Q326","megapack_rev")}-{R("X","Q326","megapack_cost")})' if b=='energy' else '');pf.cell(row,8,'B列：外部营收；C列：营收占比；D列：'+('毛利（不可直接与其他分部经营利润相加）' if b in ['automotive','energy'] else '经营利润'))
r=pf.max_row+2
pf.cell(r,1,'非控股权益及可赎回非控股权益')
pf.cell(r,6,'='+R('T','SEP26','nci')+'+'+R('T','SEP26','rnci'))
pf.cell(r,8,'不属于归母股东权益，600单列用于资产负债勾稽。')
bsnci=r
# equity bridge
own=sheet('股本与持股','股本清洗 | 已发行、可行权与未达标奖励',['项目','数值','单位','来源 / 处理'],[55,25,20,125]);OR={}
def O(k,label,formula,why,fmt=None):
 own.append([label,'='+formula,'百万股' if not fmt else '比例',why]);OR[k]=own.max_row
 if fmt:own.cell(own.max_row,2).number_format=fmt
O('tshares','Tesla已发行股份',R('OWN','SEP26','t_shares'),'T4 Ownership of Securities；2026-09-30；包含已发行未达标奖励')
O('sshares','SpaceX已发行A+B',R('OWN','SEP26','s_a')+'+'+R('OWN','SEP26','s_b'),'S3期末A8240+B5339；披露百万股四舍五入')
O('cross','Tesla持有SpaceX股份（估计）',R('OWN','MAR26','t_holds_s_pre')+'*'+R('OWN','MAR26','split'),'S1 F61：3.8百万股为拆股前，乘5得约19；假设无后续买卖')
O('sext','SpaceX外部可获换股股份','B5-B6','已发行总股数减Tesla自持，避免向自己发行')
O('tmusk','Musk可支配Tesla已发行股份',R('OWN','SEP26','t_musk'),'T4含2018奖励已行权限制股；不含423.744未达标受不可撤销代理约束股份')
O('smusk','Musk可支配SpaceX已发行股份',R('OWN','MAY26','s_musk_a')+'+'+R('OWN','MAY26','s_musk_b')+'-'+R('OWN','MAY26','s_musk_options')+'+'+AR('sdelta'),'S1表内含350可行权期权，扣除以匹配已发行分母；含1302.072已发行绩效限制股，未必已归属')
O('tbase','Musk合并前Tesla股份占比','B8/B4','统一已发行分母；与委托书19.8%口径不同','0.00%')
O('sbase','Musk合并前SpaceX股份占比','B9/B5','5月持股沿用至9月为假设；不是表决权','0.00%')
O('treported','Tesla委托书口径重算','B8/(B4-'+R('OWN','SEP26','t_unearned')+')','对应披露约19.8%；从分母剔除未达标股份','0.00%')
O('sreported','SpaceX招股书受益持股比例重算','('+R('OWN','MAY26','s_musk_a')+'+'+R('OWN','MAY26','s_musk_b')+')/('+R('OWN','MAY26','s_a_preipo')+'+'+R('OWN','MAY26','s_b_preipo')+'+'+R('OWN','MAY26','s_ipo')+'+'+R('OWN','MAY26','s_musk_options')+')','招股书IPO后无超额配售分母；加入其350可行权期权；非当前期末比例','0.00%')
O('svotes','SpaceX披露IPO后表决权',R('OWN','MAY26','s_votes'),'仅引用事实；合并后股权类别与投票倍数未知，不能自动外推','0.0%')
def ORR(k):return f"'股本与持股'!B{OR[k]}"
ex=sheet('换股比例','三种定价方法 | 每1股SpaceX换取Tesla股数',['项目','10月21日收盘','20日加权近似','假设EV/营收','口径说明'],[43,25,25,25,120]);ER={}
rows=[('sprice','SpaceX每股基础价值'),('tprice','Tesla每股基础价值'),('seq','SpaceX基础股权价值'),('teq','Tesla基础股权价值'),('sval','外部SpaceX股份交易对价'),('ratio','换股比例 r'),('newshares','新发行Tesla股份'),('total','合并后股份'),('spct','原SpaceX外部股东占比'),('tpct','原Tesla股东占比'),('musk','Musk合并后股份占比'),('deltat','较合并前Tesla变动（百分点）'),('deltas','较合并前SpaceX变动（百分点）'),('muskopt','若SpaceX350百万期权全部行权'),('muskaward','若Tesla未达标股全部转为可支配（上界）')]
for i,(k,label) in enumerate(rows,4):ER[k]=i;ex.cell(i,1,label)
for col in 'BCD':
 for k,label in rows:
  r=ER[k]
  if k=='sprice':expr="'市场数据'!B4" if col=='B' else "'市场数据'!B25" if col=='C' else f"('SpaceX财务'!G{mr['rev']}*{AR('smult')}+'SpaceX财务'!F{mr['netcash']}+{AR('sp_other')})/{ORR('sshares')}"
  elif k=='tprice':expr="'市场数据'!D4" if col=='B' else "'市场数据'!D25" if col=='C' else f"('Tesla财务'!G{mr['rev']}*{AR('tmult')}+'Tesla财务'!F{mr['netcash']}+{ORR('cross')}*D{ER['sprice']}+{AR('ts_other')})/{ORR('tshares')}"
  elif k=='seq':expr=f'{col}{ER["sprice"]}*{ORR("sshares")}'
  elif k=='teq':expr=f'{col}{ER["tprice"]}*{ORR("tshares")}'
  elif k=='sval':expr=f'{col}{ER["sprice"]}*{ORR("sext")}*(1+{AR("premium")})'
  elif k=='ratio':expr=f'{col}{ER["sprice"]}*(1+{AR("premium")})/{col}{ER["tprice"]}'
  elif k=='newshares':expr=f'{col}{ER["ratio"]}*{ORR("sext")}'
  elif k=='total':expr=f'{col}{ER["newshares"]}+{ORR("tshares")}'
  elif k=='spct':expr=f'{col}{ER["newshares"]}/{col}{ER["total"]}'
  elif k=='tpct':expr=f'{ORR("tshares")}/{col}{ER["total"]}'
  elif k=='musk':expr=f'({ORR("tmusk")}+{ORR("smusk")}*{col}{ER["ratio"]})/{col}{ER["total"]}'
  elif k=='deltat':expr=f'({col}{ER["musk"]}-{ORR("tbase")})*100'
  elif k=='deltas':expr=f'({col}{ER["musk"]}-{ORR("sbase")})*100'
  elif k=='muskopt':expr=f'({ORR("tmusk")}+({ORR("smusk")}+{R("OWN","MAY26","s_musk_options")}*{AR("sexercise")})*{col}{ER["ratio"]})/({col}{ER["total"]}+{R("OWN","MAY26","s_musk_options")}*{AR("sexercise")}*{col}{ER["ratio"]})'
  else:expr=f'({ORR("tmusk")}+{R("OWN","SEP26","t_unearned")}*{AR("tvesting")}+{ORR("smusk")}*{col}{ER["ratio"]})/{col}{ER["total"]}'
  ex[f'{col}{r}']='='+expr
  if k in ['spct','tpct','musk','muskopt','muskaward']:ex[f'{col}{r}'].number_format='0.00%'
  elif k=='ratio':ex[f'{col}{r}'].number_format='0.0000x'
  else:ex[f'{col}{r}'].number_format='#,##0.00;[Red](#,##0.00);–'
 ex.cell(ER['ratio'],5,'法律发行主体假设Tesla；每1股外部SpaceX获r股Tesla。不是反向换股。')
ex.cell(ER['seq'],5,'股权价值=每股价值×期末已发行股份；EV法加现金减债务。')
ex.cell(ER['teq'],5,'倍数法Tesla另加其SpaceX股权的同法价值；市场法市值本已含其资产。')
ex.cell(ER['musk'],5,'已发行/可支配股份口径，非投票权；持股不等于全部已归属经济利益。')
ex.cell(ER['muskopt'],5,'固定r的股数示意，非完整稀释估值；未计行权所得现金和全部员工期权。')
ex.cell(ER['muskaward'],5,'既有股份转为可支配，不重复增加分母；未计334.09美元/股抵扣，不能视为合并自动兑现。')
note(ex,22,'三种方法是交易数学情景，不是价格建议。市场法使用9月股本×10月21日价格，是近似市值而非当日股本核验；SpaceX分母为百万股取整。',5)
sen=sheet('敏感性','溢价 × 定价方法 | 公式二维敏感性',['SpaceX溢价','10月21日收盘','20日加权近似','假设EV/营收','输出指标'],[23,26,26,26,60])
for start,kind,title in [(4,'ratio','每1股SpaceX换取Tesla股数'),(12,'spct','原SpaceX外部股东占比'),(20,'musk','Musk合并后股份占比')]:
 for j,premium in enumerate([-.2,-.1,0,.1,.2,.3]):
  r=start+j;sen.cell(r,1,premium).number_format='0%';sen.cell(r,1).fill=PatternFill('solid',fgColor=yellow);sen.cell(r,5,title)
  for col in 'BCD':
   ratio=f"'换股比例'!{col}{ER['sprice']}*(1+$A{r})/'换股比例'!{col}{ER['tprice']}"
   if kind=='ratio':expr=ratio
   elif kind=='spct':expr=f'({ratio})*{ORR("sext")}/({ORR("tshares")}+({ratio})*{ORR("sext")})'
   else:expr=f'({ORR("tmusk")}+({ratio})*{ORR("smusk")})/({ORR("tshares")}+({ratio})*{ORR("sext")})'
   sen[f'{col}{r}']='='+expr;sen[f'{col}{r}'].number_format='0.0000x' if kind=='ratio' else '0.00%'
 sen.conditional_formatting.add(f'B{start}:D{start+5}',ColorScaleRule(start_type='min',start_color='E3EFFA',end_type='max',end_color='37ADAE'))
note(sen,28,'溢价为相对各方法SpaceX基础价值的加成，Tesla端不加溢价。全部单元格为明确公式，兼容无需Excel数据表重算引擎的阅读器。',5)
check=sheet('校验','模型勾稽 | 所有差额应为0',['校验项目','差额','说明'],[55,25,110])
for co in ['Tesla','SpaceX']:
 for c in 'BCDEFG':
  businesses=['automotive','energy'] if co=='Tesla' else ['space','conn','si']
  expr='+'.join(f"'分部财务'!{c}{SR[b,'rev']}" for b in businesses)+f"-'{co}财务'!{c}{mr['rev']}"
  check.append([co+' '+c+' 分部收入勾稽','='+expr,'本公司正式分部收入总和−公司收入'])
for c in 'BCD':check.append(['换股股东占比合计 '+c,f"='换股比例'!{c}{ER['spct']}+'换股比例'!{c}{ER['tpct']}-1",'原Tesla+外部SpaceX=100%'])
check.append(['Q3备考分部收入一致',f"=SUM('合并备考'!B{min(BR.values())}:B{max(BR.values())})-'合并备考'!F{PR['q','rev']}",'抵销后业务结构与备考营收一致'])
check.append(['Q3 FCF重分类无虚增',f"='合并备考'!F{PR['q','fcf']}-'合并备考'!D{PR['q','fcf']}",'同额抵销经营收款和内部资本采购'])
check.append(['9M26 FCF重分类无虚增',f"='合并备考'!F{PR['y','fcf']}-'合并备考'!D{PR['y','fcf']}",'同上'])
check.append(['备考资产负债勾稽',f"='合并备考'!F{BS['assets']}-'合并备考'!F{BS['liab']}-'合并备考'!F{BS['equity']}-'合并备考'!F{bsnci}",'总资产=总负债+归母权益+非控股权益及可赎回非控股权益'])
# global styles
for s in w:
 s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0;s.print_title_rows='1:3'
 s.auto_filter.ref=s.auto_filter.ref if s.title=='原始数据' else None
 for row in s:
  for c in row:
   c.font=Font(name='Arial',size=10,color=navy)
   c.alignment=Alignment(vertical='center',wrap_text=True)
   if c.row<=3:
    c.fill=PatternFill('solid',fgColor=navy if c.row!=2 else 'E6EFF4');c.font=Font(name='Arial',size=15 if c.row==1 else 10,bold=c.row!=2,color='FFFFFF' if c.row!=2 else muted)
   elif c.data_type=='f':
    c.fill=PatternFill('solid',fgColor=green)
    c.comment=Comment('计算公式：'+c.value+'\n引用单元格逐级追溯至原始数据（含文件全名/章节）或黄色假设。','模型路径')
    if c.number_format=='General':c.number_format='#,##0.0;[Red](#,##0.0);–'
   elif isinstance(c.value,(int,float)):
    if c.fill.fgColor.rgb not in ['00'+yellow,yellow]:c.fill=PatternFill('solid',fgColor=blue)
    if c.number_format=='General':c.number_format='#,##0.0;[Red](#,##0.0);–'
   c.border=Border(bottom=Side(style='hair',color='DCE7ED'))
  if s.row_dimensions[row[0].row].height is None:s.row_dimensions[row[0].row].height=32 if s.title not in ['原始数据','说明'] else 44
 # first 3 rows merged where suitable
 if s.title not in ['说明']:
  s.merge_cells(start_row=1,start_column=1,end_row=1,end_column=s.max_column);s.merge_cells(start_row=2,start_column=1,end_row=2,end_column=s.max_column)
# restore editable assumption distinction, including source-derived estimates
for rr in AL.values():ass.cell(rr,2).fill=PatternFill('solid',fgColor=yellow)
w.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
w.save('交付/合并模型.xlsx')
Path('交付/工作底稿/model_map.json').write_text(json.dumps(dict(mr=mr,PR={a+'_'+b:r for (a,b),r in PR.items()},BS=BS,BR=BR,ER=ER,OR=OR,AL=AL),ensure_ascii=False,indent=2))
print('Workbook created:',sum(s.max_row*s.max_column for s in w),'cells')
