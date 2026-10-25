import sys,json,zipfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parent/'python_lib'))
from openpyxl import load_workbook
from docx import Document
from pptx import Presentation
import pymupdf
out={};f=load_workbook('交付/合并模型.xlsx',data_only=False);v=load_workbook('交付/合并模型.xlsx',data_only=True)
errors=[];empty=[];formulas=0
for s in f:
 for row in s:
  for c in row:
   if c.data_type=='f':
    formulas+=1;d=v[s.title][c.coordinate]
    if d.data_type=='e':errors.append([s.title,c.coordinate,d.value])
    if d.value is None:empty.append([s.title,c.coordinate])
checks={str(v['校验'].cell(i,1).value):v['校验'].cell(i,2).value for i in range(4,v['校验'].max_row+1)}
out['excel']={'worksheets':f.sheetnames,'formula_count':formulas,'formula_errors':errors,'missing_caches':empty,'checks':checks}
out['docx']={'paragraphs':len(Document('交付/投委会备忘录.docx').paragraphs),'zip_valid':zipfile.ZipFile('交付/投委会备忘录.docx').testzip() is None}
out['pptx']={'slides':len(Presentation('交付/投委会汇报.pptx').slides),'zip_valid':zipfile.ZipFile('交付/投委会汇报.pptx').testzip() is None}
for path in Path('交付/校验').glob('*.pdf'):
 pdf=pymupdf.open(path);out[path.name]={'pages':len(pdf),'chars_per_page':[len(p.get_text()) for p in pdf]}
 for i,p in enumerate(pdf):p.get_pixmap(matrix=pymupdf.Matrix(1,1)).save(str(path.parent/(path.stem+f'_{i+1}.png')))
Path('交付/校验/文件检查.json').write_text(json.dumps(out,ensure_ascii=False,indent=2));print(json.dumps(out,ensure_ascii=False,indent=2))
assert not errors and not empty
assert all(isinstance(x,(int,float)) and abs(x)<1e-7 for x in checks.values())
