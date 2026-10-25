from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
R=Path(__file__).resolve().parents[1];M=json.loads((R/'工作底稿/映射.json').read_text())['model_values'];data=json.loads((R/'web/data.json').read_text())
p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
ink='183749';teal='168B84';muted='718892';paper='F4F7F5';blue='3976D6'
def txt(s,x,y,w,h,text,size=20,color=ink,bold=False):
 box=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));tf=box.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=0
 for i,line in enumerate(text.split('\n')):
  q=tf.paragraphs[0] if i==0 else tf.add_paragraph();q.text=line;q.font.name='等线';q.font.size=Pt(size);q.font.color.rgb=RGBColor.from_string(color);q.font.bold=bold;q.space_after=Pt(10)
 return box
def rect(s,x,y,w,h,c):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(c);sh.line.fill.background();return sh
def slide(num,title,sub):
 s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(paper)
 txt(s,.65,.3,8,.3,'FIELD NOTES  /  ENERGY × COMPUTE',10,teal,True);txt(s,.65,.85,12,1,title,31,ink,True);txt(s,.68,1.75,12,.6,sub,13,muted)
 rect(s,.65,7.0,12,.012,'D4E1E2');txt(s,.65,7.12,10,.22,'行业研究与展示 · 非投资建议 · 2026-10-22',9,muted);txt(s,12.2,7.07,.45,.3,f'{num:02}',12,teal)
 return s
def source(s,t):txt(s,.68,6.45,12,.48,t,9,muted)
s=slide(1,'AI 的电从哪来','从数据中心的资本开支，到电网、核电与聚变的商业路径')
txt(s,.75,2.65,6,1.6,'先有可交付的电，\n再有可运行的算力。',34,ink,True)
for i,(a,b) in enumerate([('485 → 950','TWh / 2025估计至2030预测'),('858.24','亿美元 / Meta 2026前三季Capex'),('0','聚变在基准供电模型中的贡献假设')]):
 x=7.3;y=2.4+i*1.17;txt(s,x,y,4,.6,a,34,teal);txt(s,x,y+.58,5,.35,b,12,muted)
source(s,'需求：IEA，2026-04-16，D03/D04；Meta：用户原件，2026-10-21，K01；聚变：A08，研究假设。')
s=slide(2,'需求翻倍，是机房总用电的预测','不是所有数据中心电量都属于 AI；2030 低、高需求均为研究敏感性。')
s.shapes.add_picture(str(R/'工作底稿/需求图.png'), Inches(.65), Inches(2.4), width=Inches(12))
source(s,'Excel：D01/D03、K14–K16。2024/2025为IEA估计；2030基准采用2026-04-16版本。')
s=slide(3,'Meta：支出进入交付密集期','公司指引与已完成投入分开；原始财务数据以所提供季报为准。')
for i,(v,l) in enumerate([('858.24 亿美元','前三季度Capex / 含租赁本金'),(f'{M["K03"]:.1%}','资本开支同比增长'),('894.99 亿美元','2026年9月末在建工程')]):
 x=.75+i*4.2;rect(s,x,2.55,3.9,1.7,'E7F0EE');txt(s,x+.22,2.8,3.5,.65,v,29,teal);txt(s,x+.22,3.55,3.5,.45,l,13,muted)
txt(s,.8,4.65,11.8,1.5,'全年指引 1,350–1,450 亿美元 → 隐含 Q4 支出 491.76–591.76 亿美元\n在建工程和资本开支不是已通电容量；不能直接换算成 AI 用电量。',20)
source(s,'F1：Meta 2026Q3，签署2026-10-21；现金流量表、Note 6、MD&A；M01/M02/M06/M07/M09，K01/K03/K04/K05。')
s=slide(4,'现实供给：电网连接多类电源','长期协议增强收入与供电确定性，但不自动消除交付瓶颈。')
rows=[('电网与燃气','接入、变压器、燃机和气源决定通电速度'),('Crane → Microsoft','约835 MW / 20年；重启、许可与接入条件仍需落实'),('Clinton → Meta','1,121 MW / 20年；主要保留存量，计划增容30 MW'),('Talen → AWS','最高1,920 MW，期限至2042年；逐步达到满合约量'),('可再生能源 + 储能','低发电成本需要叠加时段匹配、输送与保供成本')]
for i,(a,b) in enumerate(rows):
 y=2.45+i*.7;rect(s,.75,y,11.8,.58,'E8F0EF' if i%2==0 else 'F4F7F5');txt(s,.9,y+.08,3.1,.45,a,16,teal,True);txt(s,4,y+.08,8.2,.45,b,15)
source(s,'F2，2026-02-24；F3，2026-08-06；S7，2025-06-03；S8，2025-06-11。Meta买清洁属性，Clinton电力仍流入当地电网。')
s=slide(5,'聚变：投资与购电协议，尚需净电力证明','金额为所列轮次或时点；公司目标不是既成商业供电。')
rows=[('CFS','B2轮 8.63 亿美元','Google 200 MW；ARC目标2030年代初'),('Helion','G轮最终 5 亿美元','Microsoft：2028上线目标，爬坡后≥50 MW'),('TAE','新轮 >1.5 亿美元','Google投资不等于PPA；拟议并购仍在推进'),('Type One','B轮 2 亿美元','Infinity Two商业运行目标2034年')]
for i,(a,b,c) in enumerate(rows):
 y=2.5+i*.8;rect(s,.75,y,11.8,.66,'E8F0EF');txt(s,.95,y+.12,2,.4,a,17,teal,True);txt(s,3,y+.12,3.2,.4,b,17);txt(s,6.35,y+.12,5.9,.45,c,15)
source(s,'S9 2025-08-28；S10 2025-06-30；S11 2026-06-04/9月更新；S12 2023-05-10；S13 2025-06-02；S14 2026-09-30；S15 2026-10-06；S16 2026-05-06。')
s=slide(6,'交付电价，不能用发电 LCOE 代替','独立假设算例：1 GW IT功率 × 85%负载率 × 1.20 PUE × 8,760小时。')
txt(s,.8,2.7,5,1,'8.9352 TWh',43,teal);txt(s,.82,3.65,5,.5,'示例年用电 / 不是公司预测',16,muted)
for i,(a,b) in enumerate([('60','5.36'),('90','8.04'),('120','10.72')]):
 x=6.15+i*2.15;rect(s,x,2.55,1.95,2.8,'E4EEEE');txt(s,x+.2,2.85,1.55,.55,a,30,teal);txt(s,x+.2,3.46,1.55,.4,'美元/MWh',12,muted);txt(s,x+.2,4.18,1.55,.55,b,27,ink);txt(s,x+.2,4.78,1.55,.35,'亿美元 / 年',11,muted)
txt(s,.8,5.65,11.7,.6,'核电与光伏的年电量等效，不代表相同的逐时保供能力。',20)
source(s,'A09–A18：全部为研究假设，2026-10-22；K20–K27公式计算。纯电费不含芯片、土地、水、人员等成本。')
s=slide(7,'把“AI电力”落到可验证的指标','近端看通电与交付，远端看净电力与可用率。')
for i,(a,b) in enumerate([('科技公司','Capex → 在建工程转固 → 已通电GW → 实际TWh'),('供电商','合同 → 许可与接入 → 实际交付 → 实现收入与成本'),('聚变公司','科学增益 → 全厂净电 → 可维护运行 → 商业交付')]):
 y=2.6+i*1.1;txt(s,.8,y,2.5,.6,a,23,teal,True);txt(s,3.3,y+.06,9,.6,b,21)
txt(s,.8,6,11.6,.45,'配套：研究报告.docx / 数据表.xlsx / 本地 3D 网页 web/index.html',14,muted)
source(s,'上述为研究判断框架。未公开PPA价格、聚变可靠交付与完整项目成本，不用假设冒充事实。')
p.save(R/'演示摘要.pptx');print('演示摘要.pptx')
