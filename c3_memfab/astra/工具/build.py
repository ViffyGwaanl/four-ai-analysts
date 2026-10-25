from pathlib import Path
import sys,json,math
P=Path(__file__).resolve().parents[1];sys.path.insert(0,str(P/'工具/依赖'))
import xlsxwriter
from xlsxwriter.utility import xl_col_to_name as C
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
E=json.loads((P/'数据/完整导出.json').read_text()); H=E['history']; init=json.loads((P/'数据/Q01_决策前.json').read_text())['balance_sheet_now']
# Frozen browser dataset shares the authoritative export with the workbook.
(P/'web/data.js').write_text('window.FACTORY_DATA = '+json.dumps({'history':H,'initial':init,'final':E['final'],'cumulative_net_income':E['cumulative_net_income']},ensure_ascii=False)+';\n')
log=['# 16 季度决策日志','金额除特别说明外均为十亿美元；产量为千片。事实来自本地模拟接口。每季决策 note 在提交前写入并留存，结果为提交后观察；没有用未来结果改写原判断。','全部原始依据：数据/规则.txt、数据/Qxx_决策前.json、数据/Qxx_决策.json、数据/Qxx_结果.json；最终核对数据/完整导出.json 的 history。']
for q in H:
 n=q['quarter'];s=json.loads((P/f'数据/Q{n:02}_决策前.json').read_text());d=q['decision'];i=q['income_statement'];b=q['balance_sheet'];o=q['operations']
 log += [f'\n## Q{n:02}｜决策前观察',s['signals']['news'],s['signals']['consensus'],f"期初现金 {s['balance_sheet_now']['cash']:.4f}，普通债务 {s['balance_sheet_now']['debt']:.4f}，库存 {s['balance_sheet_now']['inventory_kwafers']:.1f} 千片。",'\n### 判断与理由（提交原文）',d['note'],'\n### 已提交决策',f"开工率 {d['utilization']:.2%}；销售比例 {d['sell_fraction']:.0%}；资本开支 {d['capex']:.2f}；额外 HBM 研发 {d['hbm_rnd']:.2f}；HBM 占比 {d['hbm_share']:.4%}；借款/还款 {d['borrow']:+.2f}。",'\n### 事后结果',f"普通价格指数 {q['market']['price_index']:.2f}；营收 {i['revenue']:.4f}；净利润 {i['net_income']:.4f}；经营现金流 {q['cash_flow']['operating']:.4f}；期末现金 {b['cash']:.4f}；权益 {b['equity']:.4f}。HBM 生产/销售 {o['hbm_kwafers']:.1f}/{o['hbm_sold_kwafers']:.1f} 千片，普通期末库存 {b['inventory_kwafers']:.1f} 千片。",f"出处：数据/Q{n:02}_决策前.json → signals、balance_sheet_now；数据/完整导出.json → history[{n-1}]。"]
(P/'决策日志.md').write_text('\n\n'.join(log))
# Workbook
W=xlsxwriter.Workbook(str(P/'三张表.xlsx'));W.set_properties({'title':'晶序内存｜十六季度财务账簿','subject':'虚构经营模拟；金额十亿美元','author':'经营分析'})
navy='#133B50';teal='#007C83';blue='#E8F1FF';mint='#EAF7EF';gold='#FFF0CC'
fmt={
 'title':W.add_format({'font_size':22,'bold':True,'font_color':navy}),
 'head':W.add_format({'bold':True,'bg_color':navy,'font_color':'white','border':0,'align':'center'}),
 'label':W.add_format({'font_color':navy}),
 'note':W.add_format({'font_color':'#5B6B75','font_size':10,'text_wrap':True}),
 'input':W.add_format({'bg_color':blue,'font_color':'#185D9C','num_format':'0.0000;[Red](0.0000);–'}),
 'formula':W.add_format({'bg_color':mint,'font_color':navy,'num_format':'0.0000;[Red](0.0000);–'}),
 'pct':W.add_format({'bg_color':mint,'font_color':navy,'num_format':'0.0%;[Red](0.0%);–'}),
 'assume':W.add_format({'bg_color':gold,'font_color':'#9B5C00','num_format':'0.00'}),
 'asspct':W.add_format({'bg_color':gold,'font_color':'#9B5C00','num_format':'0.0%'}),
}
S={name:W.add_worksheet(name) for name in ['导览与假设','经营分析','利润表','资产负债表','现金流量表','原始数据']}
for name,s in S.items():
 s.hide_gridlines(2);s.set_tab_color(teal);s.set_column('A:A',29);s.set_column('B:Q',13);s.set_column('R:R',58);s.freeze_panes(4,1);s.set_landscape();s.fit_to_pages(1,0);s.set_paper(9)
 s.merge_range('A1:Q2','晶序内存  /  '+name,fmt['title']);s.set_row(2,30)
 s.merge_range('A3:Q3','模拟经营 · 金额：十亿美元 USD bn；数量：千片 · 蓝色=源数据；浅绿=公式；浅黄=假设',fmt['note'])
 s.write(3,0,'项目 / 季度',fmt['head'])
 for j in range(16):s.write(3,j+1,f'Q{j+1:02}',fmt['head'])
 s.write(3,17,'来源 / 口径',fmt['head'])
raw=S['原始数据']; rawmap={};r=4
for group in ['decision','market','operations','income_statement','cash_flow','balance_sheet']:
 for key,v in H[0][group].items():
  if isinstance(v,str):continue
  rawmap[(group,key)]=r+1;raw.write(r,0,group+'.'+key,fmt['label'])
  for j,h in enumerate(H):raw.write(r,j+1,float(h[group][key]),fmt['input'])
  raw.write(r,17,f'数据/完整导出.json → history[季度-1].{group}.{key}',fmt['note']);r+=1
raw.set_column('A:A',43)
a=S['导览与假设'];a.merge_range('A5:H5','仅黄色单元格为分析假设；不改写已发生经营结果。',fmt['head'])
ass=[('季度天数（假设）',90,'用于存货周转天数，统一90天。'),('全年季度数（假设）',4,'仅用于年化ROE；不是全年实际收益。'),('收入压力比例（假设）',.2,'收入下降20%，成本/费用/产销/利息不变；无库存、营运资金响应。'),('税率（规则）',.15,'规则.txt：仅盈利征税。'),('最低现金缓冲（假设）',3,'分析目标，不是引擎0.5的紧急授信触发线。'),('期初现金（事实）',init['cash'],'Q01_决策前.json → balance_sheet_now.cash'),('期初库存账面价值（事实）',init['inventory'],'Q01_决策前.json → balance_sheet_now.inventory'),('期初股东权益（事实）',init['equity'],'Q01_决策前.json → balance_sheet_now.equity'),('期初普通借款（事实）',init['debt'],'Q01_决策前.json → balance_sheet_now.debt'),('期初固定资产净额（事实）',init['ppe_net'],'Q01_决策前.json → balance_sheet_now.ppe_net')]
for n,(lab,val,note) in enumerate(ass,6):
 a.write(n-1,0,lab,fmt['label']);a.write(n-1,1,val,fmt['asspct' if n in [8,9] else 'assume'] if n<=10 else fmt['input']);a.merge_range(n-1,2,n-1,12,note,fmt['note'])
a.data_validation('B6:B7',{'validate':'decimal','criteria':'>','value':0});a.data_validation('B8:B9',{'validate':'decimal','criteria':'between','minimum':0,'maximum':1})
notes=['读取路径：原始数据 → 三张表 → 经营分析。所有逐季事实仅来自完整导出，Q0来自首季状态。','毛利率=毛利/营业收入；存货周转天数=平均普通库存账面额/总销售成本×假设天数。引擎未分拆普通成本，故DIO为混合成本近似口径。','ROE=本季净利润/期初期末平均权益；年化ROE=单季ROE×假设全年季度数，不代表已实现年回报。','存货减值已包含在cogs中，利润表单列作备忘，不二次扣减；经营现金流采用净库存变化重构，不重复加回减值。','净负债=普通借款+紧急借款−现金；自由现金流=经营现金流−资本开支。负净负债代表净现金。','压力情景只测静态收入下降，盈利时按15%税率计税，亏损无所得税；不模拟价格/销量/产能的共同变化。','会计对账有0.0001量级舍入差异；审计容差0.0003十亿美元。累计净利润以export.cumulative_net_income为准。','本项目为虚构经营研究与展示，不构成投资建议。季度没有对应真实日历日期；2026-10-22为本次接口读取日。','没有单列应收、应付、股本、留存收益的接口字段；资产负债表忠实保留引擎简化口径，不虚构科目。']
for j,n in enumerate(notes,18):a.merge_range(j-1,0,j-1,12,n,fmt['note']);a.set_row(j-1,33)

def rawref(group,key,col):return f"'原始数据'!{col}{rawmap[(group,key)]}"
def srcrows(sheet,rows):
 for row,lab,g,k in rows:
  sh=S[sheet];sh.write(row-1,0,lab,fmt['label']);sh.write(row-1,17,f'完整导出 history[].{g}.{k}',fmt['note'])
  for j,h in enumerate(H):
   c=C(j+1);sh.write_formula(row-1,j+1,'='+rawref(g,k,c),fmt['formula'],float(h[g][k]))
srcrows('利润表',[(5,'营业收入','income_statement','revenue'),(6,'其中：普通内存收入','income_statement','revenue_std'),(7,'其中：HBM收入','income_statement','revenue_hbm'),(8,'销售成本（含减值）','income_statement','cogs'),(9,'毛利','income_statement','gross_profit'),(10,'研发费用（含额外HBM）','income_statement','rnd'),(11,'销售与管理费用','income_statement','sga'),(12,'营业利润','income_statement','operating_income'),(13,'净利息费用（负为收益）','income_statement','net_interest'),(14,'所得税','income_statement','tax'),(15,'净利润','income_statement','net_income'),(17,'备忘：存货跌价','income_statement','inventory_writedown')])
srcrows('资产负债表',[(5,'现金','balance_sheet','cash'),(6,'普通内存存货（净额）','balance_sheet','inventory'),(7,'固定资产净额','balance_sheet','ppe_net'),(8,'资产总额','balance_sheet','total_assets'),(10,'普通借款','balance_sheet','debt'),(11,'紧急授信借款','balance_sheet','emergency_debt'),(13,'股东权益','balance_sheet','equity'),(17,'备忘：库存数量（千片）','balance_sheet','inventory_kwafers'),(18,'备忘：产能（千片/季）','balance_sheet','capacity_kwafers')])
srcrows('现金流量表',[(6,'净利润','income_statement','net_income'),(7,'折旧（非现金）','cash_flow','depreciation'),(10,'经营活动现金流（接口）','cash_flow','operating'),(12,'投资活动现金流','cash_flow','investing'),(13,'融资活动现金流','cash_flow','financing'),(16,'期末现金（接口）','balance_sheet','cash'),(19,'备忘：紧急借款流入','cash_flow','emergency_borrowing')])
for j,h in enumerate(H):
 c=C(j+1);prev=C(j);b=h['balance_sheet'];i=h['income_statement'];cf=h['cash_flow'];pb=H[j-1]['balance_sheet'] if j else init
 def put(sheet,row,lab,f,v,pct=False):
  sh=S[sheet];sh.write(row-1,0,lab,fmt['label']);sh.write_formula(row-1,j+1,'='+f,fmt['pct' if pct else 'formula'],v)
 put('利润表',16,'毛利率',f'IFERROR({c}9/{c}5,0)',i['gross_profit']/i['revenue'],True)
 put('利润表',19,'收入分项差额',f'{c}5-{c}6-{c}7',i['revenue']-i['revenue_std']-i['revenue_hbm'])
 put('利润表',20,'利润重算差额',f'{c}5-{c}8-{c}10-{c}11-{c}13-{c}14-{c}15',i['revenue']-i['cogs']-i['rnd']-i['sga']-i['net_interest']-i['tax']-i['net_income'])
 put('资产负债表',12,'负债合计',f'{c}10+{c}11',b['debt']+b['emergency_debt'])
 put('资产负债表',14,'负债与权益合计',f'{c}12+{c}13',b['debt']+b['emergency_debt']+b['equity'])
 put('资产负债表',15,'资产负债平衡差额',f'{c}8-{c}14',b['total_assets']-b['debt']-b['emergency_debt']-b['equity'])
 begcash=f"'{ '资产负债表'}'!{prev}5" if j else "'导览与假设'!$B$11"
 beginv=f"'资产负债表'!{prev}6" if j else "'导览与假设'!$B$12"
 begeq=f"'资产负债表'!{prev}13" if j else "'导览与假设'!$B$13"
 put('现金流量表',5,'期初现金',begcash,pb['cash'])
 put('现金流量表',8,'存货减少 /（增加）',f"{beginv}-'资产负债表'!{c}6",pb['inventory']-b['inventory'])
 put('现金流量表',9,'经营活动现金流（重构）',f'SUM({c}6:{c}8)',i['net_income']+cf['depreciation']+pb['inventory']-b['inventory'])
 put('现金流量表',11,'经营现金流对账差额',f'{c}9-{c}10',i['net_income']+cf['depreciation']+pb['inventory']-b['inventory']-cf['operating'])
 put('现金流量表',14,'现金净增加额',f'{c}10+{c}12+{c}13',sum(cf[k] for k in ['operating','investing','financing']))
 put('现金流量表',15,'期末现金（重构）',f'{c}5+{c}14',pb['cash']+sum(cf[k] for k in ['operating','investing','financing']))
 put('现金流量表',17,'现金对账差额',f'{c}15-{c}16',pb['cash']+sum(cf[k] for k in ['operating','investing','financing'])-b['cash'])
 ar=[(5,'营业收入',f"'利润表'!{c}5",i['revenue'],False),(6,'净利润',f"'利润表'!{c}15",i['net_income'],False),(7,'毛利率',f"IFERROR('利润表'!{c}9/'利润表'!{c}5,0)",i['gross_profit']/i['revenue'],True),(8,'HBM收入占比',f"IFERROR('利润表'!{c}7/'利润表'!{c}5,0)",i['revenue_hbm']/i['revenue'],True),(9,'存货周转天数（近似）',f"IFERROR(({beginv}+'资产负债表'!{c}6)/2/'利润表'!{c}8*'导览与假设'!$B$6,0)",(pb['inventory']+b['inventory'])/2/i['cogs']*90,False),(10,'经营现金流',f"'现金流量表'!{c}10",cf['operating'],False),(11,'自由现金流',f"{c}10+'现金流量表'!{c}12",cf['operating']+cf['investing'],False),(12,'期末现金',f"'资产负债表'!{c}5",b['cash'],False),(13,'净负债',f"'资产负债表'!{c}12-'资产负债表'!{c}5",b['debt']+b['emergency_debt']-b['cash'],False),(14,'股东权益',f"'资产负债表'!{c}13",b['equity'],False),(15,'单季ROE',f"IFERROR({c}6/(({begeq}+{c}14)/2),0)",i['net_income']/((pb['equity']+b['equity'])/2),True),(16,'简单年化ROE（假设）',f"{c}15*'导览与假设'!$B$7",i['net_income']/((pb['equity']+b['equity'])/2)*4,True),(17,'开工率',rawref('decision','utilization',c),h['decision']['utilization'],True),(18,'普通库存数量（千片）',rawref('balance_sheet','inventory_kwafers',c),b['inventory_kwafers'],False),(19,'HBM生产（千片）',rawref('operations','hbm_kwafers',c),h['operations']['hbm_kwafers'],False),(20,'HBM销售（千片）',rawref('operations','hbm_sold_kwafers',c),h['operations']['hbm_sold_kwafers'],False),(21,'HBM报废（千片）',f'{c}19-{c}20',h['operations']['hbm_kwafers']-h['operations']['hbm_sold_kwafers'],False),(22,'累计净利润',f'SUM($B6:{c}6)',sum(x['income_statement']['net_income'] for x in H[:j+1]),False),(24,'压力情景：收入（假设）',f"{c}5*(1-'导览与假设'!$B$8)",i['revenue']*.8,False)]
 for row,lab,f,v,pct in ar:put('经营分析',row,lab,f,v,pct)
 pretax=i['revenue']*.8-i['cogs']-i['rnd']-i['sga']-i['net_interest'];stress=pretax-max(pretax,0)*.15
 put('经营分析',25,'压力情景：税前利润',f"{c}24-'利润表'!{c}8-'利润表'!{c}10-'利润表'!{c}11-'利润表'!{c}13",pretax)
 put('经营分析',26,'压力情景：净利润',f"{c}25-MAX({c}25,0)*'导览与假设'!$B$9",stress)
 put('经营分析',27,'压力情景：经营现金流',f'{c}10+{c}26-{c}6',cf['operating']+stress-i['net_income'])
 put('经营分析',28,'实际现金高于缓冲',f"{c}12-'导览与假设'!$B$10",b['cash']-3)
for sh,rg in [('利润表','B19:Q20'),('资产负债表','B15:Q15'),('现金流量表','B11:Q11'),('现金流量表','B17:Q17')]:
 S[sh].conditional_format(rg,{'type':'cell','criteria':'between','minimum':-.0003,'maximum':.0003,'format':W.add_format({'font_color':'#007F50'})})
for title,rows,anchor in [('利润走出周期',[(5,'收入'),(6,'净利润')],'A32'),('现金与权益',[(12,'现金'),(14,'权益')],'I32'),('产品结构与开工率',[(8,'HBM收入占比'),(17,'开工率')],'A49')]:
 chart=W.add_chart({'type':'line'})
 for k,(r,n) in enumerate(rows):chart.add_series({'name':n,'categories':['经营分析',3,1,3,16],'values':['经营分析',r-1,1,r-1,16],'line':{'color':['#087F8C','#E8A340'][k],'width':2.5}})
 chart.set_title({'name':title});chart.set_legend({'position':'bottom'});chart.set_size({'width':720,'height':310});chart.set_chartarea({'border':{'none':True}});chart.set_y_axis({'num_format':'0%' if rows[0][0]==8 else '0','major_gridlines':{'visible':True,'line':{'color':'#E6ECEF'}}});S['经营分析'].insert_chart(anchor,chart)
S['经营分析'].merge_range('I50:Q54','读图提醒：收入和利润后期的大幅增长来自模拟市场价格急升，不能外推为现实预测。压力情景用于检视收入冲击，非对历史数据的替代。',fmt['note'])
W.close()
# Board report: fixed six pages, clear evidence paths and quantified reflections.
doc=Document();sec=doc.sections[0];sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.1);sec.right_margin=Cm(2.1)
normal=doc.styles['Normal'];normal.font.name='Arial';normal._element.rPr.rFonts.set(qn('w:eastAsia'),'PingFang SC');normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(7)
for name in ['Title','Heading 1','Heading 2']:
 doc.styles[name].font.name='Arial';doc.styles[name]._element.rPr.rFonts.set(qn('w:eastAsia'),'PingFang SC');doc.styles[name].font.color.rgb=RGBColor.from_string('133B50')
header=sec.header.paragraphs[0];header.text='晶序内存  /  董事会经营复盘';header.style='Caption'
f=sec.footer.paragraphs[0];f.text='虚构经营模拟 · 研究与展示用途  |  '
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
def p(t):doc.add_paragraph(t)
def title(t):doc.add_heading(t,0)
def h(t):doc.add_heading(t,1)
def sub(t):doc.add_heading(t,2)
def table(headers,rows):
 t=doc.add_table(rows=1,cols=len(headers));t.style='Light Shading Accent 1'
 for c,v in zip(t.rows[0].cells,headers):c.text=str(v)
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 return t
def source(t):doc.add_paragraph('依据｜'+t,style='Caption')
def page():doc.add_page_break()
title('穿越下行，兑现 HBM\n十六季度经营复盘');p('致董事会  /  2026年10月22日\n虚构公司展示名称：晶序内存；不对应真实企业。')
h('期末权益增至 171.2457 十亿美元')
p('公司完成全部16个季度，未破产，未动用紧急授信。期初权益34.9600，期末171.2457，增幅389.83%；导出累计净利润136.2858。期末现金152.9207、普通借款与紧急借款均为零。权益增量与累计净利润相差0.0001，属于接口四位小数舍入。')
table(['董事会指标','结果（十亿美元，另注除外）'],[['累计收入 / HBM收入','259.0555 / 150.7872'],['累计经营现金流 / 资本开支','175.9208 / 27.0000'],['现金最低点 / 权益最低点','Q3：2.9765 / Q6：20.9828'],['首次单季盈利 / 债务归零','Q7 / Q11'],['期末产能 / 累计HBM销售','325.0千片每季 / 952.0千片']])
p('成果由两部分共同形成：早期完成研发与现金管理使公司存活并有能力接单；后期模拟市场的剧烈涨价放大了回报。末四季贡献122.2572净利润，占全期89.71%。不应把这一市场红利全部归因为管理能力。')
source('数据/完整导出.json → final、cumulative_net_income、history[].income_statement/cash_flow/balance_sheet；期初见数据/Q01_决策前.json。全篇金额统一为十亿美元（1单位=10亿美元）。')
page();h('01  /  周期如何改变了公司')
table(['阶段','实际市场与经营变化','管理动作'],[['Q1—Q3 下行','普通价格0.80→0.50；累计亏损9.4835','降低开工、Q2清空普通库存；Q1一次投3.0额外研发'],['Q4—Q6 筑底','HBM开始量产；Q5经营现金流转正','借款合计3.0支持周转；逐步增产'],['Q7—Q10 修复','Q7首次盈利；Q10 HBM收入占51.08%','满开工；两季各追加1.5扩产；开始还款'],['Q11—Q16 加速','产能300→325；Q16毛利率90.10%','Q11还清债务；逐季增加HBM占比']])
sub('价格是回报幅度的主要外部变量')
p('普通价格指数从Q3低点0.50升至Q16的4.55，达到9.10倍；HBM价格从Q4的1.40升至Q16的9.60，达到6.86倍。Q16收入57.3060，其中HBM收入41.4720，占72.37%。相较之下，期末产能仅比初始增加8.33%。')
sub('净利润与现金流的转正不同步')
p('Q5净利润仍为−1.6757，但经营现金流已为0.3993，主要因2.0750折旧并非当期现金支出。Q6仍亏0.5007，权益触底20.9828；Q7净利润才转正为0.8407。管理中同时看利润和现金，不用单一利润指标决定是否继续经营。')
source('完整导出 → history[0:16].market、operations、income_statement、cash_flow。阶段统计按所属季度求和；比率由三张表.xlsx公式计算。')
page();h('02  /  做对了什么')
sub('以研发换取进入 HBM 的时间')
p('Q1一次投入3.0000额外HBM研发；接口确认为Q4具备资格。Q4—Q16累计HBM收入150.7872，占累计收入58.21%；累计销售952.0千片。这个收入规模证明HBM成为核心业务，但不能直接当作研发净收益：它包含共享产能、成本、价格环境与普通产品机会成本。')
sub('下行期去库存，避免跌价持续侵蚀权益')
p('期初普通库存60.0千片、账面0.9600；Q1剩余30.0千片，Q2清零，此后维持零期末库存。全期计提存货跌价仅0.0095，全部在Q1。代价是部分季度普通需求未被满足；低库存并非没有机会成本。')
sub('用有限借款跨过现金低谷，再及时偿还')
p('Q3期末现金最低2.9765，仍高于规则0.5000紧急授信触发线2.4765。Q4、Q5分别借2.0与1.0，普通债务峰值9.0；Q9—Q11累计偿还9.0，Q11起零负债。全期经营现金流175.9208减资本开支27.0，得到自由现金流148.9208；再减净偿债6.0，现金较期初增加142.9207，微差来自舍入。')
sub('扩产节奏控制了资本承诺')
p('Q7、Q8分别投入3.0资本开支，各含1.5维持、1.5扩产；分别在Q11、Q12上线12.5千片/季。全期资本开支27.0，包含24.0维持与3.0扩产。没有在末期大涨时追加四季建设，避免交付期内无法验证的新项目。')
source('完整导出 → history[].decision、operations、income_statement、cash_flow、balance_sheet；Q02_决策前 → hbm.qualified_from_quarter；规则.txt。')
page();h('03  /  判断失误与未满足的需求')
sub('量产时间理解早了一个季度')
p('Q1原始note写“争取Q3取得资格”，但提交后接口明确qualified_from_quarter=4。Q2立即修正判断。没有把这一误读从日志中删除；以后先把“累计达标后再过2季”的计时口径转成明确的季度里程碑，并以接口反馈为准。')
sub('HBM并非每季都准确匹配订单')
p('全期HBM生产963.7千片，销售952.0千片，合计11.7千片未售作废，占生产1.21%。Q4浪费2.8、Q10浪费3.0千片，分别反映首次导入和增长外推偏高。相反，Q14和Q16各少供5.0千片，说明保守假设也会牺牲高价值订单。')
p('事后静态估算（假设）：若在Q14、Q16各把5千片普通产出改成HBM，且总成本、市场价格和税率不变，新增收入分别为5×0.024×(7.60−2.65)=0.5940、5×0.024×(9.60−4.55)=0.6060，税后合计约1.0200。这是局部机会成本估算，既非真实已得利润，也不是对整局最优策略的证明。单价换算0.024依据导出收入、销量与价格指数反推。')
sub('产能保守有真实代价，不能只宣称扩产克制正确')
p('16季普通市场需求减实际销量合计1600.2千片；Q16普通需求510.0、销量145.0，缺口365.0千片。这里是模拟市场需求与公司供货差额，不等同于已锁定可赚钱的客户订单。Q7以后需求持续增长，较早追加产能可能更好，但在决策当时存在周期风险，且没有第二局或反事实引擎结果，不能声称已验证最佳扩产路径。')
source('完整导出 → history[].market/operations/decision；Q01_决策.json原始note。估算公式与假设如上，和已发生结果分开。')
page();h('04  /  如果重来：改进决策机制')
table(['改进动作（建议，非已发生事实）','具体约束与检验'],[['先核对研发时钟','用接口qualified_from_quarter设置最早可销售季度；达标后停追加研发'],['把HBM预测写成区间','以最近销量与订单信号列低/中/高情景；比较报废成本和少供机会成本'],['提前进行扩产压力测试','对投产时价格和需求做下行情景；同时检验现金最低点与回收期'],['保持现金下限纪律','示例目标3.0；Q3实际2.9765略低于目标，应提前安排少量周转款'],['每季同时看四个指标','现金、平均库存、需求未满足量、HBM报废率；不只看毛利率']])
sub('示例压力测试：收入下降20%')
p('三张表.xlsx 的“导览与假设”B8可调整收入下降比例。假设Q16收入下降20%，总销售成本5.6750、研发0.8000、管理0.2500与净利息−0.8612均不变，税率15%仅对盈利征税，则收入45.8448、税前利润39.9810、净利润33.98385。相较实际净利润43.7258，下降约9.74195。')
p('对应经营现金流约37.03385，系以实际46.7758加净利润变动推导，并假设折旧和营运资金不变。这是静态压力情景，不是未来预测；未模拟订单萎缩、存货减值、债务变化、竞争新增产能或售价与销量之间的联动。')
sub('期末现金并不等于永久盈利能力')
p('期末152.9207现金和零债务提高了抗风险能力，但市场新闻已提示新产能陆续投产、价格涨幅收窄。实际经营若继续，应重新评估下一轮周期的研发、维护与扩产预算；16季成绩没有提供第17季之后的经营事实。')
source('三张表.xlsx → 导览与假设、经营分析24—28行；完整导出 → history[15].income_statement/cash_flow；Q16_决策前 → signals。')
page();h('05  /  财务口径、来源和复核边界')
p('本报告严格使用接口公开数据，没有读取引擎代码或市场未来数据。财务信息原始材料只有游戏通行证，实际财务事实来自授权本地接口；所有响应均保存于“数据/”。没有使用外部财务数字。')
table(['交付物','如何追溯'],[['三张表.xlsx','原始数据的R列标注JSON字段；三张表和经营分析用公式引用'],['决策日志.md','每季保留决策前新闻、现金状态、提交note与结果'],['web/index.html','data.js由同一完整导出生成；展示金额与Excel统一为十亿美元'],['数据/完整导出.json','history共16条；final为期末状态；bankrupt=false']])
p('引擎会计简化：只提供现金、普通库存、固定资产、普通借款、紧急借款和权益，没有独立应收、应付或权益明细。存货跌价包含在销售成本中，不重复扣除；经营现金流重构使用净库存余额变化，不重复加回减值。')
p('存货周转天数以期初期末平均普通库存账面额除总销售成本，乘90天。因为接口没有普通与HBM成本拆分，该指标是近似值；90天为假设。单季ROE使用平均权益，年化倍数4为假设。净负债定义为全部有息债务减现金。')
p('原始财务数保留四位小数，跨表重构可能有0.0001量级差异。公式校验使用0.0003十亿美元容差，不用“修正源数”强行消除舍入差。累计净利润136.2858与权益增量136.2857相差0.0001。')
p('尚未核实：策略是否全局最优、替代扩产路径的全期回报、真实行业适用性及第17季以后表现。模拟没有给出这些事实，本报告不作推断性承诺。未要求制作演示文稿，因此无PPT文件。')
p('研究与展示用途，不提供买卖持有建议、目标价或股价贵贱评价。')
source('数据/规则.txt、数据/完整导出.json、数据/Q01_决策前.json；资料读取日期2026-10-22。工具与网页验收记录见检查/验收报告.md。')
doc.save(P/'经营复盘.docx')
print('Created workbook, report, log, web dataset')
