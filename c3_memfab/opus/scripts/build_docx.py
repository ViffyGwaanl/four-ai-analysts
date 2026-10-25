"""生成 交付/经营复盘.docx（写给董事会的复盘）。数字来自 data/quarterly.json（引擎导出）、scenarios.json / review_numbers.json（复现模型估算）。"""
import json, pathlib
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = pathlib.Path(__file__).resolve().parents[1]
D = json.load(open(ROOT / "data/quarterly.json"))
SC = {s["name"]: s for s in json.load(open(ROOT / "data/scenarios.json"))["scenarios"]}
RV = json.load(open(ROOT / "data/review_numbers.json"))
Q, O = D["quarters"], D["opening"]
FIG = ROOT / "data/figures"
EA, LAT = "微软雅黑", "Arial"
NAVY, INK2 = RGBColor(0x1F, 0x38, 0x64), RGBColor(0x52, 0x51, 0x4E)

from decimal import Decimal, ROUND_HALF_UP
def _hu(x, nd):   # 十进制四舍五入（与 Excel 显示一致，避免二进制浮点把 123.705 显示成 123.70）
    return Decimal(repr(float(x))).quantize(Decimal(1).scaleb(-nd), rounding=ROUND_HALF_UP)
def f2(x): return f"{_hu(x, 2):,.2f}"
def pct(x, d=1): return f"{_hu(float(Decimal(repr(float(x))) * 100), d):.{d}f}%"

# ---------------- 关键数字（全部由数据计算） ----------------
eq0, eq1 = O["equity"], D["final"]["equity"]
cum_ni = D["cumulative_net_income"]
rev = sum(q["is_revenue"] for q in Q); rev_hbm = RV["hbm_rev_total"]; gp = RV["gp_total"]; hbm_gp = RV["hbm_gp_total"]
mult = eq1 / eq0; cagr_q = mult ** (1 / 16) - 1; cagr_a = (1 + cagr_q) ** 4 - 1
min_cash = min(q["bs_cash"] for q in Q); min_cash_q = min(Q, key=lambda q: q["bs_cash"])["quarter"]
max_debt = max(q["bs_debt"] for q in Q); max_debt_q = max(Q, key=lambda q: q["bs_debt"])["quarter"]
exp_capex = RV["expansion_capex"]; cap_add = Q[-1]["capacity_kwafers"] - O["capacity_kwafers"]
peak_inv = max(q["bs_inventory_kwafers"] for q in Q); peak_inv_q = max(Q, key=lambda q: q["bs_inventory_kwafers"])["quarter"]
scrap = sum(q["hbm_kwafers"] - q["hbm_sold_kwafers"] for q in Q)
unmet, unmet_v = RV["unmet_std_q14_16"], RV["unmet_std_value"]
S = lambda k: SC[k]["delta"]
redo = RV["redo"]
lim = RV["debt_limit_at_decision"]
assert abs(eq1 - 212.6423) < 1e-6 and abs(cum_ni - 177.6823) < 1e-6

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Cm(2.2))
sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)

def set_fonts(style, size=None, bold=None, color=None):
    style.font.name = LAT
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:eastAsia"), EA); rf.set(qn("w:ascii"), LAT); rf.set(qn("w:hAnsi"), LAT)
    if size: style.font.size = Pt(size)
    if bold is not None: style.font.bold = bold
    if color is not None: style.font.color.rgb = color

st = doc.styles
set_fonts(st["Normal"], 10.5)
st["Normal"].paragraph_format.space_after = Pt(4)
st["Normal"].paragraph_format.line_spacing = 1.3
for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11)):
    set_fonts(st[name], size, True, NAVY)
st["Heading 1"].paragraph_format.space_before = Pt(14); st["Heading 1"].paragraph_format.space_after = Pt(6)
st["Heading 2"].paragraph_format.space_before = Pt(10); st["Heading 2"].paragraph_format.space_after = Pt(4)
for name in ("List Bullet", "List Number"):
    set_fonts(st[name], 10.5)

def para(text="", bold_prefix=None, size=None, color=None, align=None, italic=False, space_after=None):
    p = doc.add_paragraph()
    if bold_prefix:
        r = p.add_run(bold_prefix); r.bold = True
        if size: r.font.size = Pt(size)
    r = p.add_run(text)
    if size: r.font.size = Pt(size)
    if color is not None: r.font.color.rgb = color
    r.italic = italic
    if align: p.alignment = align
    if space_after is not None: p.paragraph_format.space_after = Pt(space_after)
    return p

def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = p.add_run(bold_prefix); r.bold = True
    p.add_run(text)
    return p

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def table(rows, header=True, widths=None, font=9, align_num=True, zebra=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = ""
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.1
            r = p.add_run(str(v)); r.font.size = Pt(font)
            if header and i == 0:
                r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(c, "1F3864")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                if zebra and i % 2 == 0: shade(c, "F3F6FA")
                if align_num and j > 0: p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if widths:
        t.autofit = False
        tblPr0 = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr0.append(lay)
        for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(int(w / 2.54 * 1440)))
        for j, w in enumerate(widths):
            for i in range(len(rows)):
                t.cell(i, j).width = Cm(w)
    # 边框改成浅色
    tblPr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{e}"); b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4"); b.set(qn("w:color"), "D0D5DD"); borders.append(b)
    tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def figure(name, caption, width=16.4):
    doc.add_picture(str(FIG / name), width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = para(caption, size=8.5, color=INK2, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)

def callout(text, fill="EEF3FA", bold_prefix=None):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0); shade(c, fill); c.text = ""
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    if bold_prefix:
        r = p.add_run(bold_prefix); r.bold = True; r.font.size = Pt(9.5)
    r = p.add_run(text); r.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# 页脚页码
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run("经营复盘 · 第 "); run.font.size = Pt(8); run.font.color.rgb = INK2
fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
rr = OxmlElement("w:r"); tt = OxmlElement("w:t"); tt.text = "1"; rr.append(tt); fld.append(rr); fp._p.append(fld)
run = fp.add_run(" 页"); run.font.size = Pt(8); run.font.color.rgb = INK2

# ================= 封面 =================
doc.add_paragraph(style="Title").add_run("经营复盘：穿越一轮内存超级周期")
para("第 1–16 季 · 写给董事会", size=12, color=NAVY, space_after=2)
para("编制：首席执行官 · 日期：2026-10-22 · 游戏编号 memfab-opus", size=9, color=INK2, space_after=8)
callout("金额单位都是十亿美元（与经营模拟引擎一致，例如 212.64 即 2126.4 亿美元），片数单位是千片/季。"
        "没有标注的数字都来自引擎导出数据（交付/data/raw/export_final.json）；标注「估算」的数字来自我按推断参数重建的复现模型"
        "（16 季逐季复现引擎结果，误差小于 0.001），是假设性的情景计算，不是引擎输出。市场、公司和数字都是虚构的，本报告不构成任何投资建议。",
        bold_prefix="读法说明　")

# ================= 1 结论 =================
doc.add_heading("一、结论先行", level=1)
bullet(f"期末股东权益 {f2(eq1)}，是期初 {f2(eq0)} 的 {mult:.2f} 倍（季度复合增长 {pct(cagr_q)}，折合年化 {pct(cagr_a)}）；16 季累计净利润 {f2(cum_ni)}。"
       f"全程没有断过资金链、没有动用过紧急授信；期末零借款、零库存，现金 {f2(Q[-1]['bs_cash'])}。", "结果：")
bullet(f"第 1 季一次投满 HBM 研发 3.00。HBM 累计收入 {f2(rev_hbm)}，占全部营收 {pct(rev_hbm/rev)}；HBM 毛利 {f2(hbm_gp)}，占全部毛利 {pct(hbm_gp/gp)}。"
       f"估算：如果不做 HBM，期末权益少 {f2(-S('S2 不投HBM'))}。", "最值钱的决定：")
bullet(f"在价格底部和复苏早期（第 5–10 季）投入扩产资金 {f2(exp_capex)}，产能从 300 增加到 {Q[-1]['capacity_kwafers']:.0f} 千片/季（+{pct(cap_add/300,0)}）。"
       f"估算：完全不扩产，期末权益少 {f2(-S('S3 不扩产(借款同实际)'))}。", "第二值钱：")
bullet(f"下行期满开工、把多出的片子存进仓库（库存最高 {peak_inv:.0f} 千片），复苏后在第 14 季前全部卖掉。估算：像同业一样减产到 70%，期末权益少 {f2(-S('S1 下行期减产到70%(Q2-Q5)'))}。", "第三：")
bullet(f"对顶部的判断早了两三个季度：第 11–12 季停止扩产、第 12–13 季没有惜售。第 14–16 季普通内存需求比我们的供给多 {unmet:.1f} 千片，"
       f"按当季价格约 {f2(unmet_v)} 的收入没有接住。按只用当时信息就能想到的改进规则重来，估算期末权益还能多约 {f2(redo['R3 R2+Q11/Q12每季多扩3.6']['delta'])}"
       f"（约为期末权益的 {pct(redo['R3 R2+Q11/Q12每季多扩3.6']['delta']/eq1,0)}）。", "最大的错：")

# ================= 2 发生了什么 =================
doc.add_heading("二、这 16 个季度发生了什么", level=1)
doc.add_heading("2.1 市场：一轮完整的超级周期", level=2)
para(f"普通内存价格指数从 1.00 跌到第 3 季的 0.50，随后一路上涨到第 16 季的 {Q[-1]['price_index']:.2f}；HBM 价格指数从 1.6 跌到 1.3 后涨到 {Q[-1]['hbm_price_index']:.1f}。"
     f"我们的普通内存需求从 240 千片跌到 205 千片，最后升到 {Q[-1]['demand_kwafers']} 千片；HBM 需求从 5 千片涨到 {Q[-1]['hbm_demand_kwafers']} 千片。"
     "每季我只能看到上一季价格、一条新闻和分析师一致预期，未来走势要靠判断。")
figure("fig1_price.png", "图 1　价格指数（1.00 = 每千片 0.024 十亿美元，推断）。来源：引擎导出 market.price_index / hbm_price_index。")
doc.add_heading("2.2 四个阶段", level=2)
ph = RV["phases"]
rows = [["阶段", "价格指数（首→末）", "营业收入", "其中 HBM", "净利润", "毛利率", "资本开支", "期末现金", "期末借款", "期末权益"]]
SHORT = {"下行（Q1-Q4）": "下行 Q1–4", "复苏与布局（Q5-Q8）": "复苏布局 Q5–8", "加速（Q9-Q12）": "加速 Q9–12", "超级周期（Q13-Q16）": "超级周期 Q13–16"}
for k, v in ph.items():
    rows.append([SHORT.get(k, k), f"{v['price_start']:.2f}→{v['price_end']:.2f}", f2(v["rev"]), f2(v["hbm"]), f2(v["ni"]), pct(v["gm"]), f2(v["capex"]), f2(v["cash_end"]), f2(v["debt_end"]), f2(v["eq_end"])])
table(rows, widths=[2.9, 2.0, 1.55, 1.45, 1.5, 1.35, 1.45, 1.45, 1.4, 1.55], font=8)
p1, p2, p3, p4 = list(ph.values())
bullet(f"下行（Q1–Q4）：价格一年内腰斩，四个季度累计亏损 {f2(-p1['ni'])}，存货跌价共 {f2(RV['writedown_total'])}，股东权益从 {f2(eq0)} 降到 {f2(p1['eq_end'])}。"
       f"我们没有减产，库存从 60 千片升到 {p1['inv_end']:.0f} 千片；借款从 6.00 增到 {f2(p1['debt_end'])}。", "")
bullet(f"复苏与布局（Q5–Q8）：价格指数 0.60→0.92，第 6 季净利润转正。HBM 需求从 12 千片涨到 36 千片。这 4 个季度资本开支 {f2(p2['capex'])}，"
       f"其中扩产 11.40；借款用到上限附近，现金在第 {min_cash_q} 季最低到 {f2(min_cash)}。", "")
bullet(f"加速（Q9–Q12）：新产能陆续投产（325→395 千片），HBM 收入 {f2(p3['hbm'])}，第 12 季首次超过普通内存；四季净利润 {f2(p3['ni'])}，借款在第 12 季全部还清。", "")
bullet(f"超级周期（Q13–Q16）：价格指数 1.70→4.55，四季营收 {f2(p4['rev'])}、净利润 {f2(p4['ni'])}，占 16 季累计净利润的 {pct(p4['ni']/cum_ni,0)}。"
       f"从第 14 季起需求大于我们的产能，库存清零。", "")
figure("fig2_revenue.png", "图 2　营收结构与净利润（十亿美元）。来源：引擎导出 income_statement。")
figure("fig3_balance.png", "图 3　股东权益、现金与借款（十亿美元，期末数）。来源：引擎导出 balance_sheet。")

# ================= 3 做对了什么 =================
doc.add_heading("三、哪些决定做对了", level=1)
doc.add_heading("3.1 第 1 季一次投满 HBM 研发", level=2)
para("当时：一致预期本季价格跌 15%～25%，HBM 需求只有 5 千片。我判断 HBM 是给 AI 芯片用的高价产品，研发是「门票」，越早拿到越好，所以在亏损季一次投入 3.00，第 4 季拿到量产资格。")
para(f"结果：第 7 季新闻说「HBM 认证名额紧张，只有少数厂商通过」，我们已经在里面。HBM 价格指数从 1.3 涨到 9.6，16 季 HBM 收入 {f2(rev_hbm)}、毛利 {f2(hbm_gp)}。"
     f"按每千片计算，第 16 季 HBM 售价 0.230，是普通内存（0.109）的 2.1 倍，而单位成本相同。")
para(f"需要说明：如果把研发推迟到第 5 季（第 8 季才能量产），估算期末权益反而多 {f2(S('S2b HBM研发推迟到Q5'))}，因为早期 HBM 量很小，省下的产能做成普通内存后来按高价卖掉了；"
     "但这个估算假设晚进入者拿到同样多的 HBM 订单，而当时的新闻恰恰说认证名额紧张，而且那条路径要动用紧急授信。我仍然认为早投是对的，只是可以分两季投、减轻下行期的现金压力。", color=INK2)
doc.add_heading("3.2 在底部和复苏早期逆周期扩产", level=2)
para(f"当时：第 4–5 季的新闻是「两家同业减产超过 20%」「渠道补库开始」「云厂商测试新一代 AI 服务器」。扩产要 4 个季度才投产，所以我从第 5 季开始每季扩产，到第 10 季共投入 {f2(exp_capex)}，"
     f"新增 {cap_add:.0f} 千片/季产能，分别在第 9–14 季投产。")
para(f"结果：新增产能正好赶上 HBM 放量和普通内存涨价。估算如果完全不扩产，期末权益少 {f2(-S('S3 不扩产(借款同实际)'))}，"
     f"约为扩产投入的 {(-S('S3 不扩产(借款同实际)'))/exp_capex:.1f} 倍。另外，从报表看出扩产投资在投产当季才开始折旧，所以后期投入的产能对 16 季内利润的拖累很小（这是推断）。")
doc.add_heading("3.3 下行期满开工，把产能存进仓库", level=2)
para("当时：第 2 季我用前两季的报表反推出成本结构——每季固定成本约 3.00（折旧 2.00 + 工厂费用 1.00），每千片变动成本只有 0.005；"
     "而价格指数跌到 0.50 时，每千片售价仍有 0.012。也就是说，少产一千片只省 0.005 现金，固定成本却照付。")
para(f"结果：第 2–5 季我满开工，库存在第 {peak_inv_q} 季升到 {peak_inv:.0f} 千片；第 7–13 季需求超过产量后，这些库存按 0.83～1.70 的价格指数陆续卖掉，第 14 季清零。"
     f"估算：第 2–5 季减产到 70%，期末权益少 {f2(-S('S1 下行期减产到70%(Q2-Q5)'))}；减产到 50%，少 {f2(-S('S1b 下行期减产到50%(Q2-Q5)'))}。代价是下行期存货跌价共 {f2(RV['writedown_total'])}（会计上提前确认，没有额外现金流出）。")
figure("fig4_capacity.png", "图 4　产量分配与总需求（千片）。产大于需的部分进库存，第 14–16 季需求大于产能。来源：引擎导出 operations、market。")
figure("fig5_inventory.png", "图 5　期末库存（千片）。来源：引擎导出 balance_sheet.inventory_kwafers。", width=15.5)
doc.add_heading("3.4 流动性纪律与去杠杆", level=2)
para(f"借款上限按股东权益的 60% 计算，下行期权益下降，上限从第 1 季的 {lim[0]:.2f} 降到第 6 季的 {lim[5]:.2f}。我每季按「季末现金不低于约 1.5」来定扩产规模和借款，"
     f"借款最高 {f2(max_debt)}（第 {max_debt_q} 季），现金最低 {f2(min_cash)}（第 {min_cash_q} 季），从未触发 3% 的紧急授信。"
     f"经营现金流转强后，第 10–12 季还清全部借款。第 1–12 季净利息支出合计 {f2(RV['interest_paid_q1_q12'])}；之后现金利息收入超过了它，16 季净利息合计为净收入 {f2(-RV['net_interest_total'])}。")
doc.add_heading("3.5 没有把库存留过终点", level=2)
para("期末库存只按成本入账（每千片约 0.015），远低于第 16 季每千片 0.109 的售价。最后两季我坚持全部出货，期末库存为零，没有让价值躺在仓库里。")

# ================= 4 做错了什么 =================
doc.add_heading("四、哪些决定做错了", level=1)
doc.add_heading("4.1 顶部判断早了两三个季度（最大的错）", level=2)
para("当时：第 9 季新闻说「普通内存涨势放缓」，第 11 季说「同业新产能一年后才出来」，我据此判断第 14–15 季会见顶，于是第 11 季起只做维持性资本开支，第 12–13 季也没有惜售。"
     "实际上价格指数从第 12 季的 1.28 一路涨到第 16 季的 4.55。")
para(f"代价（事后看）：第 14、15、16 季普通内存需求缺口分别为 78.9、182.4、233.3 千片，合计 {unmet:.1f} 千片，按当季价格约 {f2(unmet_v)} 的收入。"
     f"估算：第 11、12 季各多投 3.60（各 +30 千片），期末权益多 {f2(S('S4 顶部前再扩产(Q11/Q12各+36亿)'))}；各多投 7.20，多 {f2(S('S4b 顶部前再扩产(Q11/Q12各+72亿)'))}。"
     f"第 12–13 季只卖当季产出、把库存留到第 14 季以后卖，多 {f2(S('S5 Q12-Q13惜售,库存留到后面'))}。")
para("我的反思：我用了「叙事」信号（新闻里的「放缓」「新产能」）去判断顶部，却没有重视「量」的信号——一致预期连续上调、我们的库存天数从 75 天降到 13 天、需求增速快于产能增速。"
     "而且投产滞后 4 个季度、折旧从投产才开始，意味着第 11–12 季下单的产能在第 15–16 季投产时几乎没有拖累，判断错了代价也有限。", color=INK2)
doc.add_heading("4.2 HBM 的安全余量在普通内存紧缺时偏大", level=2)
para(f"HBM 按订单生产，多产作废。我每季按报童模型给 HBM 留了余量，16 季共作废 {scrap:.1f} 千片。前期普通内存过剩时，这个余量几乎没有成本；"
     f"但第 13–16 季普通内存也缺货，作废的每一千片 HBM 都是一千片能按 0.041～0.109 卖掉的普通内存。估算：按需求精确排产，期末权益多 {f2(S('S6 HBM按需求精确排产'))}。"
     "错误在于我没有及时把「被挤掉的普通内存的价值」调高到市场价，余量应该压到 1%～2%。")
doc.add_heading("4.3 第 1 季减产到 85%", level=2)
para(f"第 1 季还不知道成本结构，我按「客户谨慎」减产 15%。少产的 45 千片后来本可以在缺货季按高价卖掉。估算：满开工期末权益多 {f2(S('S7 Q1满开工'))}。")
doc.add_heading("4.4 没有在借款上限高的时候先锁定流动性", level=2)
para(f"借款上限是顺周期的：第 1 季 {lim[0]:.2f}，第 6 季只剩 {lim[5]:.2f}。第 5–8 季我的扩产规模被流动性卡住（第 6–8 季每季都借到上限）。"
     f"如果第 1 季就把借款一次加到 20.90（在当时 {lim[0]:.2f} 的上限以内，比实际多借 14.90，之后第 3–9 季不再新借；估算多付净利息约 "
     f"{f2(redo['R0 只做：Q1满开工']['delta'] - redo['R1 Q1满开工+Q1预借14.9']['delta'])}），第 5–7 季就有钱每季多扩产 1.50。估算扣掉利息后，这一项能让期末权益多约 "
     f"{f2(redo['R2 R1+Q5-Q7每季多扩1.5']['delta'] - redo['R0 只做：Q1满开工']['delta'])}。"
     "这里有一个没能核实的假设：权益下降后已有借款超过新上限时，引擎不会强制还款（规则没写，我也没有试过）。", color=None)
figure("fig6_decisions.png", "图 6　关键决定的价值（估算）：实际决定相对替代方案对期末股东权益的影响。蓝色是做对的，橙色是事后看错失的。来源：复现模型情景计算（交付/scripts/sim.py）。")
rows = [["决定", "替代方案", "期末权益差（估算）", "说明"]]
for name, alt, key, note in [
    ("第 1 季投 HBM 研发", "完全不做 HBM", "S2 不投HBM", "替代方案需要动用紧急授信"),
    ("第 5–10 季扩产 19.80", "每季只做 1.50 维持", "S3 不扩产(借款同实际)", "产能停在 300 千片/季"),
    ("第 2–5 季满开工", "减产到 70%", "S1 下行期减产到70%(Q2-Q5)", "复苏后少了库存可卖"),
    ("第 2–5 季满开工", "减产到 50%", "S1b 下行期减产到50%(Q2-Q5)", "同上"),
    ("第 1 季开工率 85%", "满开工", "S7 Q1满开工", "事后看"),
    ("第 12–13 季按需出货", "只卖当季产出，库存留后", "S5 Q12-Q13惜售,库存留到后面", "事后看"),
    ("第 11–12 季不扩产", "各多投 3.60", "S4 顶部前再扩产(Q11/Q12各+36亿)", "事后看"),
    ("第 11–12 季不扩产", "各多投 7.20", "S4b 顶部前再扩产(Q11/Q12各+72亿)", "事后看；需多借款"),
    ("HBM 留安全余量", "按需求精确排产", "S6 HBM按需求精确排产", "需要完美预知需求"),
]:
    d = -S(key) if key in ("S2 不投HBM", "S3 不扩产(借款同实际)", "S1 下行期减产到70%(Q2-Q5)", "S1b 下行期减产到50%(Q2-Q5)") else S(key)
    sign = "实际多赚 " if key in ("S2 不投HBM", "S3 不扩产(借款同实际)", "S1 下行期减产到70%(Q2-Q5)", "S1b 下行期减产到50%(Q2-Q5)") else "替代方案多赚 "
    rows.append([name, alt, sign + f2(d), note])
table(rows, widths=[3.5, 4.4, 3.3, 4.8], font=8.5, align_num=False)

# ================= 5 如果重来 =================
doc.add_heading("五、如果重来，我会怎么改", level=1)
para("下面五条都是只用当时能看到的信息就能执行的规则，不依赖事后才知道的价格：")
bullet("开局就满开工。变动成本只有每千片 0.005，即使在价格底部也只占售价的四成左右；减产只在售价低于变动成本时才有意义。", "1. 产量按边际成本定，不按情绪定。")
bullet(f"借款上限按权益的 60% 计算、会随亏损下降。下行一开始就把借款加到上限附近（第 1 季上限 {lim[0]:.2f}），多付约 0.7%/季的净利差，换来底部扩产的弹药。", "2. 在上限高的时候锁定流动性。")
bullet("投产滞后 4 季、折旧从投产才开始。只要「4 季后的预期需求缺口」存在、单千片产能回本期短于剩余季数，就继续扩产；用一致预期是否连续上调、库存天数是否快速下降来判断，不用新闻里的定性词判断顶部。", "3. 扩产按「量」的信号滚动决策。")
bullet("当库存天数低于约 30 天、而一致预期仍在上调时，只卖当季产出，把库存留给后面的高价季度；最后两季再全部清空。", "4. 惜售要有触发条件。")
bullet("HBM 的安全余量要跟着普通内存的机会成本走：普通内存过剩时可以多留，普通内存缺货时压到 1%～2%。", "5. HBM 排产的机会成本要按市价算。")
rows = [["方案（累加）", "期末权益（估算）", "比实际多", "最低现金", "最高借款", "紧急授信"]]
for k, lab in [("R0 只做：Q1满开工", "R0：第 1 季满开工"), ("R1 Q1满开工+Q1预借14.9", "R1：+ 第 1 季预借 14.90"),
               ("R2 R1+Q5-Q7每季多扩1.5", "R2：+ 第 5–7 季每季多扩 1.50"), ("R3 R2+Q11/Q12每季多扩3.6", "R3：+ 第 11–12 季每季多扩 3.60")]:
    v = redo[k]
    rows.append([lab, f2(v["final_equity"]), f2(v["delta"]), f2(v["min_cash"]), f2(v["max_debt"]), "未动用" if v["em_peak"] < 1e-9 else f2(v["em_peak"])])
table(rows, widths=[5.6, 2.4, 2.0, 2.0, 2.0, 2.0], font=8.5)
para(f"按这四步累加，估算期末权益约 {f2(redo['R3 R2+Q11/Q12每季多扩3.6']['final_equity'])}，比实际多 {f2(redo['R3 R2+Q11/Q12每季多扩3.6']['delta'])}，且不需要紧急授信。"
     f"如果再加上 HBM 精确排产（第 5 条的理想情况）还能多约 {f2(S('S6 HBM按需求精确排产'))}。这些都是模型估算，依赖下文的推断参数。", color=INK2)

# ================= 附录 =================
doc.add_heading("附录 A　逐季关键数字", level=1)
rows = [["季度", "价格指数", "HBM价格指数", "产能", "开工率", "HBM产量", "营业收入", "净利润", "毛利率", "现金", "借款", "股东权益", "库存(千片)"]]
for q in Q:
    rows.append([f"Q{q['quarter']}", f"{q['price_index']:.2f}", f"{_hu(q['hbm_price_index'], 1):.1f}", f"{q['capacity_kwafers']:.0f}", pct(q["utilization"], 0),
                 f"{_hu(q['hbm_kwafers'], 1):.1f}", f2(q["is_revenue"]), f2(q["is_net_income"]), pct(q["is_gross_margin"], 1), f2(q["bs_cash"]), f2(q["bs_debt"]),
                 f2(q["bs_equity"]), f"{_hu(q['bs_inventory_kwafers'], 1):.1f}"])
table(rows, widths=[1.0, 1.2, 1.4, 1.0, 1.1, 1.3, 1.4, 1.3, 1.3, 1.3, 1.1, 1.5, 1.4], font=7.5)
para("来源：交付/data/raw/export_final.json；与交付/三张表.xlsx 的利润表、资产负债表、现金流量表一致（工作簿里的公式逐项核对，差额为 0）。", size=8.5, color=INK2)

doc.add_heading("附录 B　推断参数与方法（假设）", level=1)
para("规则只给了框架。下面的参数是我从报表反推、再用 16 季全部数据逐季核对得到的，属于推断，不是引擎公开的事实：")
rows = [["参数", "数值", "依据"],
        ["价格指数 1.00 对应的普通内存售价", "每千片 0.024", "第 1 季收入 4.608 ÷ 销量 240 ÷ 指数 0.80"],
        ["HBM 售价", "每千片 0.024 × HBM 价格指数", "第 4 季起逐季核对"],
        ["每千片变动成本", "0.005", "第 1 季（产 255）与第 2 季（产 300）单位成本对比"],
        ["工厂固定费用（不含折旧）", "每季 1.00", "同上；扩产后不变"],
        ["折旧", "2.00 + 已投产资本开支 ÷ 20", "资本开支 4 季后投产并开始折旧"],
        ["固定研发、销售及管理费用", "每季 0.80、0.25", "16 季均相同"],
        ["净利息", "1.5%×期初借款 − 0.8%×期初现金", "16 季逐季核对"]]
table(rows, widths=[5.0, 4.2, 6.8], font=8.5, align_num=False)
para("情景估算用的复现模型（交付/scripts/sim.py，Excel「模型复现与情景」页是同一套公式）另外做了三条假设：HBM 产量 = 产量 × HBM 占比（我从没在开工率低于 100% 时生产过 HBM，无法验证）；"
     "现金低于 0.5 时紧急授信补足到 0.5；HBM 需求不受我们何时拿到资格影响。模型把实际 16 季的权益、现金、净利润逐季复现到 0.001 以内，但情景结果仍然是估算。")
doc.add_heading("附录 C　数据出处", level=1)
for t in ["交付/data/raw/state_before_Q01.json：开局状态（期初资产负债、借款上限、信号）。",
          "交付/data/raw/decide_Q01.json … decide_Q16.json：每季提交后引擎的返回（当季三张表、下一季信号）。",
          "交付/data/raw/export_final.json：16 季结束后的导出，本报告所有「实际」数字的出处。",
          "交付/data/raw/rules.md：引擎规则原文。",
          "交付/决策日志.md：每季看到了什么、怎么判断、做了什么、为什么。",
          "交付/三张表.xlsx：逐季三张表、分析指标和情景模型（全部公式联动）。"]:
    bullet(t)

out = ROOT / "经营复盘.docx"
doc.save(out)
print("saved", out)
