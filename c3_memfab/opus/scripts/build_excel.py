"""生成 交付/三张表.xlsx。
- 「原始数据」页：引擎导出数字（输入格，黄底蓝字），逐项写明来源字段。
- 「假设」页：规则给定的参数和从报表反推的参数（输入格），每项标明性质（规则/推断/假设）。
- 三张表、分析、模型复现与情景：全部是公式（白底黑字），引用原始数据和假设；核对行（绿底）显示与引擎数字的差额。
运行：python3 交付/scripts/build_excel.py
"""
import json, pathlib
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, Font as DFont

ROOT = pathlib.Path(__file__).resolve().parents[1]
D = json.load(open(ROOT / "data/quarterly.json"))
Q, O = D["quarters"], D["opening"]
NQ = len(Q)
C0 = 3                      # Q1 所在列（C）
QC = [L(C0 + i) for i in range(NQ)]   # C..R
TOT = L(C0 + NQ)            # S 合计列
FIRST, LAST = QC[0], QC[-1]

# ---------- 样式 ----------
FONT = "Microsoft YaHei"
F_IN = PatternFill("solid", start_color="FFF4CC", end_color="FFF4CC")      # 输入格：浅黄
F_CALC = PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF")    # 公式格：白
F_CHK = PatternFill("solid", start_color="E3F1DF", end_color="E3F1DF")     # 核对格：浅绿
F_HDR = PatternFill("solid", start_color="1F3864", end_color="1F3864")
F_SEC = PatternFill("solid", start_color="DCE8F5", end_color="DCE8F5")
F_KEY = PatternFill("solid", start_color="F3F0FA", end_color="F3F0FA")     # 关键结果：浅紫
F_NA = PatternFill("solid", start_color="EFEFEF", end_color="EFEFEF")
FT_IN = Font(name=FONT, color="1F4E99", size=10)
FT_CALC = Font(name=FONT, color="111111", size=10)
FT_BOLD = Font(name=FONT, color="111111", size=10, bold=True)
FT_HDR = Font(name=FONT, color="FFFFFF", size=10, bold=True)
FT_TITLE = Font(name=FONT, color="1F3864", size=14, bold=True)
FT_NOTE = Font(name=FONT, color="595959", size=9)
FT_CHK = Font(name=FONT, color="2E7D32", size=9)
thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
NF = {"money": '#,##0.00;[Red]-#,##0.00', "money3": '#,##0.000;[Red]-#,##0.000', "pct": '0.0%;[Red]-0.0%', "idx": '0.00',
      "kw": '#,##0.0', "days": '0.0', "musd": '#,##0.0', "x": '0.00"x"', "chk": '0.0000;[Red]-0.0000;0.0000', "int": '0'}

wb = Workbook()

def title(ws, text, sub=None):
    ws["A1"] = text; ws["A1"].font = FT_TITLE
    if sub:
        ws["A2"] = sub; ws["A2"].font = FT_NOTE
    ws.sheet_view.showGridLines = False

def header(ws, row, first_label="项目", second_label="单位/口径", opening=False, total=None):
    ws.cell(row, 1, first_label); ws.cell(row, 2, second_label if not opening else "期初")
    for i in range(NQ):
        ws.cell(row, C0 + i, f"Q{i+1}")
    if total:
        ws.cell(row, C0 + NQ, total)
    last = C0 + NQ if total else C0 + NQ - 1
    for c in range(1, last + 1):
        cell = ws.cell(row, c); cell.fill = F_HDR; cell.font = FT_HDR; cell.alignment = Alignment(horizontal="center", vertical="center")

def style_row(ws, row, kind, nf, c_from=C0, c_to=C0 + NQ - 1):
    fill, font = {"in": (F_IN, FT_IN), "calc": (F_CALC, FT_CALC), "chk": (F_CHK, FT_CHK), "key": (F_KEY, FT_BOLD)}[kind]
    for c in range(c_from, c_to + 1):
        cell = ws.cell(row, c); cell.fill = fill; cell.font = font; cell.number_format = NF[nf]; cell.border = BORDER

def label(ws, row, text, unit="", bold=False, indent=0):
    a = ws.cell(row, 1, text); a.font = FT_BOLD if bold else FT_CALC; a.alignment = Alignment(indent=indent)
    b = ws.cell(row, 2, unit); b.font = FT_NOTE
    for c in (1, 2):
        ws.cell(row, c).border = BORDER

def section(ws, row, text, width=C0 + NQ):
    for c in range(1, width + 1):
        ws.cell(row, c).fill = F_SEC
    ws.cell(row, 1, text).font = FT_BOLD

def widths(ws, a=30, b=16, q=9.5, extra=None):
    ws.column_dimensions["A"].width = a; ws.column_dimensions["B"].width = b
    for i in range(NQ + 1):
        ws.column_dimensions[L(C0 + i)].width = q
    for k, v in (extra or {}).items():
        ws.column_dimensions[k].width = v

def ref(sheet, cell):
    return f"'{sheet}'!{cell}"

# =====================================================================
# 1. 说明
# =====================================================================
ws0 = wb.active; ws0.title = "说明"
title(ws0, "内存芯片公司 16 季经营：三张表与分析", "游戏编号 memfab-opus · 数据导出日期 2026-10-22 · 金额单位：十亿美元 · 片数单位：千片/季")
lines = [
    ("这本工作簿里有什么", None),
    ("原始数据", "引擎导出的逐季数字，原样录入（输入格）。来源：交付/data/raw/export_final.json 的 history[i] 各字段；期初数来自 state_before_Q01.json。"),
    ("假设", "规则给定的参数（如税率15%、借款季利率1.5%）和我从报表反推的参数（如每千片变动成本0.005），每项注明性质和出处。"),
    ("利润表 / 资产负债表 / 现金流量表", "逐季三张表。小计、合计、期末余额全部用公式从原始数据算出，三张表相互勾稽；底部绿色核对行显示与引擎数字的差额（应为0）。"),
    ("分析", "毛利率、存货周转天数、经营现金流、自由现金流、净负债、ROE等指标，全部是公式；下方配图表。改「假设」页的每季天数、年化倍数等，结果自动变。"),
    ("模型复现与情景", "用「假设」页的参数和每季决策，从头重算产量、成本、存货、收入、利润、现金和权益。默认复现实际结果（核对行≈0）；在黄色情景格里改决策或价格系数，期末权益等结果自动重算。"),
    ("决策备注", "每季提交给引擎的决策备注原文。"),
    ("", None),
    ("颜色约定", None),
    ("浅黄底、蓝字", "输入格：引擎导出的原始数字、规则参数、推断参数、情景输入。可以改。"),
    ("白底、黑字", "公式格：由输入格计算而来，不要手改。"),
    ("浅绿底、绿字", "核对格：公式结果与引擎数字的差额，正常应为0（四舍五入误差在±0.01以内）。"),
    ("浅紫底", "关键结果（公式）。"),
    ("", None),
    ("口径说明", None),
    ("营业成本", "引擎的营业成本含存货跌价损失（期末按可变现净值=0.95×市价减记，不转回），表中单列「其中：存货跌价」。HBM 片按订单生产，多产作废，其成本计入当季营业成本。"),
    ("存货周转天数", "公式：（期初存货+期末存货）/2 ÷ 当季营业成本 × 每季天数（默认91.25天，假设）。另给片数口径：期末库存千片 ÷ 当季普通内存销量 × 每季天数。"),
    ("净负债", "公式：借款+紧急授信−现金；负数表示净现金。"),
    ("ROE", "公式：季度ROE=净利润÷（期初权益+期末权益）/2；年化=季度×4（年化倍数是假设，可改）。"),
    ("推断参数", "「假设」页里标为「推断」的数字不是规则原文，是我用前两季报表反推、再用16季数据逐季核对得到的（模型复现核对行≈0）。它们是假设，不是引擎公开的事实。"),
    ("免责声明", "这是经营模拟的研究和展示材料，市场、公司和数字都是虚构的，不构成任何投资建议。"),
]
r = 4
for k, v in lines:
    if v is None:
        ws0.cell(r, 1, k).font = FT_BOLD; ws0.cell(r, 1).fill = F_SEC; ws0.cell(r, 2).fill = F_SEC
    else:
        ws0.cell(r, 1, k).font = FT_CALC
        c = ws0.cell(r, 2, v); c.font = FT_CALC; c.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
ws0["A11"].fill = F_IN; ws0["A12"].fill = F_CALC; ws0["A13"].fill = F_CHK; ws0["A14"].fill = F_KEY
ws0["A11"].font = FT_IN; ws0["A13"].font = FT_CHK
ws0.column_dimensions["A"].width = 30; ws0.column_dimensions["B"].width = 120
for rr in range(5, r):
    ws0.row_dimensions[rr].height = 30

# =====================================================================
# 2. 假设
# =====================================================================
wa = wb.create_sheet("假设")
title(wa, "假设与参数", "黄色为输入格。性质：规则=规则原文给定；推断=从报表反推并用16季数据核对；假设=分析口径。")
for c, t in enumerate(["参数", "数值", "单位", "性质", "出处/说明"], start=1):
    cell = wa.cell(4, c, t); cell.fill = F_HDR; cell.font = FT_HDR
PARAMS = [
    # name, label, value, unit, nature, source, numfmt
    ("P_BASE", "价格指数1.00对应的普通内存售价", 0.024, "十亿美元/千片", "推断", "Q1 收入4.608÷销量240÷指数0.80；16季逐季核对", "money3"),
    ("V_COST", "每千片变动成本", 0.005, "十亿美元/千片", "推断", "Q1（产255）与Q2（产300）单位成本对比反推", "money3"),
    ("PLANT_FIXED", "工厂固定费用（不含折旧）", 1.0, "十亿美元/季", "推断", "同上；扩产后不变（Q9起验证）", "money"),
    ("DEP_BASE", "期初固定资产每季折旧", 2.0, "十亿美元/季", "推断", "Q1-Q4 折旧均为2.0", "money"),
    ("DEP_LIFE", "资本开支折旧期", 20, "季", "规则", "规则：资本开支按20个季度直线折旧", "int"),
    ("LAG_Q", "资本开支投产（并开始折旧）滞后", 4, "季", "规则/推断", "规则：超出维持部分4个季度后形成新产能；折旧也从投产季开始（推断）", "int"),
    ("CAP_COST", "每千片/季新增产能投资", 0.12, "十亿美元", "规则", "规则：每1.2亿美元增加1千片/季", "money"),
    ("MAINT", "每季维持性资本开支", 1.5, "十亿美元/季", "规则", "规则：每季先要15亿美元维持现有产能", "money"),
    ("RND_BASE", "每季固定研发费用", 0.8, "十亿美元/季", "推断", "Q1 研发3.8=固定0.8+HBM研发3.0；其余各季0.8", "money"),
    ("SGA_Q", "每季销售及管理费用", 0.25, "十亿美元/季", "推断", "16季均为0.25", "money"),
    ("SELL_COST", "存货减记测试中的销售费用率", 0.05, "比例", "规则", "规则：市价扣5%销售费用后低于成本则减记", "pct"),
    ("TAX_RATE", "所得税率（只对盈利征）", 0.15, "比例", "规则", "规则", "pct"),
    ("RATE_DEBT", "借款季利率", 0.015, "比例/季", "规则", "规则", "pct"),
    ("RATE_EMERG", "紧急授信季利率", 0.03, "比例/季", "规则", "规则", "pct"),
    ("EMERG_FLOOR", "触发紧急授信的现金下限", 0.5, "十亿美元", "规则", "规则：现金低于5亿美元自动动用紧急授信（补足到下限为假设）", "money"),
    ("RATE_CASH", "现金季收益率", 0.008, "比例/季", "推断", "净利息=1.5%×期初借款−0.8%×期初现金，16季逐季核对", "pct"),
    ("DEBT_RATIO", "借款上限占股东权益比例", 0.6, "比例", "规则", "规则：约为股东权益的60%", "pct"),
    ("DEBT_FLOOR", "借款上限下限", 8.0, "十亿美元", "规则", "规则：不低于80亿美元", "money"),
    ("HBM_NEED", "HBM研发累计门槛", 3.0, "十亿美元", "规则", "规则：累计满30亿美元后再过2个季度可量产", "money"),
    ("DAYS_Q", "每季天数（周转天数用）", 91.25, "天", "假设", "365÷4", "days"),
    ("ANNUAL", "年化倍数（ROE年化用）", 4, "倍", "假设", "一年4个季度，简单年化", "int"),
    ("PRICE_SCEN", "情景：普通内存价格系数", 1.0, "倍", "情景输入", "只影响「模型复现与情景」页；1.0=实际价格", "idx"),
    ("HBM_PRICE_SCEN", "情景：HBM价格系数", 1.0, "倍", "情景输入", "只影响「模型复现与情景」页；1.0=实际价格", "idx"),
]
for i, (name, lab, val, unit, nat, src, nf) in enumerate(PARAMS):
    r = 5 + i
    wa.cell(r, 1, lab).font = FT_CALC
    c = wa.cell(r, 2, val); c.fill = F_IN; c.font = FT_IN; c.number_format = NF[nf]; c.border = BORDER
    wa.cell(r, 3, unit).font = FT_NOTE; wa.cell(r, 4, nat).font = FT_NOTE; wa.cell(r, 5, src).font = FT_NOTE
    wb.defined_names[name] = DefinedName(name, attr_text=f"'假设'!$B${r}")
wa.column_dimensions["A"].width = 34; wa.column_dimensions["B"].width = 12; wa.column_dimensions["C"].width = 15
wa.column_dimensions["D"].width = 11; wa.column_dimensions["E"].width = 70

# =====================================================================
# 3. 原始数据
# =====================================================================
wr = wb.create_sheet("原始数据")
title(wr, "引擎导出数据（原样录入，输入格）", "来源：交付/data/raw/export_final.json → history[第i季]；B列为字段路径。金额：十亿美元；片数：千片/季。")
header(wr, 4, "项目", "来源字段")
RAW_FIELDS = [
    ("决策", None, None, None),
    ("utilization", "开工率", "decision.utilization", "pct"),
    ("sell_fraction", "出货比例", "decision.sell_fraction", "pct"),
    ("capex", "资本开支", "decision.capex", "money"),
    ("hbm_rnd", "HBM 研发投入", "decision.hbm_rnd", "money"),
    ("hbm_share", "HBM 产能占比", "decision.hbm_share", "pct"),
    ("borrow", "新增借款（还款为负）", "decision.borrow", "money"),
    ("市场", None, None, None),
    ("price_index", "普通内存价格指数", "market.price_index", "idx"),
    ("hbm_price_index", "HBM 价格指数", "market.hbm_price_index", "idx"),
    ("demand_kwafers", "普通内存需求（千片）", "market.demand_kwafers", "kw"),
    ("hbm_demand_kwafers", "HBM 需求（千片）", "market.hbm_demand_kwafers", "kw"),
    ("运营", None, None, None),
    ("capacity_kwafers", "产能（千片/季）", "operations.capacity_kwafers", "kw"),
    ("produced_kwafers", "产量（千片）", "operations.produced_kwafers", "kw"),
    ("hbm_kwafers", "其中 HBM 产量（千片）", "operations.hbm_kwafers", "kw"),
    ("std_sold_kwafers", "普通内存销量（千片）", "operations.std_sold_kwafers", "kw"),
    ("hbm_sold_kwafers", "HBM 销量（千片）", "operations.hbm_sold_kwafers", "kw"),
    ("unit_cost_musd", "单位生产成本（百万美元/千片）", "operations.unit_cost_musd", "musd"),
    ("hbm_qualified", "HBM 量产资格（1=有）", "operations.hbm_qualified", "int"),
    ("利润表（引擎）", None, None, None),
    ("is_revenue", "营业收入", "income_statement.revenue", "money"),
    ("is_revenue_std", "普通内存收入", "income_statement.revenue_std", "money"),
    ("is_revenue_hbm", "HBM 收入", "income_statement.revenue_hbm", "money"),
    ("is_cogs", "营业成本（含存货跌价）", "income_statement.cogs", "money"),
    ("is_inventory_writedown", "存货跌价损失", "income_statement.inventory_writedown", "money"),
    ("is_gross_profit", "毛利", "income_statement.gross_profit", "money"),
    ("is_rnd", "研发费用", "income_statement.rnd", "money"),
    ("is_sga", "销售及管理费用", "income_statement.sga", "money"),
    ("is_operating_income", "营业利润", "income_statement.operating_income", "money"),
    ("is_net_interest", "净利息支出", "income_statement.net_interest", "money"),
    ("is_tax", "所得税", "income_statement.tax", "money"),
    ("is_net_income", "净利润", "income_statement.net_income", "money"),
    ("现金流量（引擎）", None, None, None),
    ("cf_operating", "经营活动现金流", "cash_flow.operating", "money"),
    ("cf_investing", "投资活动现金流", "cash_flow.investing", "money"),
    ("cf_financing", "筹资活动现金流（不含紧急授信）", "cash_flow.financing", "money"),
    ("cf_depreciation", "折旧", "cash_flow.depreciation", "money"),
    ("cf_emergency_borrowing", "紧急授信提取", "cash_flow.emergency_borrowing", "money"),
    ("资产负债（引擎，期末）", None, None, None),
    ("bs_cash", "现金", "balance_sheet.cash", "money"),
    ("bs_inventory", "存货", "balance_sheet.inventory", "money"),
    ("bs_ppe_net", "固定资产净值", "balance_sheet.ppe_net", "money"),
    ("bs_total_assets", "资产合计", "balance_sheet.total_assets", "money"),
    ("bs_debt", "借款", "balance_sheet.debt", "money"),
    ("bs_emergency_debt", "紧急授信余额", "balance_sheet.emergency_debt", "money"),
    ("bs_equity", "股东权益", "balance_sheet.equity", "money"),
    ("bs_inventory_kwafers", "库存（千片）", "balance_sheet.inventory_kwafers", "kw"),
]
RAW = {}
r = 5
for key, lab, src, nf in RAW_FIELDS:
    if lab is None:
        section(wr, r, key); r += 1; continue
    label(wr, r, lab, src)
    for i, q in enumerate(Q):
        v = q[key]
        if isinstance(v, bool):
            v = 1 if v else 0
        wr.cell(r, C0 + i, v)
    style_row(wr, r, "in", nf)
    RAW[key] = r
    r += 1
wr.freeze_panes = "C5"
widths(wr, a=30, b=34, q=9.5)

def R(key, col):
    return f"'原始数据'!{col}{RAW[key]}"

# =====================================================================
# 4. 利润表
# =====================================================================
wi = wb.create_sheet("利润表")
title(wi, "利润表（逐季）", "单位：十亿美元。所有数字为公式：引用「原始数据」页，小计由公式计算；底部为与引擎数字的核对。")
header(wi, 4, total="16季合计")
IS = {}
rows_is = [
    ("rev_std", "普通内存收入", "", lambda c: f"={R('is_revenue_std', c)}", "money", False),
    ("rev_hbm", "HBM 收入", "", lambda c: f"={R('is_revenue_hbm', c)}", "money", False),
    ("rev", "营业收入", "", None, "money", True),
    ("cogs", "营业成本（含存货跌价）", "", lambda c: f"={R('is_cogs', c)}", "money", False),
    ("wd", "  其中：存货跌价损失", "", lambda c: f"={R('is_inventory_writedown', c)}", "money", False),
    ("gp", "毛利", "", None, "money", True),
    ("gm", "毛利率", "%", None, "pct", False),
    ("rnd", "研发费用", "", lambda c: f"={R('is_rnd', c)}", "money", False),
    ("rnd_hbm", "  其中：HBM 专项研发", "", lambda c: f"={R('hbm_rnd', c)}", "money", False),
    ("sga", "销售及管理费用", "", lambda c: f"={R('is_sga', c)}", "money", False),
    ("oi", "营业利润", "", None, "money", True),
    ("int", "净利息支出（收入为负）", "", lambda c: f"={R('is_net_interest', c)}", "money", False),
    ("pbt", "税前利润", "", None, "money", True),
    ("tax", "所得税", "", lambda c: f"={R('is_tax', c)}", "money", False),
    ("ni", "净利润", "", None, "money", True),
    ("cum_ni", "累计净利润", "", None, "money", False),
]
r = 5
for key, lab, unit, f, nf, bold in rows_is:
    IS[key] = r; label(wi, r, lab, unit, bold=bold); r += 1
for i, c in enumerate(QC):
    prev = QC[i - 1] if i else None
    for key, lab, unit, f, nf, bold in rows_is:
        rr = IS[key]
        if f:
            wi[f"{c}{rr}"] = f(c)
    wi[f"{c}{IS['rev']}"] = f"={c}{IS['rev_std']}+{c}{IS['rev_hbm']}"
    wi[f"{c}{IS['gp']}"] = f"={c}{IS['rev']}-{c}{IS['cogs']}"
    wi[f"{c}{IS['gm']}"] = f"=IF({c}{IS['rev']}=0,0,{c}{IS['gp']}/{c}{IS['rev']})"
    wi[f"{c}{IS['oi']}"] = f"={c}{IS['gp']}-{c}{IS['rnd']}-{c}{IS['sga']}"
    wi[f"{c}{IS['pbt']}"] = f"={c}{IS['oi']}-{c}{IS['int']}"
    wi[f"{c}{IS['ni']}"] = f"={c}{IS['pbt']}-{c}{IS['tax']}"
    wi[f"{c}{IS['cum_ni']}"] = f"={c}{IS['ni']}" if i == 0 else f"={prev}{IS['cum_ni']}+{c}{IS['ni']}"
for key, lab, unit, f, nf, bold in rows_is:
    style_row(wi, IS[key], "key" if key in ("ni",) else "calc", nf, c_to=C0 + NQ)
    if key not in ("gm", "cum_ni"):
        wi[f"{TOT}{IS[key]}"] = f"=SUM({FIRST}{IS[key]}:{LAST}{IS[key]})"
wi[f"{TOT}{IS['gm']}"] = f"={TOT}{IS['gp']}/{TOT}{IS['rev']}"
wi[f"{TOT}{IS['cum_ni']}"] = f"={LAST}{IS['cum_ni']}"
for key in ("rev", "gp", "oi", "pbt", "ni"):
    for cc in range(1, 3):
        wi.cell(IS[key], cc).font = FT_BOLD
r = IS["cum_ni"] + 2
section(wi, r, "核对（公式结果 − 引擎数字，应为 0）"); r += 1
chk_is = [
    ("chk_rev", "营业收入差额", lambda c: f"={c}{IS['rev']}-{R('is_revenue', c)}"),
    ("chk_gp", "毛利差额", lambda c: f"={c}{IS['gp']}-{R('is_gross_profit', c)}"),
    ("chk_oi", "营业利润差额", lambda c: f"={c}{IS['oi']}-{R('is_operating_income', c)}"),
    ("chk_ni", "净利润差额", lambda c: f"={c}{IS['ni']}-{R('is_net_income', c)}"),
    ("chk_tax", "所得税 − 税率×MAX(0,税前利润)", lambda c: f"={c}{IS['tax']}-TAX_RATE*MAX(0,{c}{IS['pbt']})"),
]
for key, lab, f in chk_is:
    IS[key] = r; label(wi, r, lab, "应为0")
    for c in QC:
        wi[f"{c}{r}"] = f(c)
    style_row(wi, r, "chk", "chk"); r += 1
wi.freeze_panes = "C5"; widths(wi, a=28, b=8, q=9.5, extra={TOT: 11})

# =====================================================================
# 5. 现金流量表（先建，资产负债表的现金引用它）
# =====================================================================
wc = wb.create_sheet("现金流量表")
wbs = wb.create_sheet("资产负债表")
wb.move_sheet("资产负债表", offset=-1)   # 顺序：利润表、资产负债表、现金流量表
title(wc, "现金流量表（逐季，间接法）", "单位：十亿美元。经营现金流=净利润+折旧−存货增加（本模拟没有应收应付）；期末现金滚动计算后与引擎核对。")
header(wc, 4, total="16季合计")
CF = {}
rows_cf = [
    ("ni", "净利润", True), ("dep", "加：折旧", False), ("dinv", "减：存货增加（减少为正）", False), ("ocf", "经营活动现金流", True), (None, None, None),
    ("capex", "资本开支", False), ("icf", "投资活动现金流", True), (None, None, None),
    ("borrow", "新增借款（还款为负）", False), ("emerg", "紧急授信提取", False), ("fcf_fin", "筹资活动现金流", True), (None, None, None),
    ("net", "现金净增加", True), ("beg", "期初现金", False), ("end", "期末现金", True), (None, None, None),
    ("fcf", "自由现金流（经营+投资）", True),
]
r = 5
for key, lab, bold in rows_cf:
    if key:
        CF[key] = r; label(wc, r, lab, "", bold=bold)
    r += 1
# 资产负债表行号（先定好，现金流量表要引用存货）
BS = {}
bs_rows = [("cash", "现金", False), ("inv", "存货", False), ("ppe", "固定资产净值", False), ("ta", "资产合计", True), (None, None, None),
           ("debt", "借款", False), ("emerg", "紧急授信", False), ("tl", "负债合计", True), ("eq", "股东权益", True), ("tle", "负债和股东权益合计", True)]
r = 5
for key, lab, bold in bs_rows:
    if key:
        BS[key] = r
    r += 1
for i, c in enumerate(QC):
    p = QC[i - 1] if i else "B"   # 资产负债表期初在B列
    wc[f"{c}{CF['ni']}"] = f"='利润表'!{c}{IS['ni']}"
    wc[f"{c}{CF['dep']}"] = f"={R('cf_depreciation', c)}"
    wc[f"{c}{CF['dinv']}"] = f"=-('资产负债表'!{c}{BS['inv']}-'资产负债表'!{p}{BS['inv']})"
    wc[f"{c}{CF['ocf']}"] = f"=SUM({c}{CF['ni']}:{c}{CF['dinv']})"
    wc[f"{c}{CF['capex']}"] = f"=-{R('capex', c)}"
    wc[f"{c}{CF['icf']}"] = f"={c}{CF['capex']}"
    wc[f"{c}{CF['borrow']}"] = f"={R('borrow', c)}"
    wc[f"{c}{CF['emerg']}"] = f"={R('cf_emergency_borrowing', c)}"
    wc[f"{c}{CF['fcf_fin']}"] = f"={c}{CF['borrow']}+{c}{CF['emerg']}"
    wc[f"{c}{CF['net']}"] = f"={c}{CF['ocf']}+{c}{CF['icf']}+{c}{CF['fcf_fin']}"
    wc[f"{c}{CF['beg']}"] = f"='资产负债表'!B{BS['cash']}" if i == 0 else f"={QC[i-1]}{CF['end']}"
    wc[f"{c}{CF['end']}"] = f"={c}{CF['beg']}+{c}{CF['net']}"
    wc[f"{c}{CF['fcf']}"] = f"={c}{CF['ocf']}+{c}{CF['icf']}"
for key, lab, bold in rows_cf:
    if not key:
        continue
    style_row(wc, CF[key], "key" if key in ("ocf", "fcf") else "calc", "money", c_to=C0 + NQ)
    if key in ("beg",):
        wc[f"{TOT}{CF[key]}"] = f"={FIRST}{CF[key]}"
    elif key in ("end",):
        wc[f"{TOT}{CF[key]}"] = f"={LAST}{CF[key]}"
    else:
        wc[f"{TOT}{CF[key]}"] = f"=SUM({FIRST}{CF[key]}:{LAST}{CF[key]})"
r = CF["fcf"] + 2
section(wc, r, "核对（公式结果 − 引擎数字，应为 0）"); r += 1
for key, lab, f in [
    ("chk_ocf", "经营活动现金流差额", lambda c: f"={c}{CF['ocf']}-{R('cf_operating', c)}"),
    ("chk_icf", "投资活动现金流差额", lambda c: f"={c}{CF['icf']}-{R('cf_investing', c)}"),
    ("chk_fin", "筹资活动现金流差额（引擎筹资不含紧急授信）", lambda c: f"={c}{CF['borrow']}-{R('cf_financing', c)}"),
    ("chk_end", "期末现金差额", lambda c: f"={c}{CF['end']}-{R('bs_cash', c)}"),
]:
    CF[key] = r; label(wc, r, lab, "应为0")
    for c in QC:
        wc[f"{c}{r}"] = f(c)
    style_row(wc, r, "chk", "chk"); r += 1
wc.freeze_panes = "C5"; widths(wc, a=34, b=6, q=9.5, extra={TOT: 11})

# =====================================================================
# 6. 资产负债表
# =====================================================================
title(wbs, "资产负债表（逐季期末）", "单位：十亿美元。B列期初数来自 state_before_Q01.json（输入格）；之后各季：现金取自现金流量表，固定资产=上期+资本开支−折旧，借款=上期+新增借款，股东权益=上期+净利润。")
header(wbs, 4, opening=True)
wbs.cell(4, 2).value = "期初"
for key, lab, bold in bs_rows:
    if key:
        label(wbs, BS[key], lab, "", bold=bold)
        wbs.cell(BS[key], 2).value = None
open_vals = {"cash": O["cash"], "inv": O["inventory"], "ppe": O["ppe_net"], "debt": O["debt"], "emerg": O["emergency_debt"], "eq": O["equity"]}
for k, v in open_vals.items():
    cell = wbs.cell(BS[k], 2, v); cell.fill = F_IN; cell.font = FT_IN; cell.number_format = NF["money"]; cell.border = BORDER
for key, f in [("ta", lambda c: f"={c}{BS['cash']}+{c}{BS['inv']}+{c}{BS['ppe']}"), ("tl", lambda c: f"={c}{BS['debt']}+{c}{BS['emerg']}"),
               ("tle", lambda c: f"={c}{BS['tl']}+{c}{BS['eq']}")]:
    cell = wbs.cell(BS[key], 2, f("B")); cell.fill = F_CALC; cell.font = FT_BOLD; cell.number_format = NF["money"]; cell.border = BORDER
for i, c in enumerate(QC):
    p = QC[i - 1] if i else "B"
    wbs[f"{c}{BS['cash']}"] = f"='现金流量表'!{c}{CF['end']}"
    wbs[f"{c}{BS['inv']}"] = f"={R('bs_inventory', c)}"
    wbs[f"{c}{BS['ppe']}"] = f"={p}{BS['ppe']}+{R('capex', c)}-{R('cf_depreciation', c)}"
    wbs[f"{c}{BS['ta']}"] = f"={c}{BS['cash']}+{c}{BS['inv']}+{c}{BS['ppe']}"
    wbs[f"{c}{BS['debt']}"] = f"={p}{BS['debt']}+{R('borrow', c)}"
    wbs[f"{c}{BS['emerg']}"] = f"={p}{BS['emerg']}+{R('cf_emergency_borrowing', c)}"
    wbs[f"{c}{BS['tl']}"] = f"={c}{BS['debt']}+{c}{BS['emerg']}"
    wbs[f"{c}{BS['eq']}"] = f"={p}{BS['eq']}+'利润表'!{c}{IS['ni']}"
    wbs[f"{c}{BS['tle']}"] = f"={c}{BS['tl']}+{c}{BS['eq']}"
for key, lab, bold in bs_rows:
    if key:
        style_row(wbs, BS[key], "key" if key == "eq" else "calc", "money")
r = BS["tle"] + 2
section(wbs, r, "备查（片数与单位价值）"); r += 1
BS["kw"] = r; label(wbs, r, "库存（千片）", "千片")
c = wbs.cell(r, 2, O["inventory_kwafers"]); c.fill = F_IN; c.font = FT_IN; c.number_format = NF["kw"]; c.border = BORDER
for col in QC:
    wbs[f"{col}{r}"] = f"={R('bs_inventory_kwafers', col)}"
style_row(wbs, r, "calc", "kw"); r += 1
BS["unit_inv"] = r; label(wbs, r, "每千片库存账面价值", "百万美元/千片")
for col in ["B"] + QC:
    wbs[f"{col}{r}"] = f"=IF({col}{BS['kw']}>0,{col}{BS['inv']}/{col}{BS['kw']}*1000,0)"
style_row(wbs, r, "calc", "musd", c_from=2); r += 1
BS["cap"] = r; label(wbs, r, "产能（千片/季）", "千片/季")
c = wbs.cell(r, 2, O["capacity_kwafers"]); c.fill = F_IN; c.font = FT_IN; c.number_format = NF["kw"]; c.border = BORDER
for col in QC:
    wbs[f"{col}{r}"] = f"={R('capacity_kwafers', col)}"
style_row(wbs, r, "calc", "kw"); r += 2
section(wbs, r, "核对（应为 0）"); r += 1
for key, lab, f in [
    ("chk_bal", "资产 − 负债 − 股东权益", lambda c: f"={c}{BS['ta']}-{c}{BS['tle']}"),
    ("chk_cash", "现金 − 引擎期末现金", lambda c: f"={c}{BS['cash']}-{R('bs_cash', c)}"),
    ("chk_ppe", "固定资产 − 引擎数", lambda c: f"={c}{BS['ppe']}-{R('bs_ppe_net', c)}"),
    ("chk_ta", "资产合计 − 引擎数", lambda c: f"={c}{BS['ta']}-{R('bs_total_assets', c)}"),
    ("chk_debt", "借款 − 引擎数", lambda c: f"={c}{BS['debt']}-{R('bs_debt', c)}"),
    ("chk_eq", "股东权益 − 引擎数", lambda c: f"={c}{BS['eq']}-{R('bs_equity', c)}"),
]:
    BS[key] = r; label(wbs, r, lab, "应为0")
    for col in QC:
        wbs[f"{col}{r}"] = f(col)
    style_row(wbs, r, "chk", "chk"); r += 1
wbs.freeze_panes = "C5"; widths(wbs, a=26, b=10, q=9.5)

# =====================================================================
# 7. 分析
# =====================================================================
wn = wb.create_sheet("分析")
title(wn, "经营分析（全部为公式）", "引用三张表、原始数据和「假设」页参数（每季天数、年化倍数、价格基准等）。改假设，指标和图表自动更新。")
header(wn, 4, "指标", "口径", total="全程")
AN = {}
IS_ = lambda key, c: f"'利润表'!{c}{IS[key]}"
BS_ = lambda key, c: f"'资产负债表'!{c}{BS[key]}"
CF_ = lambda key, c: f"'现金流量表'!{c}{CF[key]}"
def prevc(i):
    return QC[i - 1] if i else "B"
AN_ROWS = [
    ("盈利能力", None),
    ("rev", "营业收入", "十亿美元", "money", lambda c, i: f"={IS_('rev', c)}", f"=SUM({FIRST}{{r}}:{LAST}{{r}})"),
    ("rev_g", "营业收入环比增速", "%", "pct", lambda c, i: None if i == 0 else f"={c}{{rev}}/{QC[i-1]}{{rev}}-1", None),
    ("gm", "毛利率", "毛利÷营收", "pct", lambda c, i: f"={IS_('gp', c)}/{IS_('rev', c)}", f"='利润表'!{TOT}{IS['gp']}/'利润表'!{TOT}{IS['rev']}"),
    ("om", "营业利润率", "营业利润÷营收", "pct", lambda c, i: f"={IS_('oi', c)}/{IS_('rev', c)}", f"='利润表'!{TOT}{IS['oi']}/'利润表'!{TOT}{IS['rev']}"),
    ("nm", "净利率", "净利润÷营收", "pct", lambda c, i: f"={IS_('ni', c)}/{IS_('rev', c)}", f"='利润表'!{TOT}{IS['ni']}/'利润表'!{TOT}{IS['rev']}"),
    ("ebitda", "EBITDA", "营业利润+折旧", "money", lambda c, i: f"={IS_('oi', c)}+{CF_('dep', c)}", f"=SUM({FIRST}{{r}}:{LAST}{{r}})"),
    ("hbm_share_rev", "HBM 收入占比", "HBM收入÷营收", "pct", lambda c, i: f"={IS_('rev_hbm', c)}/{IS_('rev', c)}", f"='利润表'!{TOT}{IS['rev_hbm']}/'利润表'!{TOT}{IS['rev']}"),
    ("单位经济（每千片，百万美元）", None),
    ("p_std", "普通内存售价", "价格指数×基准价×1000", "musd", lambda c, i: f"={R('price_index', c)}*P_BASE*1000", None),
    ("p_hbm", "HBM 售价", "HBM价格指数×基准价×1000", "musd", lambda c, i: f"={R('hbm_price_index', c)}*P_BASE*1000", None),
    ("uc", "单位生产成本（含折旧）", "引擎单位成本", "musd", lambda c, i: f"={R('unit_cost_musd', c)}", None),
    ("vc", "单位变动成本", "假设页参数", "musd", lambda c, i: "=V_COST*1000", None),
    ("cash_cost", "单位现金成本", "(工厂费用+变动成本+固定研发+管理费)÷产量", "musd", lambda c, i: f"=(PLANT_FIXED+V_COST*{R('produced_kwafers', c)}+RND_BASE+SGA_Q)/{R('produced_kwafers', c)}*1000", None),
    ("std_margin", "普通内存单片毛利", "售价−单位生产成本", "musd", lambda c, i: f"={c}{{p_std}}-{c}{{uc}}", None),
    ("hbm_premium", "HBM 售价 ÷ 普通内存售价", "倍", "x", lambda c, i: f"={c}{{p_hbm}}/{c}{{p_std}}", None),
    ("运营", None),
    ("util", "产能利用率", "产量÷产能", "pct", lambda c, i: f"={R('produced_kwafers', c)}/{R('capacity_kwafers', c)}", None),
    ("fill", "普通内存需求满足率", "销量÷需求", "pct", lambda c, i: f"={R('std_sold_kwafers', c)}/{R('demand_kwafers', c)}", f"=SUM('原始数据'!{FIRST}{RAW['std_sold_kwafers']}:{LAST}{RAW['std_sold_kwafers']})/SUM('原始数据'!{FIRST}{RAW['demand_kwafers']}:{LAST}{RAW['demand_kwafers']})"),
    ("gap", "普通内存需求缺口（千片）", "需求−销量", "kw", lambda c, i: f"={R('demand_kwafers', c)}-{R('std_sold_kwafers', c)}", f"=SUM({FIRST}{{r}}:{LAST}{{r}})"),
    ("hbm_scrap", "HBM 作废量（千片）", "HBM产量−HBM销量", "kw", lambda c, i: f"={R('hbm_kwafers', c)}-{R('hbm_sold_kwafers', c)}", f"=SUM({FIRST}{{r}}:{LAST}{{r}})"),
    ("dio", "存货周转天数（金额口径）", "平均存货÷营业成本×每季天数", "days", lambda c, i: f"=AVERAGE({BS_('inv', prevc(i))},{BS_('inv', c)})/{IS_('cogs', c)}*DAYS_Q", None),
    ("dsi_kw", "库存可售天数（片数口径）", "期末库存千片÷普通内存销量×每季天数", "days", lambda c, i: f"={R('bs_inventory_kwafers', c)}/{R('std_sold_kwafers', c)}*DAYS_Q", None),
    ("现金与杠杆", None),
    ("ocf", "经营活动现金流", "现金流量表", "money", lambda c, i: f"={CF_('ocf', c)}", f"=SUM({FIRST}{{r}}:{LAST}{{r}})"),
    ("fcf", "自由现金流", "经营+投资", "money", lambda c, i: f"={CF_('fcf', c)}", f"=SUM({FIRST}{{r}}:{LAST}{{r}})"),
    ("ocf_ni", "经营现金流 ÷ 净利润", "倍", "x", lambda c, i: f"=IF({IS_('ni', c)}>0,{CF_('ocf', c)}/{IS_('ni', c)},0)", None),
    ("nd", "净负债", "借款+紧急授信−现金", "money", lambda c, i: f"={BS_('tl', c)}-{BS_('cash', c)}", None),
    ("nd_eq", "净负债 ÷ 股东权益", "%", "pct", lambda c, i: f"={c}{{nd}}/{BS_('eq', c)}", None),
    ("lev", "资产负债率", "负债合计÷资产合计", "pct", lambda c, i: f"={BS_('tl', c)}/{BS_('ta', c)}", None),
    ("limit", "借款上限（按期末权益）", "MAX(下限,比例×权益)", "money", lambda c, i: f"=MAX(DEBT_FLOOR,DEBT_RATIO*{BS_('eq', c)})", None),
    ("limit_use", "借款占上限比例", "借款÷借款上限", "pct", lambda c, i: f"={BS_('debt', c)}/{c}{{limit}}", None),
    ("回报", None),
    ("avg_eq", "平均股东权益", "(期初+期末)÷2", "money", lambda c, i: f"=AVERAGE({BS_('eq', prevc(i))},{BS_('eq', c)})", None),
    ("roe_q", "ROE（季度）", "净利润÷平均权益", "pct", lambda c, i: f"={IS_('ni', c)}/{c}{{avg_eq}}", None),
    ("roe_a", "ROE（年化）", "季度ROE×年化倍数", "pct", lambda c, i: f"={c}{{roe_q}}*ANNUAL", None),
    ("cum_roe", "累计回报（相对期初权益）", "累计净利润÷期初权益", "pct", lambda c, i: f"={IS_('cum_ni', c)}/'资产负债表'!$B${BS['eq']}", None),
    ("eq_x", "股东权益 ÷ 期初权益", "倍", "x", lambda c, i: f"={BS_('eq', c)}/'资产负债表'!$B${BS['eq']}", None),
]
r = 5
for row in AN_ROWS:
    if row[1] is None:
        section(wn, r, row[0], width=C0 + NQ); r += 1; continue
    AN[row[0]] = r; r += 1
for row in AN_ROWS:
    if row[1] is None:
        continue
    key, lab, unit, nf, f, tot = row
    rr = AN[key]
    label(wn, rr, lab, unit)
    for i, c in enumerate(QC):
        fx = f(c, i)
        if fx is None:
            wn[f"{c}{rr}"] = "—"; wn[f"{c}{rr}"].fill = F_NA; wn[f"{c}{rr}"].alignment = Alignment(horizontal="center")
            continue
        wn[f"{c}{rr}"] = fx.replace("{rev}", str(AN.get("rev", 0))).format(**{k: v for k, v in AN.items()}) if "{" in fx else fx
    style_row(wn, rr, "calc", nf, c_to=C0 + NQ)
    if wn[f"{FIRST}{rr}"].value == "—":
        wn[f"{FIRST}{rr}"].fill = F_NA
    if tot:
        wn[f"{TOT}{rr}"] = tot.format(r=rr)
    else:
        wn[f"{TOT}{rr}"] = "—"; wn[f"{TOT}{rr}"].fill = F_NA; wn[f"{TOT}{rr}"].alignment = Alignment(horizontal="center")
wn.column_dimensions["B"].width = 30
# 全程汇总
r = max(AN.values()) + 2
section(wn, r, "全程汇总（关键结果）", width=4); r += 1
SUMROWS = [
    ("期末股东权益", f"='资产负债表'!{LAST}{BS['eq']}", "money"),
    ("期初股东权益", f"='资产负债表'!B{BS['eq']}", "money"),
    ("16季累计净利润", f"='利润表'!{TOT}{IS['ni']}", "money"),
    ("16季累计营业收入", f"='利润表'!{TOT}{IS['rev']}", "money"),
    ("  其中 HBM 收入", f"='利润表'!{TOT}{IS['rev_hbm']}", "money"),
    ("全程毛利率", f"='利润表'!{TOT}{IS['gp']}/'利润表'!{TOT}{IS['rev']}", "pct"),
    ("16季累计经营现金流", f"='现金流量表'!{TOT}{CF['ocf']}", "money"),
    ("16季累计资本开支", f"=-'现金流量表'!{TOT}{CF['capex']}", "money"),
    ("16季累计自由现金流", f"='现金流量表'!{TOT}{CF['fcf']}", "money"),
    ("最低季末现金", f"=MIN('资产负债表'!{FIRST}{BS['cash']}:{LAST}{BS['cash']})", "money"),
    ("最高季末借款", f"=MAX('资产负债表'!{FIRST}{BS['debt']}:{LAST}{BS['debt']})", "money"),
    ("最高净负债", f"=MAX({FIRST}{AN['nd']}:{LAST}{AN['nd']})", "money"),
    ("紧急授信累计提取", f"=SUM('现金流量表'!{FIRST}{CF['emerg']}:{LAST}{CF['emerg']})", "money"),
    ("季度复合权益增长率", f"=('资产负债表'!{LAST}{BS['eq']}/'资产负债表'!B{BS['eq']})^(1/{NQ})-1", "pct"),
    ("折合年化权益增长率", "=(1+{prev})^ANNUAL-1", "pct"),
    ("扩产新增产能（千片/季）", f"='原始数据'!{LAST}{RAW['capacity_kwafers']}-'资产负债表'!B{BS['cap']}", "kw"),
    ("HBM 累计作废（千片）", f"={TOT}{AN['hbm_scrap']}", "kw"),
    ("普通内存累计需求缺口（千片）", f"={TOT}{AN['gap']}", "kw"),
]
SUMMARY = {}
for lab, fx, nf in SUMROWS:
    wn.cell(r, 1, lab).font = FT_CALC; wn.cell(r, 1).border = BORDER
    if "{prev}" in fx:
        fx = fx.format(prev=f"B{r-1}")
    c = wn.cell(r, 2, fx); c.fill = F_KEY; c.font = FT_BOLD; c.number_format = NF[nf]; c.border = BORDER
    SUMMARY[lab] = r
    r += 1
wn.freeze_panes = "C5"
widths(wn, a=30, b=30, q=9.5, extra={TOT: 11})
CHART_TOP = r + 2

# ---------- 图表（数据都来自上面的公式行；单轴，不用双Y轴） ----------
PAL = {"blue": "2A78D6", "orange": "EB6834", "aqua": "1BAF7A", "violet": "4A3AA7", "grid": "E1E0D9", "axis": "C3C2B7", "ink": "52514E"}
def _axis_style(ch, ytitle, nf="General"):
    ch.y_axis.title = ytitle
    ch.x_axis.delete = False; ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=PAL["grid"], w=6350))
    ch.y_axis.number_format = nf
    ch.y_axis.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=PAL["axis"], w=9525))
    if ch.legend is not None:
        ch.legend.position = "b"
    ch.x_axis.tickLblPos = "low"      # 有负值时类别标签放在图底部，不压在零线上
    ch.width, ch.height = 17.5, 8.2
def _cats():
    return Reference(wn, min_col=C0, max_col=C0 + NQ - 1, min_row=4)
def _series(row, name, color, line=False, ws=wn):
    s = Series(Reference(ws, min_col=C0, max_col=C0 + NQ - 1, min_row=row), title=name)
    if line:
        s.graphicalProperties.line.solidFill = color; s.graphicalProperties.line.width = 25400; s.smooth = False
        s.marker.symbol = "circle"; s.marker.size = 5
        s.marker.graphicalProperties = GraphicalProperties(solidFill=color, ln=LineProperties(solidFill="FFFFFF"))
    else:
        s.graphicalProperties.solidFill = color; s.graphicalProperties.line.solidFill = "FFFFFF"
    return s

charts = []
# 1 营收结构与净利润（同一单位：十亿美元）
rs_row = 4  # 利用利润表行直接作图
c1 = BarChart(); c1.type = "col"; c1.grouping = "stacked"; c1.overlap = 100; c1.gapWidth = 60
c1.title = "营收结构与净利润（十亿美元）"
c1.series.append(_series(IS["rev_std"], "普通内存收入", PAL["blue"], ws=wi))
c1.series.append(_series(IS["rev_hbm"], "HBM 收入", PAL["violet"], ws=wi))
l1 = LineChart(); l1.series.append(_series(IS["ni"], "净利润", PAL["orange"], line=True, ws=wi))
c1.set_categories(Reference(wi, min_col=C0, max_col=C0 + NQ - 1, min_row=4)); l1.set_categories(Reference(wi, min_col=C0, max_col=C0 + NQ - 1, min_row=4))
_axis_style(c1, "十亿美元", "0"); c1 += l1
charts.append(c1)
# 2 毛利率与年化ROE（同一单位：%）
c2 = LineChart(); c2.title = "毛利率与年化 ROE"
c2.series.append(_series(AN["gm"], "毛利率", PAL["blue"], line=True)); c2.series.append(_series(AN["roe_a"], "ROE（年化）", PAL["orange"], line=True))
c2.set_categories(_cats()); _axis_style(c2, "%", "0%"); charts.append(c2)
# 3 现金、借款与股东权益
c3 = LineChart(); c3.title = "现金、借款与股东权益（十亿美元）"
c3.series.append(_series(BS["eq"], "股东权益", PAL["blue"], line=True, ws=wbs))
c3.series.append(_series(BS["debt"], "借款", PAL["orange"], line=True, ws=wbs))
c3.series.append(_series(BS["cash"], "现金", PAL["aqua"], line=True, ws=wbs))
c3.set_categories(Reference(wbs, min_col=C0, max_col=C0 + NQ - 1, min_row=4)); _axis_style(c3, "十亿美元", "0"); charts.append(c3)
# 4 库存（千片）
c4 = BarChart(); c4.type = "col"; c4.title = "期末库存（千片）"; c4.gapWidth = 60
c4.series.append(_series(BS["kw"], "期末库存", PAL["blue"], ws=wbs)); c4.legend = None
c4.set_categories(Reference(wbs, min_col=C0, max_col=C0 + NQ - 1, min_row=4)); _axis_style(c4, "千片", "0"); c4.legend = None; charts.append(c4)
# 5 存货周转天数
c5 = LineChart(); c5.title = "存货周转天数（金额口径，天）"
c5.series.append(_series(AN["dio"], "存货周转天数", PAL["blue"], line=True)); c5.set_categories(_cats()); _axis_style(c5, "天", "0"); c5.legend = None; charts.append(c5)
# 6 经营现金流与自由现金流
c6 = BarChart(); c6.type = "col"; c6.title = "经营现金流与自由现金流（十亿美元）"; c6.gapWidth = 60
c6.series.append(_series(AN["ocf"], "经营活动现金流", PAL["blue"])); c6.series.append(_series(AN["fcf"], "自由现金流", PAL["orange"]))
c6.set_categories(_cats()); _axis_style(c6, "十亿美元", "0"); charts.append(c6)
# 7 净负债
c7 = BarChart(); c7.type = "col"; c7.title = "净负债（负数=净现金，十亿美元）"; c7.gapWidth = 60
c7.series.append(_series(AN["nd"], "净负债", PAL["blue"])); c7.set_categories(_cats()); _axis_style(c7, "十亿美元", "0"); c7.legend = None; charts.append(c7)
# 8 价格指数
c8 = LineChart(); c8.title = "价格指数（1.00 = 每千片 0.024 十亿美元）"
c8.series.append(_series(RAW["price_index"], "普通内存价格指数", PAL["blue"], line=True, ws=wr))
c8.series.append(_series(RAW["hbm_price_index"], "HBM 价格指数", PAL["violet"], line=True, ws=wr))
c8.set_categories(Reference(wr, min_col=C0, max_col=C0 + NQ - 1, min_row=4)); _axis_style(c8, "指数", "0.0"); charts.append(c8)

for k, ch in enumerate(charts):
    col = "A" if k % 2 == 0 else "J"
    row = CHART_TOP + (k // 2) * 18
    wn.add_chart(ch, f"{col}{row}")
wn.cell(CHART_TOP - 1, 1, "图表（数据随上方公式自动更新；所有图均为单一纵轴）").font = FT_BOLD

# =====================================================================
# 8. 模型复现与情景
# =====================================================================
wm = wb.create_sheet("模型复现与情景")
title(wm, "模型复现与情景（按推断参数从头重算）", "黄色情景格留空=使用实际决策；填入数字即覆盖该季决策。价格系数在「假设」页。HBM 产量=产量×HBM占比（假设开工率同样作用于HBM）；现金低于下限时紧急授信补足到下限（假设）。")
header(wm, 4, "项目〔算法〕", "期初", opening=True)
wm.cell(4, 2).value = "期初"
M = {}
r = 5
def mrow(key, lab, unit=""):
    global r
    M[key] = r; label(wm, r, lab + (f"  〔{unit}〕" if unit else ""), ""); r += 1
section(wm, r, "情景输入（留空=实际决策）"); r += 1
for key, lab in [("o_util", "开工率（情景）"), ("o_sell", "出货比例（情景）"), ("o_capex", "资本开支（情景）"), ("o_share", "HBM 产能占比（情景）"), ("o_borrow", "新增借款（情景）")]:
    mrow(key, lab, "留空=实际")
    for c in QC:
        cell = wm[f"{c}{M[key]}"]; cell.fill = F_IN; cell.font = FT_IN; cell.border = BORDER
        cell.number_format = NF["pct"] if key in ("o_util", "o_sell", "o_share") else NF["money"]
section(wm, r, "实际生效的决策"); r += 1
for key, lab, rawkey, nf in [("util", "开工率", "utilization", "pct"), ("sell", "出货比例", "sell_fraction", "pct"), ("capex", "资本开支", "capex", "money"),
                             ("share", "HBM 产能占比", "hbm_share", "pct"), ("borrow", "新增借款", "borrow", "money")]:
    mrow(key, lab)
    ok = "o_" + key
    for c in QC:
        wm[f"{c}{M[key]}"] = f'=IF({c}{M[ok]}="",{R(rawkey, c)},{c}{M[ok]})'
    style_row(wm, M[key], "calc", nf)
section(wm, r, "产能、产量与成本"); r += 1
for key, lab, unit, nf in [("cap", "产能（千片/季）", "300+已投产扩产", "kw"), ("dep", "折旧", "基础+已投产资本开支÷折旧期", "money"),
                           ("prod", "产量（千片）", "产能×开工率", "kw"), ("qual", "HBM 量产资格", "取实际", "int"), ("hbm", "HBM 产量（千片）", "产量×HBM占比×资格", "kw"),
                           ("std", "普通内存产量（千片）", "产量−HBM", "kw"), ("uc", "单位成本（十亿美元/千片）", "(折旧+工厂费用+变动成本×产量)÷产量", "money3")]:
    mrow(key, lab, unit)
section(wm, r, "存货与销售"); r += 1
for key, lab, unit, nf in [("beg_kw", "期初库存（千片）", "", "kw"), ("beg_inv", "期初存货金额", "", "money"), ("pool_kw", "可售片数", "期初+普通产量", "kw"),
                           ("pool_v", "可售成本池", "期初金额+普通产量×单位成本", "money"), ("avg", "加权平均成本（十亿美元/千片）", "", "money3"),
                           ("sold", "普通内存销量（千片）", "MIN(出货比例×可售,需求)", "kw"), ("hsold", "HBM 销量（千片）", "MIN(HBM产量,HBM需求)", "kw"),
                           ("end_kw", "期末库存（千片）", "", "kw"), ("nrv", "可变现净值（十亿美元/千片）", "(1−销售费用率)×价格", "money3"),
                           ("wd", "存货跌价", "MAX(0,期末片×(平均成本−净值))", "money"), ("end_inv", "期末存货金额", "", "money")]:
    mrow(key, lab, unit)
section(wm, r, "利润"); r += 1
for key, lab, unit in [("rev", "营业收入", ""), ("cogs", "营业成本（含跌价）", ""), ("rnd", "研发费用", "固定+HBM研发"), ("oi", "营业利润", ""),
                       ("int", "净利息支出", "借款×1.5%+紧急×3%−现金×0.8%"), ("pbt", "税前利润", ""), ("tax", "所得税", ""), ("ni", "净利润", "")]:
    mrow(key, lab, unit)
section(wm, r, "现金与资本"); r += 1
for key, lab, unit in [("ocf", "经营活动现金流", ""), ("cash_pre", "期末现金（紧急授信前）", ""), ("emb", "紧急授信提取", "补足到下限"),
                       ("cash", "期末现金", ""), ("debt", "期末借款", ""), ("em", "紧急授信余额", ""), ("eq", "期末股东权益", "")]:
    mrow(key, lab, unit)
# 期初列
for key, val, nf in [("beg_kw", None, "kw"), ("end_kw", O["inventory_kwafers"], "kw"), ("end_inv", O["inventory"], "money"), ("cash", O["cash"], "money"),
                     ("debt", O["debt"], "money"), ("em", O["emergency_debt"], "money"), ("eq", O["equity"], "money"), ("cap", O["capacity_kwafers"], "kw")]:
    if val is None:
        continue
    cell = wm.cell(M[key], 2, f"='资产负债表'!B{BS[{'end_kw':'kw','end_inv':'inv','cash':'cash','debt':'debt','em':'emerg','eq':'eq','cap':'cap'}[key]]}")
    cell.fill = F_CALC; cell.font = FT_CALC; cell.number_format = NF[nf]; cell.border = BORDER
for i, c in enumerate(QC):
    t = i + 1
    p = QC[i - 1] if i else "B"
    if t > 4:
        lagc = QC[t - 5]   # 第 t-4 季所在列
        wm[f"{c}{M['cap']}"] = f"=$B${M['cap']}+(SUM(${FIRST}${M['capex']}:{lagc}{M['capex']})-MAINT*{t-4})/CAP_COST"
        wm[f"{c}{M['dep']}"] = f"=DEP_BASE+SUM(${FIRST}${M['capex']}:{lagc}{M['capex']})/DEP_LIFE"
    else:
        wm[f"{c}{M['cap']}"] = f"=$B${M['cap']}"
        wm[f"{c}{M['dep']}"] = "=DEP_BASE"
    wm[f"{c}{M['prod']}"] = f"={c}{M['cap']}*{c}{M['util']}"
    wm[f"{c}{M['qual']}"] = f"={R('hbm_qualified', c)}"
    wm[f"{c}{M['hbm']}"] = f"={c}{M['prod']}*{c}{M['share']}*{c}{M['qual']}"
    wm[f"{c}{M['std']}"] = f"={c}{M['prod']}-{c}{M['hbm']}"
    wm[f"{c}{M['uc']}"] = f"=({c}{M['dep']}+PLANT_FIXED+V_COST*{c}{M['prod']})/{c}{M['prod']}"
    wm[f"{c}{M['beg_kw']}"] = f"={p}{M['end_kw']}"
    wm[f"{c}{M['beg_inv']}"] = f"={p}{M['end_inv']}"
    wm[f"{c}{M['pool_kw']}"] = f"={c}{M['beg_kw']}+{c}{M['std']}"
    wm[f"{c}{M['pool_v']}"] = f"={c}{M['beg_inv']}+{c}{M['std']}*{c}{M['uc']}"
    wm[f"{c}{M['avg']}"] = f"=IF({c}{M['pool_kw']}>0,{c}{M['pool_v']}/{c}{M['pool_kw']},0)"
    wm[f"{c}{M['sold']}"] = f"=MIN({c}{M['sell']}*{c}{M['pool_kw']},{R('demand_kwafers', c)})"
    wm[f"{c}{M['hsold']}"] = f"=MIN({c}{M['hbm']},{R('hbm_demand_kwafers', c)})"
    wm[f"{c}{M['end_kw']}"] = f"={c}{M['pool_kw']}-{c}{M['sold']}"
    wm[f"{c}{M['nrv']}"] = f"=(1-SELL_COST)*{R('price_index', c)}*PRICE_SCEN*P_BASE"
    wm[f"{c}{M['wd']}"] = f"=IF({c}{M['end_kw']}>0.000001,MAX(0,{c}{M['end_kw']}*({c}{M['avg']}-{c}{M['nrv']})),0)"
    wm[f"{c}{M['end_inv']}"] = f"={c}{M['end_kw']}*{c}{M['avg']}-{c}{M['wd']}"
    wm[f"{c}{M['rev']}"] = f"={c}{M['sold']}*{R('price_index', c)}*PRICE_SCEN*P_BASE+{c}{M['hsold']}*{R('hbm_price_index', c)}*HBM_PRICE_SCEN*P_BASE"
    wm[f"{c}{M['cogs']}"] = f"={c}{M['sold']}*{c}{M['avg']}+{c}{M['hbm']}*{c}{M['uc']}+{c}{M['wd']}"
    wm[f"{c}{M['rnd']}"] = f"=RND_BASE+{R('hbm_rnd', c)}"
    wm[f"{c}{M['oi']}"] = f"={c}{M['rev']}-{c}{M['cogs']}-{c}{M['rnd']}-SGA_Q"
    wm[f"{c}{M['int']}"] = f"=RATE_DEBT*{p}{M['debt']}+RATE_EMERG*{p}{M['em']}-RATE_CASH*{p}{M['cash']}"
    wm[f"{c}{M['pbt']}"] = f"={c}{M['oi']}-{c}{M['int']}"
    wm[f"{c}{M['tax']}"] = f"=TAX_RATE*MAX(0,{c}{M['pbt']})"
    wm[f"{c}{M['ni']}"] = f"={c}{M['pbt']}-{c}{M['tax']}"
    wm[f"{c}{M['ocf']}"] = f"={c}{M['ni']}+{c}{M['dep']}-({c}{M['end_inv']}-{c}{M['beg_inv']})"
    wm[f"{c}{M['cash_pre']}"] = f"={p}{M['cash']}+{c}{M['ocf']}-{c}{M['capex']}+{c}{M['borrow']}"
    wm[f"{c}{M['emb']}"] = f"=MAX(0,EMERG_FLOOR-{c}{M['cash_pre']})"
    wm[f"{c}{M['cash']}"] = f"={c}{M['cash_pre']}+{c}{M['emb']}"
    wm[f"{c}{M['debt']}"] = f"={p}{M['debt']}+{c}{M['borrow']}"
    wm[f"{c}{M['em']}"] = f"={p}{M['em']}+{c}{M['emb']}"
    wm[f"{c}{M['eq']}"] = f"={p}{M['eq']}+{c}{M['ni']}"
NFM = {"cap": "kw", "dep": "money", "prod": "kw", "qual": "int", "hbm": "kw", "std": "kw", "uc": "money3", "beg_kw": "kw", "beg_inv": "money", "pool_kw": "kw",
       "pool_v": "money", "avg": "money3", "sold": "kw", "hsold": "kw", "end_kw": "kw", "nrv": "money3", "wd": "money", "end_inv": "money", "rev": "money",
       "cogs": "money", "rnd": "money", "oi": "money", "int": "money3", "pbt": "money", "tax": "money", "ni": "money", "ocf": "money", "cash_pre": "money",
       "emb": "money", "cash": "money", "debt": "money", "em": "money", "eq": "money"}
for key, nf in NFM.items():
    style_row(wm, M[key], "key" if key in ("ni", "eq", "cash") else "calc", nf)
section(wm, r, "核对：模型 − 引擎（情景格留空、价格系数=1时应约为0）"); r += 1
for key, lab, f in [("chk_rev", "营业收入差额", lambda c: f"={c}{M['rev']}-{R('is_revenue', c)}"),
                    ("chk_ni", "净利润差额", lambda c: f"={c}{M['ni']}-{R('is_net_income', c)}"),
                    ("chk_cash", "期末现金差额", lambda c: f"={c}{M['cash']}-{R('bs_cash', c)}"),
                    ("chk_eq", "股东权益差额", lambda c: f"={c}{M['eq']}-{R('bs_equity', c)}"),
                    ("chk_kw", "期末库存片数差额", lambda c: f"={c}{M['end_kw']}-{R('bs_inventory_kwafers', c)}")]:
    mrow(key, lab, "应约为0")
    for c in QC:
        wm[f"{c}{M[key]}"] = f(c)
    style_row(wm, M[key], "chk", "chk")
r += 1
section(wm, r, "情景结果", width=4); r += 1
for lab, fx, nf in [("模型期末股东权益", f"={LAST}{M['eq']}", "money"), ("实际期末股东权益", f"={R('bs_equity', LAST)}", "money"),
                    ("差额（情景 − 实际）", f"=B{r}-B{r+1}", "money"), ("模型16季累计营业收入", f"=SUM({FIRST}{M['rev']}:{LAST}{M['rev']})", "money"),
                    ("模型最低季末现金", f"=MIN({FIRST}{M['cash']}:{LAST}{M['cash']})", "money"), ("模型紧急授信累计提取", f"=SUM({FIRST}{M['emb']}:{LAST}{M['emb']})", "money"),
                    ("模型期末库存（千片）", f"={LAST}{M['end_kw']}", "kw")]:
    wm.cell(r, 1, lab).font = FT_CALC; wm.cell(r, 1).border = BORDER
    c = wm.cell(r, 2, fx); c.fill = F_KEY; c.font = FT_BOLD; c.number_format = NF[nf]; c.border = BORDER
    M["res_" + lab] = r; r += 1
r += 1
wm.cell(r, 1, "用法示例：在「开工率（情景）」行的 D～G 列（第2-5季）填 0.7，可以看到下行期减产对期末权益的影响（复盘报告里的情景S1）。").font = FT_NOTE
wm.freeze_panes = "C5"; widths(wm, a=52, b=10, q=9.5)

# =====================================================================
# 9. 决策备注
# =====================================================================
wd_ = wb.create_sheet("决策备注")
title(wd_, "每季决策与备注原文", "来源：export_final.json → history[i].decision；数字为输入格（原样）。")
hdrs = ["季度", "开工率", "出货比例", "资本开支", "HBM研发", "HBM占比", "借款", "提交给引擎的备注原文"]
for j, h in enumerate(hdrs, start=1):
    cell = wd_.cell(4, j, h); cell.fill = F_HDR; cell.font = FT_HDR
for i, q in enumerate(Q):
    rr = 5 + i
    wd_.cell(rr, 1, f"Q{i+1}").font = FT_BOLD
    for j, (k, nf) in enumerate([("utilization", "pct"), ("sell_fraction", "pct"), ("capex", "money"), ("hbm_rnd", "money"), ("hbm_share", "pct"), ("borrow", "money")], start=2):
        cell = wd_.cell(rr, j, q[k]); cell.fill = F_IN; cell.font = FT_IN; cell.number_format = NF[nf]; cell.border = BORDER
    cell = wd_.cell(rr, 8, q["note"]); cell.alignment = Alignment(wrap_text=True, vertical="top"); cell.font = FT_CALC
    wd_.row_dimensions[rr].height = 120
for k, w in zip("ABCDEFGH", [7, 8, 9, 9, 9, 9, 8, 110]):
    wd_.column_dimensions[k].width = w

# 打印设置：横向、按页宽缩放
for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
# 打开时重算
wb.calculation.fullCalcOnLoad = True
out = ROOT / "三张表.xlsx"
wb.save(out)
print("saved", out)
