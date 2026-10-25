import sys,json,re,math,zipfile
from pathlib import Path
sys.path.insert(0,'交付/工作底稿/依赖')
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries,get_column_letter
from docx import Document
P=Path('交付/伯克希尔账本.xlsx');wb=load_workbook(P,data_only=False);cached=load_workbook(P,data_only=True)
class Calc:
 def __init__(self,overrides=None):self.overrides=overrides or {};self.values={};self.stack=set()
 def get(self,s,c):
  c=c.replace('$','');key=(s,c)
  if key in self.overrides:return self.overrides[key]
  if key in self.values:return self.values[key]
  if key in self.stack:raise ValueError('Circular '+str(key))
  self.stack.add(key);v=wb[s][c].value
  if isinstance(v,str) and v.startswith('='):v=self.expr(v[1:],s)
  if v is None:v=0
  self.values[key]=v;self.stack.remove(key);return v
 def expr(self,x,s):
  if x.startswith('IF(') and x.endswith(')'):
   args=[];level=0;start=3;quote=False
   for i,ch in enumerate(x[3:-1],3):
    if ch=='"':quote=not quote
    if not quote:
     if ch=='(':level+=1
     if ch==')':level-=1
     if ch==',' and level==0:args.append(x[start:i]);start=i+1
   args.append(x[start:-1]);return self.expr(args[1] if self.expr(args[0],s) else args[2],s)
  def summed(m):
   a,b,c,d=range_boundaries(m[1]);return str(sum(self.get(s,f'{get_column_letter(i)}{j}') for i in range(a,c+1) for j in range(b,d+1)))
  x=re.sub(r'SUM\(([A-Z]+\d+:[A-Z]+\d+)\)',summed,x)
  x=re.sub(r"'([^']+)'!([A-Z]+\d+)",lambda m:repr(self.get(m[1],m[2])),x)
  x=re.sub(r'(?<![\w\'])\b([A-Z]+\d+)\b',lambda m:repr(self.get(s,m[1])),x)
  return eval(x,{'__builtins__':{}},{'ABS':abs})
calc=Calc();count=0;errors=[]
for s in wb:
 for row in s:
  for c in row:
   if c.data_type=='f':
    count+=1
    try:
     v=calc.get(s.title,c.coordinate);cv=cached[s.title][c.coordinate].value
     if isinstance(v,(int,float)):
      if not isinstance(cv,(int,float)) or not math.isclose(v,cv,rel_tol=1e-11,abs_tol=1e-7):errors.append([s.title,c.coordinate,v,cv])
     elif v!=cv:errors.append([s.title,c.coordinate,v,cv])
    except Exception as e:errors.append([s.title,c.coordinate,str(e)])
   elif c.data_type=='e':errors.append([s.title,c.coordinate,c.value])
for sheet,rows in [('利润',[18]),('股票组合',[11])]:
 for r in rows:
  for col in 'BCDEFGH':assert calc.get(sheet,f'{col}{r}')==0
expected=[32466,34775,43637,53936,51714,24993,26743]
for col,v in zip('BCDEFGH',expected):assert calc.get('法定分部',f'{col}14')==v
perturb=Calc({('假设','B6'):-.02});base=calc.get('假设','B13');new=perturb.get('假设','B13');assert math.isclose(new,base*2);assert math.isclose(perturb.get('变化','G6'),new)
for file in ['交付/伯克希尔账本.xlsx','交付/研究备忘录.docx']:
 with zipfile.ZipFile(file) as z:assert z.testzip() is None
memo=Document('交付/研究备忘录.docx');assert len(memo.paragraphs)>40;assert len(memo.tables)==4
result={'xlsx_opened':True,'sheets':wb.sheetnames,'formulas_evaluated':count,'formula_errors':errors,'cached_results_match':not errors,'balance_checks':'利润归母对账 / 行业组合 / 法定分部合计通过','assumption_test':{'yield_change_before':-.01,'yield_change_after':-.02,'after_tax_before':base,'after_tax_after':new,'dashboard_updates':True},'docx_opened':True,'docx_paragraphs':len(memo.paragraphs),'docx_tables':len(memo.tables),'scope':'使用openpyxl/python-docx打开包；独立公式求值器复算全部公式，不声称已用Microsoft Office GUI或LibreOffice验证分页。'}
Path('交付/验收/财务与文档验收.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False,indent=2));assert not errors
