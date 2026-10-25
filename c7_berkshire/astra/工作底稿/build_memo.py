import sys,json
from pathlib import Path
sys.path.insert(0,'交付/工作底稿/依赖')
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
D=json.loads(Path('交付/工作底稿/financial_data.json').read_text());doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.1);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
style=doc.styles['Normal'];style.font.name='Arial';style._element.rPr.rFonts.set(qn('w:eastAsia'),'等线');style.font.size=Pt(10.5);style.font.color.rgb=RGBColor.from_string('28434A');style.paragraph_format.space_after=Pt(8);style.paragraph_format.line_spacing=1.2
for name,size,col in [('Title',31,'183D46'),('Heading 1',21,'183D46'),('Heading 2',13,'138B83'),('Heading 3',11,'138B83')]:
 s=doc.styles[name];s.font.name='Arial';s._element.rPr.rFonts.set(qn('w:eastAsia'),'等线');s.font.size=Pt(size);s.font.color.rgb=RGBColor.from_string(col)
header=sec.header.paragraphs[0];header.text='THE LONG VIEW  /  伯克希尔研究';header.style='Caption';header.runs[0].font.color.rgb=RGBColor.from_string('138B83')
foot=sec.footer.paragraphs[0];foot.text='研究展示 · 不构成投资建议                                      ';field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
def p(text,style=None):return doc.add_paragraph(text,style)
def h(text):doc.add_heading(text,2)
def title(kicker,text):p(kicker,'Subtitle');doc.add_heading(text,0)
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1, cols=len(headers));t.autofit=False
 for i,x in enumerate(headers):t.rows[0].cells[i].text=x
 for row in rows:
  cells=t.add_row().cells
  for c,x in zip(cells,row):c.text=str(x)
 for ri,row in enumerate(t.rows):
  for ci,c in enumerate(row.cells):
   tcPr=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'183D46' if ri==0 else ('F0F6F3' if ri%2 else 'FFFFFF'));tcPr.append(shade)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(6);pp.paragraph_format.space_before=Pt(6)
    for r in pp.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '28434A');r.bold=ri==0
   if widths:c.width=Cm(widths[ci])
  trPr=row._tr.get_or_add_trPr();el=OxmlElement('w:cantSplit');trPr.append(el)
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'));return t
def page():doc.add_page_break()
ops=[sum(D[k][j] for k in ['underwriting','insurance_income','rail','energy','msr','other']) for j in range(7)]
cash=[D['cash_ins'][j]+D['cash_rail'][j]+D['bills'][j] for j in range(7)]
bv=[D['equity'][j]*1e6/(D['A'][j]+D['B'][j]/1500) for j in range(7)]
title('研究备忘录  /  2026 H1','交班之后，\n账本先说话。')
p('伯克希尔·哈撒韦｜2021—2025 年度与 2026 年上半年','Subtitle')
p('研究基准：提供的2025年报、2026年二季报及业绩公告；补充历史资料访问日：2026年10月22日。观察窗口截止2026年6月30日，另标明已披露期后事项。')
h('核心判断')
p('交接后的第一个半年，同时出现了经营利润上升、股票投资由净卖出转为净买入，以及恢复回购。现金与短期国债仍然庞大。不过，汇兑反转解释了经营利润增量的大部分；此前签订的收购也在这一时期交割。因此，账本显示资本开始更积极地投入，却不足以把所有变化归因于新任首席执行官。')
table(['观察指标','2025H1','2026H1','读法'],[
 ['经营分项利润','20,801','24,329',f'同比 +{(ops[6]/ops[5]-1)*100:.1f}%'],
 ['现金＋短期国债','344,091','365,514','合并期末口径'],
 ['权益证券净买入','−4,500','11,625','负数表示净卖出'],
 ['回购入账额','0','4,762','现金支付为4,444'],
], [4.5,3,3,5.5])
p('表内金额：百万美元。利润为税后归属股东；现金含铁路能源，不扣未结算国债应付款。来源：L2损益、资产负债表及现金流量表；2025H1现金补充E3。', 'Caption')
h('三条贯穿全文的线索')
p('利润：先分经营与投资，再分业务改善与汇兑影响。\n资金：现金、浮存金和股东权益承担不同角色，不能相加成“可花的钱”。\n交接：区分签约、交割和报表确认的时间，再讨论管理层的资本配置。')
p('这是一份研究与历史学习材料，不提供买卖、持有、目标价或股价贵贱判断。小游戏的契合度也不是投资能力评分。')
page();doc.add_heading('01  增长是真的，成分也重要',1)
h('先搭好利润的桥')
p('2026H1归母净利润357.73亿美元，由经营分项利润243.29亿美元与投资收益114.44亿美元组成。2025H1的169.73亿美元净利润则受到37.60亿美元权益法投资减值和0.68亿美元投资亏损影响。净利润同比大幅跳升，不能直接等同于经营能力同步提高。[L2，Results of Operations；L3，业绩汇总]')
rows=[]
for k,label in [('underwriting','保险承保'),('insurance_income','保险投资收益'),('rail','BNSF'),('energy','BHE'),('msr','制造服务零售'),('other','其他')]:
 a,b=D[k][5:7];rows.append([label,f'{a:,}',f'{b:,}',f'{(b/a-1)*100:+.1f}%',f'{b-a:+,}'])
table(['百万美元','2025H1','2026H1','同比','增量'],rows,[4.7,3,3,2.5,2.8])
p('来源：L2 Item 2 Results of Operations，p33；与L3业绩公告交叉核对。', 'Caption')
h('增量中，汇兑占了约六成')
p('经营分项利润增加35.28亿美元。“其他”增加24.61亿美元，其中外币债务汇兑从亏损15.90亿美元转为收益5.75亿美元，反转21.65亿美元，约占全部经营增量的61.4%。仅扣除这项汇兑影响，两期经营分项利润为223.91亿和237.54亿美元，增幅约6.1%。这只是研究拆解，仍包含投资收入、收购并表及其他因素，不是有机增长率。[L2，Other，p47；账本「对账」公式]')
h('把五年放在一起')
p('统一分类、使用重述后数据的经营分项利润，2021至2025年分别为275.97亿、308.53亿、373.50亿、474.37亿、444.86亿美元。2025年下降约6.2%，主要受到保险承保和保险投资收益回落影响；因此，2026H1的增长也发生在经历调整之后。[L1，K-34；E2，K-35；账本「利润」]')
p('2021—2022数据使用2023年报ASU 2018-12重述版本；原非控制业务归入“其他”，2023年的Pilot税后收益归入制造服务零售。这样做是分类桥接，不是新增利润。')
page();doc.add_heading('02  现金还厚，投入的方向变了',1)
h('现金库略降，仍不等于“闲置余额”')
p(f'合并现金及等价物加单列短期国债，从2025年末3,733.11亿美元降至2026年6月末3,655.14亿美元，减少77.97亿美元，降幅约{(1-cash[6]/cash[4])*100:.1f}%。其中合并现金及等价物为406.09亿美元，单列短期国债为3,249.05亿美元，占该合计的88.9%。现金等价物中已经包含部分短期国债，不能重复加总。[L1/L2，Consolidated Balance Sheets]')
p('管理层另报的“保险及其他业务、扣除未结算国债应付款”口径为3,592亿美元，2025年末约3,690亿美元。它不含铁路能源现金，也扣除了特定应付款，不能与主表的合并毛额混用。所有现金也并非母公司随时可分配资金，子公司经营、保险赔付和监管需求仍然存在。[L1/L2，Financial Condition]')
table(['现金用途 / 百万美元','2025H1','2026H1'],[['经营现金净流入','20,988','21,653'],['资本支出（含租赁设备）','9,139','10,631'],['购入权益证券','7,092','39,405'],['出售权益证券所得','11,592','27,780'],['企业收购净支付','101','9,704'],['回购现金支付','0','4,444']],[8.5,4,4])
p('来源：L2 Consolidated Statements of Cash Flows。表为部分重要现金流，不构成完整现金余额变动桥。', 'Caption')
h('从净卖出转向净买入，但不要把市值当交易量')
p('2026H1权益证券净买入116.25亿美元；去年同期净卖出45亿美元。组合公允价值从2025年末2,977.78亿美元增至3,237.79亿美元。商业、工业及其他类升至1,288.16亿美元，占组合39.8%，高于年末33.1%；银行保险金融与消费品分别占29.2%与31.0%。这些是市值权重，受买卖与估值共同影响。[L2，Note 4及现金流量表；账本「股票组合」]')
p('组合采用Note 4行业分类，包含该表中的优先股与认股权证，排除Kraft Heinz、Occidental普通股等权益法投资。2022年Occidental普通股转入权益法影响历史分类可比性。')
page();doc.add_heading('03  保险、铁路、能源：分别看条件',1)
h('保险：承保小幅上升，仍须看损失组成')
p('承保利润34.48亿美元，同比增长3.6%。但2026H1没有重大灾害事件损失，而2025H1包含约8.50亿美元税后重大灾害损失。仅把这项损失加回2025H1，比较基数便从33.28亿变成41.78亿美元，高于本期。该观察不能取代完整的承保质量分析：准备金释放、业务结构和损失趋势仍需分别看。[L2，Results of Operations、Insurance—Underwriting]')
p('材料指出GEICO承保利润回落，再保险和其他初级保险有所抵消；部分再保险业绩也受以前事故年度损失估计下调支持。保险投资收益为57.38亿美元，下降8.3%，主要受到利率下降影响。此前保险子公司向母公司的资本分配，也改变了收入在保险投资与“其他”之间的归属，跨年分析须同时看两处。[L1，Insurance—Investment Income；L2，p33、p38及Other]')
p('6月末浮存金约1,775亿美元，季报披露较年末增加约11亿美元。年报只写约1,760亿美元，两个约数相减得到15亿美元，与季报增量不完全相同；账本保留原披露并注释差异。浮存金是待支付保单义务形成的可投资资金，不是额外的股东权益；承保盈利使其当期成本为负，仍不能保证未来。[L1，K-41；L2，Insurance—Investment Income]')
h('BNSF：运量和效率支撑利润，投入不可忽略')
p('BNSF税后利润29.35亿美元，同比增长9.5%。材料将改善归因于运量增加和经营效率提升，部分被燃料成本及有效税率上升抵消。铁路持续需要资本支出，所以利润不能全部当作可上缴母公司的现金。[L2，BNSF及Financial Condition；历史语境见2023年股东信BNSF段]')
h('BHE：利润改善，与长期负担并存')
p('BHE税后归属股东利润20.05亿美元，同比增长11.5%，主要来自美国公用事业和天然气管道业务。税前法定分部利润只有13.48亿美元，税后数更高并非录入错误：分部存在税收抵免形成的所得税收益，还须考虑少数股东归属。[L2，BHE；Note 24 Business segment data]')
p('PacifiCorp山火诉讼、监管成本回收及电网投入仍值得持续跟踪。BNSF与BHE上半年合计资本支出约67亿美元，材料预计2026年剩余时间还将投入约86亿美元；后者是管理层展望，不是已经发生的事实。[L2，Note 22 Commitments and contingencies；Financial Condition，p49]')
page();doc.add_heading('04  把交接与报表确认分开',1)
h('交割在新一年，决策未必都在新一年')
p('OxyChem于2026年1月2日交割，现金对价约94亿美元，但协议签于2025年10月1日。这项投入跨越交接期。OxyChem上半年收入约26亿美元、税前利润1.21亿美元，包含收购会计折旧摊销及过渡成本。不能仅凭半年并表利润判断长期成效，更不能把交易全部归入新任CEO的独立决策。[L2，Note 2 Business acquisitions；Manufacturing]')
p('材料还披露Taylor Morrison于2026年7月24日完成收购，现金对价约68亿美元。这是6月末之后的事项，未把该现金支出纳入2026H1模型。[L2，Note 2及Financial Condition]')
h('恢复回购，对每股权益的作用可计算')
p(f'2025全年未回购；2026H1入账回购47.62亿美元，其中第二季45.27亿美元。现金流量表支付44.44亿美元，差额3.18亿美元；缺少逐笔结算明细，本研究不把差额全部归因于某一种原因。季报的回购程序由CEO在与董事长协商后判断，不能仅凭结果推断某个人的独立决策。[L1，Financial Condition；L2，权益变动表、现金流量表及Note 18]')
p(f'归属股东权益从7,174.19亿美元增至7,479.10亿美元；按期末A股加B股/1500换算，每股A类账面价值从{bv[4]:,.2f}美元升至{bv[6]:,.2f}美元，增长约{(bv[6]/bv[4]-1)*100:.1f}%。B类对应约{bv[6]/1500:,.2f}美元。权益增长与股份减少共同影响每股数，账面价值仍不等于内在价值。[L1/L2，资产负债表、Common stock；账本公式]')
h('后续研究清单')
table(['问题','应继续观察的证据'],[['保险盈利的可持续性','保费定价、GEICO保单趋势、灾害损失与准备金释放分别贡献多少。'],['降息如何传导到利润','国债余额与实际收益率；保险投资收入与母公司投资收入的合并观察。'],['实业增长有多少来自并购','OxyChem经营与收购会计影响，工业制造、服务业务的可比经营数据。'],['能源资产的现金需求','山火赔偿与监管成本回收，资本支出、税收抵免及子公司负债。'],['资本配置有没有改变约束','股票净买入、并购、回购与现金余额共同变化，而非单看某一笔交易。']],[5.1,11.4])
page();doc.add_heading('05  方法、假设与证据边界',1)
h('可修改假设：利率下降的机械影响')
p('账本「假设」使用2026年6月末现金加国债3,655.14亿美元为固定基数，假设其中80%同步重定价、收益率下降100基点、作用半年、边际税率21%。由公式得到税前影响−14.62亿美元、税后−11.55亿美元，约为已披露上半年经营分项利润的−4.7%。这不是公司预测。它忽略资金流、期限结构、海外税差和管理层应对；改变橙色输入格会自动更新结果和变化页。')
h('能核验什么，不能推断什么')
p('已核对：六项利润合计与归母净利润桥、行业组合加总、期末股份换算、回购两种金额、游戏十段原文。全部关键输入在Excel中附文件及章节，集中列入「取数索引」。')
p('未独立核实：所给2025/2026材料的外部真实性；回购入账与现金支付差额的逐笔成因；全部股票的交易明细及未披露投资意图；浮存金精确年末未取整数。研究不将这些空白写成事实。部分早期HTML材料不可读，游戏避开这些文件并使用可核验原文。')
p('所给2025股东信由Gregory E. Abel署名。小游戏只从更早年份的巴菲特股东信选题；1967、1972、1994及2009等决定采用事后回顾并同时标注两种年份。2011年的110%账面价值回购上限仅用于历史题，不是现行政策说明。')
h('来源对照')
for a,b in [('L1','材料/Berkshire_2025年报__brka-20251231.htm。章节：Item 7、合并财务报表、Note 4、Common stock、Note 26。'),('L2','材料/Berkshire_2026Q2季报__brka-20260630.htm。章节：Item 2、合并财务报表、Note 2、4、18、22、24。'),('L3','材料/Berkshire_2026Q2业绩公告__d159922dex991.htm。表：经营分项、净利润、回购及浮存金。'),('E1','官方2022年报： https://www.berkshirehathaway.com/2022ar/2022ar.pdf'),('E2','官方2023年报： https://www.berkshirehathaway.com/2023ar/2023ar.pdf'),('E3','官方2025二季报： https://www.berkshirehathaway.com/qtrly/2ndqtr25.pdf')]:
 pp=p(a+'  '+b);pp.paragraph_format.space_after=Pt(6)
 for r in pp.runs:r.font.size=Pt(8.5)
p('补充网络来源访问日均为2026-10-22；数字仅用于填补本地材料未覆盖年份或比较期。股东信逐题来源见web/decisions.json及账本「历史决策」。','Caption')
doc.save('交付/研究备忘录.docx')
# An HTML reading proof is generated from the document content to inspect readable text and table flow.
import html
out=['<!doctype html><html lang="zh"><meta charset="utf-8"><title>备忘录阅读校样</title><style>body{max-width:900px;margin:50px auto;font-family:Arial,"PingFang SC",sans-serif;color:#28434a;line-height:1.9}h1{color:#183d46}h2{color:#138b83}table{border-collapse:collapse;width:100%;font-size:14px}td{padding:12px;border:1px solid #dce5df}tr:first-child{background:#183d46;color:white}.page{break-before:page;height:20px;border-top:1px solid #ddd;margin-top:40px}</style>']
from docx.table import Table
from docx.text.paragraph import Paragraph
for child in doc.element.body:
 if child.tag==qn('w:p'):
  pp=Paragraph(child,doc)
  if child.xpath('.//w:br[@w:type="page"]'):out.append('<div class="page"></div>')
  elif pp.text:
   tag='h1' if pp.style.name in ('Title','Heading 1') else 'h2' if pp.style.name=='Heading 2' else 'p';out.append(f'<{tag}>'+html.escape(pp.text).replace('\n','<br>')+f'</{tag}>')
 elif child.tag==qn('w:tbl'):
  tab=Table(child,doc);out.append('<table>'+''.join('<tr>'+''.join('<td>'+html.escape(cc.text)+'</td>' for cc in row.cells)+'</tr>' for row in tab.rows)+'</table>')
Path('交付/验收/备忘录阅读校样.html').write_text(''.join(out)+'</html>')
print('Memo created; BV A',bv[6], 'growth',bv[6]/bv[4]-1)
